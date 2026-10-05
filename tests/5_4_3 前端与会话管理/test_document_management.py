"""5.4.3 前端与会话管理：TestDocumentManagement。"""

import sys
from pathlib import Path

# 从其他目录直接运行测试文件时，也能定位 src 和共用样例。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import tempfile
import unittest
from unittest.mock import patch


class TestDocumentManagement(unittest.TestCase):
    """回收边界与实际加载任务验证；无需生成模型。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.raw = Path(self.directory.name) / "raw"
        self.index = Path(self.directory.name) / "index"
        self.data = b"Six encoder layers"
        import hashlib
        self.doc_id = hashlib.sha256(self.data).hexdigest()
        self.folder = self.raw / self.doc_id
        self.folder.mkdir(parents=True)
        (self.folder / "paper.txt").write_bytes(self.data)

    def test_unindexed_list_delete_restore_without_creating_database(self):
        from src.frontend.components.documents import list_documents, delete_document, restore_document
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("不可加载模型")):
            self.assertEqual(list_documents(self.raw, self.index)[0]["chunks"], 0)
            self.assertEqual(delete_document(self.raw, self.index, self.doc_id), 0)
            self.assertEqual(list_documents(self.raw, self.index), [])
            tasks = restore_document(self.raw, self.doc_id)
        self.assertEqual(tasks[0]["data"], self.data)
        self.assertEqual(tasks[0]["status"], "pending")
        self.assertFalse(self.index.exists())

    def test_reject_path_and_symlink(self):
        from src.frontend.components.documents import delete_document, list_documents
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, "../outside")
        link = self.raw / ("a" * 64)
        link.symlink_to(self.folder, target_is_directory=True)
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, link.name)
        self.assertEqual(len(list_documents(self.raw, self.index)), 1)
        self.assertTrue((self.folder / "paper.txt").exists())

    def test_archive_conflict_preserves_source(self):
        from src.frontend.components.documents import delete_document
        (self.raw / ".trash" / self.doc_id).mkdir(parents=True)
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, self.doc_id)
        self.assertEqual((self.folder / "paper.txt").read_bytes(), self.data)

    def test_restore_rejects_changed_content_and_no_overwrite(self):
        from src.frontend.components.documents import delete_document, restore_document
        delete_document(self.raw, self.index, self.doc_id)
        archived = self.raw / ".trash" / self.doc_id / "paper.txt"
        archived.write_bytes(b"Changed")
        with self.assertRaises(ValueError):
            restore_document(self.raw, self.doc_id)
        self.assertTrue(archived.exists())
        archived.write_bytes(self.data)
        self.folder.mkdir()
        with self.assertRaises(FileExistsError):
            restore_document(self.raw, self.doc_id)
        self.assertTrue(archived.exists())

    def test_missing_source_never_deletes_index(self):
        from src.frontend.components.documents import delete_document
        with patch("src.frontend.components.documents.VectorStore") as store:
            with self.assertRaises(FileNotFoundError):
                delete_document(self.raw, self.index, "b" * 64)
            store.assert_not_called()

    def test_read_complete_text_and_markdown_without_index(self):
        from src.frontend.components.documents import read_document_content
        with patch("src.frontend.components.documents.VectorStore", side_effect=AssertionError("不可读取索引")):
            parts = read_document_content(self.raw, self.doc_id)
        self.assertEqual([part.page_content for part in parts], [self.data.decode()])
        (self.folder / "paper.txt").rename(self.folder / "paper.md")
        self.assertEqual(read_document_content(self.raw, self.doc_id)[0].page_content, self.data.decode())

    def test_read_word_preserves_paragraph_and_table(self):
        from hashlib import sha256
        from docx import Document
        from src.frontend.components.documents import read_document_content
        path = Path(self.directory.name) / "word.docx"
        word = Document()
        word.add_paragraph("论文方法完整说明")
        table = word.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("指标", "结果"), ("准确率", "90%")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        word.save(path)
        identifier = sha256(path.read_bytes()).hexdigest()
        folder = self.raw / identifier
        folder.mkdir()
        path.rename(folder / path.name)
        parts = read_document_content(self.raw, identifier)
        self.assertEqual(parts[0].page_content, "论文方法完整说明")
        self.assertIn("90%", parts[1].page_content)
        self.assertEqual(parts[1].metadata["table_index"], 1)

    def test_read_pdf_keeps_both_pages(self):
        import pymupdf
        from hashlib import sha256
        from src.frontend.components.documents import read_document_content
        pdf = pymupdf.open()
        for text in ("First original page", "Second original page"):
            pdf.new_page().insert_text((72, 72), text)
        data = pdf.tobytes()
        pdf.close()
        identifier = sha256(data).hexdigest()
        folder = self.raw / identifier
        folder.mkdir()
        (folder / "paper.pdf").write_bytes(data)
        parts = read_document_content(self.raw, identifier)
        self.assertEqual([part.metadata["page_number"] for part in parts], [1, 2])
        self.assertEqual([part.page_content for part in parts], ["First original page", "Second original page"])

    def test_read_rejects_missing_changed_and_symlink_sources(self):
        from src.frontend.components.documents import read_document_content
        with self.assertRaises(ValueError):
            read_document_content(self.raw, "../outside")
        with self.assertRaises(FileNotFoundError):
            read_document_content(self.raw, "b" * 64)
        path = self.folder / "paper.txt"
        path.write_bytes(b"Changed original")
        with self.assertRaises(ValueError):
            read_document_content(self.raw, self.doc_id)
        path.unlink()
        path.symlink_to(Path(self.directory.name) / "outside.txt")
        with self.assertRaises(FileNotFoundError):
            read_document_content(self.raw, self.doc_id)

if __name__ == "__main__":
    unittest.main()
