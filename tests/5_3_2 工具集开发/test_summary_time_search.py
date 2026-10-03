"""5.3.2 工具集开发：TestSummaryTimeSearch。"""

import sys
from pathlib import Path

# 从其他目录直接运行测试文件时，也能定位 src 和共用样例。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
import tempfile
import unittest
from unittest.mock import patch
from urllib.error import URLError
from src.agent.react_loop import run_react
from src.agent.tools import AVAILABLE_TOOLS, execute_tool
from src.agent.tools import current_time, get_available_tools, paper_summary, web_search
from src.utils.config import load_config
from tests.helpers import isolated_agent_logs


def setUpModule():
    """保持既有 Agent 测试的模块级日志隔离。"""
    global _log_context
    _log_context = isolated_agent_logs()
    _log_context.__enter__()


def tearDownModule():
    _log_context.__exit__(None, None, None)


class TestSummaryTimeSearch(unittest.TestCase):
    """上传与引用使用真实实现；模型/外部HTTP响应为显式构造的协议样例。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.config = deepcopy(load_config())
        self.config["paths"]["raw_documents"] = self.directory.name
        self.addCleanup(patch.stopall)
        for module in ("src.agent.tools", "src.agent.react_loop", "src.generation.rag_pipeline", "src.chunking"):
            patch(module + ".load_config", return_value=self.config).start()
        self.doc_id = self.upload("论文.md", "背景：翻译任务。\n方法：Transformer。\n结果：BLEU 28.4。\n结论：可用于翻译。".encode())
        self.sections = {key: {"text": text, "reference_ids": [1]} for key, text in zip(
            ("background", "method", "results", "conclusion"), ("研究机器翻译。", "使用Transformer。", "BLEU为28.4。", "可用于翻译。"))}
        self.packet = {"model": "qwen2.5:7b", "done": True, "done_reason": "stop",
                       "prompt_eval_count": 400, "eval_count": 150, "message": {"content": json.dumps(self.sections)}}

    def upload(self, name, data):
        from src.data_loader import batch_import, create_import_tasks
        import hashlib
        tasks = create_import_tasks([(name, data)])
        list(batch_import(tasks, self.directory.name))
        self.assertEqual(tasks[0]["status"], "success", tasks[0]["error"])
        return hashlib.sha256(data).hexdigest()

    def summarize(self, packet=None, doc_id=None):
        with patch("src.generation.rag_pipeline.urlopen", return_value=BytesIO(json.dumps(
                self.packet if packet is None else packet).encode())) as http:
            result = paper_summary.invoke({"doc_id": doc_id or self.doc_id})
        return result, http

    def search(self, html, status=200):
        import httpx
        self.config["agent"]["online_search_enabled"] = True
        with patch("httpx.Client") as client:
            client.return_value.__enter__.return_value.get.return_value = httpx.Response(
                status, text=html, request=httpx.Request("GET", "https://html.duckduckgo.com/html/"))
            result = web_search.invoke({"query": "Transformer 论文"})
        return result, client

    def test_summary_generates_four_sections_with_true_sources_and_usage(self):
        result, http = self.summarize()
        self.assertEqual(result["sections"], self.sections)
        self.assertEqual(result["status"], "answered")
        self.assertIn("### 背景", result["answer"])
        self.assertIn("论文.md；行", result["answer"])
        self.assertEqual(result["citations"][0]["metadata"]["doc_id"], self.doc_id)
        self.assertNotIn("score", result["references"][0])
        self.assertFalse(result["input_truncated"])
        self.assertEqual(result["usage"]["eval_count"], 150)
        payload = json.loads(http.call_args.args[0].data)
        self.assertEqual(http.call_count, 1)
        self.assertEqual(payload["model"], self.config["llm"]["model"])
        self.assertEqual(payload["format"]["required"], list(self.sections))
        self.assertEqual(payload["messages"][0]["role"], "system")
        self.assertIn("Transformer", payload["messages"][1]["content"])
        self.assertEqual(payload["options"]["num_predict"], 512)

    def test_summary_preserves_pdf_physical_page_and_word_paragraph(self):
        import fitz
        from docx import Document
        with fitz.open() as pdf:
            pdf.new_page().insert_text((40, 50), "Transformer Translation BLEU 28.4 Conclusion")
            pdf_id = self.upload("论文.pdf", pdf.tobytes())
        word, stream = Document(), BytesIO()
        word.add_paragraph("Transformer 用于翻译，BLEU 28.4。")
        word.save(stream)
        word_id = self.upload("论文.docx", stream.getvalue())
        pdf_result, _ = self.summarize(doc_id=pdf_id)
        word_result, _ = self.summarize(doc_id=word_id)
        self.assertIn("第1页（物理页码）", pdf_result["answer"])
        self.assertIn("段落1", word_result["answer"])
        self.assertNotIn("page_number", word_result["citations"][0]["metadata"])

    def test_summary_prioritizes_conclusion_and_reports_truncation(self):
        doc_id = self.upload("长文.txt", ("Abstract\nTransformer用于翻译。\n" + "中间研究讨论。\n" * 1800 +
                                          "7 Conclusion\n最终结论保留标记：可用于翻译。\n").encode())
        result, http = self.summarize(doc_id=doc_id)
        self.assertTrue(result["input_truncated"])
        self.assertIn("最终结论保留标记", json.loads(http.call_args.args[0].data)["messages"][1]["content"])
        self.assertTrue(result["warnings"])
        self.assertTrue(all(ref["metadata"]["doc_id"] == doc_id for ref in result["references"]))

    def test_missing_section_is_explicit_and_cannot_complete_task(self):
        packet = deepcopy(self.packet)
        sections = deepcopy(self.sections)
        sections["results"] = {"text": "", "reference_ids": []}
        packet["message"]["content"] = json.dumps(sections)
        result, _ = self.summarize(packet)
        self.assertEqual(result["status"], "insufficient_evidence")
        self.assertEqual(result["missing_fields"], ["结果"])
        self.assertIn("### 结果\n\n资料不足", result["answer"])

    def test_summary_omits_author_only_chunks_before_abstract(self):
        doc_id = self.upload("含作者.txt", ("作者目录标记" * 100 + "\nAbstract\n研究机器翻译，采用Transformer。\n"
                                           "Conclusion\n支持翻译任务。\n").encode())
        result, http = self.summarize(doc_id=doc_id)
        source = json.loads(http.call_args.args[0].data)["messages"][1]["content"]
        self.assertIn("Abstract", source)
        self.assertNotIn("作者目录标记" * 30, source)
        self.assertTrue(result["input_truncated"])

    def test_summary_rejects_unknown_or_changed_uploaded_ids_before_model(self):
        with patch("src.generation.rag_pipeline.urlopen") as http:
            for doc_id in ("非法路径", "0" * 64):
                with self.subTest(doc_id=doc_id):
                    self.assertEqual(execute_tool("paper_summary", {"doc_id": doc_id}, AVAILABLE_TOOLS)["status"], "error")
            (Path(self.directory.name) / self.doc_id / "论文.md").write_text("已改变原文")
            self.assertEqual(execute_tool("paper_summary", {"doc_id": self.doc_id}, AVAILABLE_TOOLS)["status"], "error")
        http.assert_not_called()

    def test_summary_rejects_empty_loaded_content_before_model(self):
        from langchain_core.documents import Document
        with patch("src.data_loader.load_document", return_value=[Document(page_content=" ")]), \
                patch("src.generation.rag_pipeline.urlopen") as http:
            with self.assertRaisesRegex(ValueError, "没有可用文本"):
                paper_summary.invoke({"doc_id": self.doc_id})
        http.assert_not_called()

    def test_summary_rejects_malformed_sections_and_fabricated_references(self):
        bad = [None, {"background": self.sections["background"]},
               {**self.sections, "extra": {}}, {**self.sections, "results": None}]
        for section in ({"text": "结果", "reference_ids": [99]}, {"text": "结果", "reference_ids": [True]},
                        {"text": "结果", "reference_ids": [1, 1]}, {"text": "结果", "reference_ids": []},
                        {"text": "", "reference_ids": [1]}, {"text": "[参考文档1]", "reference_ids": [1]},
                        {"text": "长" * 101, "reference_ids": [1]}, {"text": 123, "reference_ids": [1]}):
            bad.append({**self.sections, "results": section})
        for sections in bad:
            with self.subTest(sections=sections):
                with self.assertRaises(ValueError):
                    self.summarize({**self.packet, "message": {"content": json.dumps(sections)}})

    def test_summary_rejects_incomplete_or_invalid_model_packets(self):
        for update in ({"done": False}, {"done_reason": "length"}, {"error": "失败"},
                       {"model": ""}, {"message": None}, {"message": {"content": "不是JSON"}}):
            with self.subTest(update=update), self.assertRaises(ValueError):
                self.summarize({**self.packet, **update})

    def test_summary_unknown_tokens_are_not_zero(self):
        packet = {key: value for key, value in self.packet.items() if key not in ("prompt_eval_count", "eval_count")}
        result, _ = self.summarize(packet)
        self.assertEqual(result["usage"], {"prompt_eval_count": None, "eval_count": None})

    def test_summary_api_failure_is_tool_error_without_retry(self):
        with patch("src.generation.rag_pipeline.urlopen", side_effect=URLError("断开")) as http:
            event = execute_tool("paper_summary", {"doc_id": self.doc_id}, AVAILABLE_TOOLS)
        self.assertEqual(event["status"], "error")
        self.assertIsNone(event["result"])
        self.assertEqual(http.call_count, 1)

    def test_current_time_is_real_timezone_aware_and_uses_no_model(self):
        before = datetime.now().astimezone() - timedelta(seconds=1)
        with patch("src.generation.rag_pipeline.urlopen") as http:
            event = execute_tool("current_time", {}, AVAILABLE_TOOLS)
        after = datetime.now().astimezone()
        now = datetime.fromisoformat(event["result"]["system_time"])
        self.assertEqual(event["status"], "success")
        self.assertIsNotNone(now.utcoffset())
        self.assertLessEqual(before, now)
        self.assertLessEqual(now, after)
        self.assertEqual(event["result"]["usage"]["eval_count"], 0)
        http.assert_not_called()

    def test_current_time_uses_system_timezone_without_hardcoding(self):
        instant = datetime(2026, 10, 2, 9, 30, 40, tzinfo=timezone(timedelta(hours=-5)))
        with patch("src.agent.tools.datetime") as clock:
            clock.now.return_value.astimezone.return_value = instant
            result = current_time.invoke({})
        self.assertEqual(result["system_time"], "2026-10-02T09:30:40-05:00")

    def test_disabled_search_is_not_registered_and_cannot_call_http(self):
        with patch("httpx.Client") as client:
            self.assertNotIn(web_search, get_available_tools())
            event = execute_tool("web_search", {"query": "Transformer"}, [web_search])
        self.assertEqual(event["status"], "error")
        self.assertIn("未启用", event["error"])
        client.assert_not_called()

    def test_search_registration_reads_current_boolean_switch(self):
        self.config["agent"]["online_search_enabled"] = True
        self.assertEqual(get_available_tools(), [*AVAILABLE_TOOLS, web_search])
        self.assertNotIn(web_search, AVAILABLE_TOOLS)
        self.config["agent"]["online_search_enabled"] = "false"
        self.assertEqual(get_available_tools(), AVAILABLE_TOOLS)

    def test_default_agent_registers_enabled_search_and_explicit_empty_stays_empty(self):
        self.config["agent"]["online_search_enabled"] = True
        with patch("src.agent.react_loop.route_question", return_value=None), \
                patch("src.agent.react_loop.think", side_effect=ValueError("测试终止")) as think_mock:
            list(run_react("搜索论文"))
            self.assertIn(web_search, think_mock.call_args.args[1])
            list(run_react("搜索论文", tools=[]))
            self.assertEqual(think_mock.call_args.args[1], [])

    def test_search_parses_top_five_titles_snippets_and_unwraps_redirects(self):
        html = "".join(f'<div class="result"><a class="result__a" href="/l/?uddg=https%3A%2F%2Fexample.com%2F{i}">'
                       f'<b>论文{i}</b></a><div class="result__snippet">摘要 <b>{i}</b></div></div>' for i in range(7))
        result, client = self.search(html)
        self.assertEqual(result["status"], "results")
        self.assertEqual(len(result["results"]), 5)
        self.assertEqual(result["results"][0], {"title": "论文0", "snippet": "摘要 0", "url": "https://example.com/0"})
        call = client.return_value.__enter__.return_value.get.call_args
        self.assertEqual(call.kwargs["params"], {"q": "Transformer 论文", "kl": "cn-zh"})
        self.assertEqual(client.return_value.__enter__.return_value.get.call_count, 1)
        self.assertEqual(result["usage"]["eval_count"], 0)

    def test_search_skips_invalid_links_and_allows_missing_snippet(self):
        result, _ = self.search('<div class="result"><a class="result__a" href="javascript:alert(1)">坏链接</a></div>'
                                '<div class="result"><a class="result__a" href="https://arxiv.org/">论文</a></div>')
        self.assertEqual(result["results"], [{"title": "论文", "snippet": "", "url": "https://arxiv.org/"}])

    def test_search_distinguishes_legitimate_empty_results(self):
        result, _ = self.search('<div class="no-results"><div class="no-results__message">No results found</div></div>')
        self.assertEqual(result["status"], "no_results")
        self.assertEqual(result["results"], [])
        self.assertIn("未找到", result["message"])

    def test_search_challenge_or_unrecognized_page_is_error(self):
        for html, status in (("<form id='challenge-form'>验证</form>", 200), ("待验证", 202), ("<html></html>", 200)):
            with self.subTest(status=status, html=html), self.assertRaisesRegex(RuntimeError, "联网搜索失败"):
                self.search(html, status)

    def test_search_http_error_is_not_empty_result(self):
        with self.assertRaisesRegex(RuntimeError, "403"):
            self.search("拒绝访问", 403)

    def test_search_timeout_is_execution_error_without_retry(self):
        import httpx
        self.config["agent"]["online_search_enabled"] = True
        with patch("httpx.Client") as client:
            client.return_value.__enter__.return_value.get.side_effect = httpx.ReadTimeout("超时")
            event = execute_tool("web_search", {"query": "论文"}, get_available_tools())
        self.assertEqual(event["status"], "error")
        self.assertIn("检查网络", event["error"])
        self.assertEqual(client.return_value.__enter__.return_value.get.call_count, 1)

    def test_search_empty_query_rejected_before_network(self):
        self.config["agent"]["online_search_enabled"] = True
        with patch("httpx.Client") as client, self.assertRaisesRegex(ValueError, "不能为空"):
            web_search.invoke({"query": " "})
        client.assert_not_called()


if __name__ == "__main__":
    unittest.main()
