"""5.2.1 Prompt工程与生成策略：TestCitations。"""

import sys
from pathlib import Path

# 从其他目录直接运行测试文件时，也能定位 src 和共用样例。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from copy import deepcopy
import tempfile
import unittest
from unittest.mock import patch
from langchain_core.documents import Document
from src.generation.rag_pipeline import build_context, resolve_citations


class TestCitations(unittest.TestCase):
    """校验引用编号和原文映射，不把映射通过等同于语义支持。"""

    def setUp(self):
        self.config = {"generation": {"max_context_chars": 6000, "max_prompt_chars": 12000}}
        config_patch = patch("src.generation.rag_pipeline.load_config", return_value=self.config)
        config_patch.start()
        self.addCleanup(config_patch.stop)
        self.documents = [Document(page_content=f"证据{i}：编码器的结构。", metadata={
            "source_file": "attention.pdf", "page_number": i + 2,
            "doc_id": "paper-attention", "chunk_id": f"chunk-{i}"}) for i in (1, 2, 3)]
        self.results = [(document, 1 - i / 10) for i, document in enumerate(self.documents)]
        self.context = build_context("编码器是什么结构？", self.results)

    def test_context_references_follow_sorting_and_budget(self):
        self.config["generation"]["max_context_chars"] = build_context("问题", self.results[:1])["context_chars"]
        context = build_context("问题", list(reversed(self.results)))
        self.assertEqual(len(context["references"]), context["chunk_count"])
        self.assertEqual([ref["id"] for ref in context["references"]], [1])
        self.assertEqual(context["references"][0]["metadata"]["chunk_id"], "chunk-1")
        self.assertEqual(context["references"][0]["text"], self.documents[0].page_content)
        self.assertEqual(context["dropped_count"], 2)

    def test_reference_metadata_is_an_independent_snapshot(self):
        document = Document(page_content="原文", metadata={"source_file": "原.pdf", "page_number": 3,
                                                        "detail": {"tag": "原始"}})
        context = build_context("问题", [(document, 0.8)])
        document.metadata["source_file"] = "已改.pdf"
        document.metadata["detail"]["tag"] = "已修改"
        ref = context["references"][0]
        self.assertEqual(ref["source_file"], "原.pdf")
        self.assertEqual(ref["metadata"]["detail"]["tag"], "原始")

    def test_inline_pdf_location_and_authoritative_footer(self):
        answer = "## 回答\n结构如原文所述[参考文档1]。\n\n## 参考来源\n伪造.pdf，第999页"
        resolved = resolve_citations(answer, self.context)
        self.assertIn("结构如原文所述[参考文档1：attention.pdf；第3页（物理页码）]。", resolved["answer"])
        self.assertNotIn("伪造.pdf", resolved["answer"])
        self.assertNotIn("999", resolved["answer"])
        self.assertEqual(resolved["answer"].count("## 参考来源"), 1)
        self.assertEqual(resolved["citations"][0]["text"], self.documents[0].page_content)
        self.assertEqual(resolved["citations"][0]["metadata"]["chunk_id"], "chunk-1")
        self.assertEqual(resolved["warnings"], [])

    def test_used_order_repeated_citations_and_unused_blocks(self):
        resolved = resolve_citations("先看[参考文档2]，再看[参考文档1]，再次看[参考文档2]。", self.context)
        self.assertEqual([ref["id"] for ref in resolved["citations"]], [2, 1])
        footer = resolved["answer"].split("## 参考来源", 1)[1]
        self.assertEqual(footer.count("[参考文档2]"), 1)
        self.assertNotIn("参考文档3", resolved["answer"])
        self.assertIn("第4页", footer)
        self.assertIn("第3页", footer)
        # 同一文件不同页/块不能按文件名合并成一个引用。
        self.assertEqual(len(resolved["citations"]), 2)

    def test_cross_page_table_location_is_kept(self):
        document = Document(page_content="跨页表格证据", metadata={"source_file": "table.pdf",
                            "page_number": 3, "page_end": 5, "content_type": "table"})
        resolved = resolve_citations("实验结果[参考文档1]。", build_context("问题", [(document, 0.8)]))
        self.assertIn("table.pdf；第3–5页（物理页码）", resolved["answer"])
        self.assertEqual(resolved["citations"][0]["metadata"]["page_end"], 5)

    def test_word_and_text_do_not_get_invented_pdf_pages(self):
        for metadata, position in (({"source_file": "word.docx", "paragraph_index": 2}, "段落2"),
                                   ({"source_file": "word.docx", "table_index": 1}, "表格1"),
                                   ({"source_file": "text.md", "line_start": 2, "line_end": 6}, "行2–6")):
            with self.subTest(metadata=metadata):
                context = build_context("问题", [(Document(page_content="证据", metadata=metadata), 0.8)])
                resolved = resolve_citations("结论[参考文档1]。", context)
                self.assertIn(position, resolved["answer"])
                self.assertNotIn("物理页码", resolved["answer"])
                self.assertEqual(resolved["warnings"], [])

    def test_truncated_evidence_does_not_include_unseen_tail(self):
        document = Document(page_content="已提供的事实。" * 100 + "未读到的尾部结论",
                            metadata={"source_file": "长.pdf", "page_number": 3, "chunk_id": "long"})
        self.config["generation"]["max_context_chars"] = 100
        context = build_context("问题", [(document, 0.9), (self.documents[0], 0.8)])
        resolved = resolve_citations("已提供事实[参考文档1]；其他结论[参考文档2]。", context)
        evidence = resolved["citations"][0]
        self.assertTrue(evidence["truncated"])
        self.assertTrue(document.page_content.startswith(evidence["text"]))
        self.assertIn(evidence["text"], context["context"])
        self.assertNotIn("未读到的尾部结论", evidence["text"])
        self.assertNotIn("[正文已截断]", evidence["text"])
        self.assertIn("仅展示已送入上下文的证据", resolved["answer"])
        self.assertEqual(resolved["invalid_citation_ids"], [2])

    def test_unknown_ids_are_marked_and_never_assigned_sources(self):
        resolved = resolve_citations("[参考文档0] [参考文档99] [参考文档99] [参考文档1]", self.context)
        self.assertEqual(resolved["invalid_citation_ids"], [0, 99])
        self.assertIn("[无效引用：参考文档99]", resolved["answer"])
        self.assertEqual([ref["id"] for ref in resolved["citations"]], [1])
        self.assertEqual(len(resolved["warnings"]), 2)

    def test_uncited_answer_does_not_add_retrieved_sources(self):
        resolved = resolve_citations("## 回答\n没有引用的结论。\n## 参考来源\n[参考文档1]", self.context)
        self.assertTrue(resolved["missing_citations"])
        self.assertEqual(resolved["citations"], [])
        self.assertNotIn("attention.pdf", resolved["answer"])
        self.assertIn("无可引用来源", resolved["answer"])
        self.assertEqual(len(resolved["warnings"]), 1)

    def test_empty_context_cannot_resolve_a_document(self):
        context = build_context("问题", [])
        resolved = resolve_citations("通用回答[参考文档1]。", context)
        self.assertEqual(resolved["citations"], [])
        self.assertEqual(resolved["invalid_citation_ids"], [1])
        plain = resolve_citations("当前知识库中未找到相关文档。", context)
        self.assertEqual(plain["warnings"], [])
        self.assertFalse(plain["missing_citations"])

    def test_missing_source_information_is_flagged(self):
        for metadata in ({}, {"source_file": "未知位置.pdf"}, {"page_number": 2}):
            with self.subTest(metadata=metadata):
                context = build_context("问题", [(Document(page_content="证据", metadata=metadata), 1.0)])
                resolved = resolve_citations("结论[参考文档1]。", context)
                self.assertEqual(len(resolved["warnings"]), 1)
                self.assertIn("无法完整溯源", resolved["warnings"][0])
                self.assertNotIn("第1页", resolved["answer"])

    def test_code_examples_and_escaped_markers_are_not_evidence(self):
        answer = ("代码 `[参考文档99]`、``[参考文档98]`` 和转义 \\[参考文档97]。\n"
                  "```text\n## 参考来源\n[参考文档96]\n```\n"
                  "~~~text\n[参考文档95]\n~~~\n正文结论[参考文档1]。")
        resolved = resolve_citations(answer, self.context)
        self.assertEqual([ref["id"] for ref in resolved["citations"]], [1])
        self.assertEqual(resolved["invalid_citation_ids"], [])
        self.assertIn("```text\n## 参考来源\n[参考文档96]\n```", resolved["answer"])
        self.assertIn("[参考文档1：attention.pdf", resolved["answer"])
        self.assertEqual(resolved["warnings"], [])

    def test_model_supplied_link_is_not_a_source_link(self):
        resolved = resolve_citations("结论[参考文档1](https://example.invalid/fake.pdf)。", self.context)
        self.assertNotIn("example.invalid", resolved["answer"])
        self.assertIn("[参考文档1：attention.pdf", resolved["answer"])

    def test_markdown_characters_in_filename_are_escaped(self):
        document = Document(page_content="证据", metadata={"source_file": "论文_[A]<草稿>.pdf", "page_number": 2})
        resolved = resolve_citations("事实[参考文档1]。", build_context("问题", [(document, 1.0)]))
        self.assertIn("论文\\_\\[A\\]\\<草稿\\>.pdf", resolved["answer"])
        self.assertEqual(resolved["citations"][0]["source_file"], "论文_[A]<草稿>.pdf")

    def test_context_is_not_modified_by_resolution_or_returned_evidence(self):
        before = deepcopy(self.context)
        resolved = resolve_citations("事实[参考文档1]。", self.context)
        resolved["citations"][0]["metadata"]["page_number"] = 999
        resolved["citations"][0]["text"] = "调用方改变结果"
        self.assertEqual(self.context, before)

    def test_blank_answer_is_rejected(self):
        for answer in ("", " \n\t"):
            with self.subTest(answer=answer), self.assertRaisesRegex(ValueError, "答案不能为空"):
                resolve_citations(answer, self.context)

    def test_real_multi_format_loaders_and_chunks_keep_traceable_positions(self):
        # 临时生成实际文件，再走真实加载/分块；这些是功能样例，不是模型生成答案。
        import pymupdf
        from docx import Document as WordDocument
        from src.data_loader import load_document
        from src.chunking import split_documents

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = pymupdf.open()
            for text in ("Evidence from page one.", "Evidence from page two."):
                pdf.new_page().insert_text((72, 72), text)
            pdf.save(root / "paper.pdf")
            pdf.close()
            word = WordDocument()
            word.add_paragraph("正文证据。")
            word.add_table(rows=1, cols=1).cell(0, 0).text = "表格证据"
            word.save(root / "paper.docx")
            for filename in ("notes.txt", "notes.md"):
                (root / filename).write_text("第一行证据\n第二行证据", encoding="utf-8")
            chunks = []
            for filename in ("paper.pdf", "paper.docx", "notes.txt", "notes.md"):
                chunks.extend(split_documents(load_document(root / filename)))
            self.assertEqual(len(chunks), 6)
            context = build_context("功能核验", [(doc, 1 - i / 10) for i, doc in enumerate(chunks)])
            answer = "；".join(f"功能样例[参考文档{i}]" for i in range(1, 7))
            resolved = resolve_citations(answer, context)
            self.assertEqual(len(resolved["citations"]), 6)
            self.assertEqual(resolved["warnings"], [])
            for citation, chunk in zip(resolved["citations"], chunks):
                self.assertEqual(citation["metadata"]["chunk_id"], chunk.metadata["chunk_id"])
                self.assertEqual(citation["metadata"]["doc_id"], chunk.metadata["doc_id"])
                self.assertEqual(citation["text"], chunk.page_content)
                self.assertTrue(Path(citation["metadata"]["source"]).is_file())
            for location in ("第1页", "第2页", "段落1", "表格1", "行1–2"):
                self.assertIn(location, resolved["answer"])


if __name__ == "__main__":
    unittest.main()
