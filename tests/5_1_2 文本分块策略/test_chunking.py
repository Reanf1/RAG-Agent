"""5.1.2 文本分块策略：TestChunking。"""

import sys
from pathlib import Path

# 从其他目录直接运行测试文件时，也能定位 src 和共用样例。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import tempfile
import json
import unittest
from unittest.mock import patch
import pymupdf
from docx import Document as WordDocument
from langchain_core.documents import Document
from src.data_loader.pdf_loader import load_pdf
from src.data_loader.docx_loader import load_docx
from src.data_loader.text_loader import load_text
from src.chunking import split_documents
from src.chunking.fixed_chunk import split_fixed
from src.chunking.recursive_chunk import split_recursive
from src.chunking.semantic_chunk import split_semantic


class TestChunking(unittest.TestCase):
    """用真实切分器检查内容、边界与溯源，不依赖模型和检索库。"""

    strategies = (split_fixed, split_recursive, split_semantic)

    def test_short_document(self):
        """短文档保持原文，包括缩进与尾部换行。"""
        text = "  中文论文摘要。\r\n"
        for splitter in self.strategies:
            with self.subTest(strategy=splitter.__name__):
                chunks = splitter([Document(page_content=text)], 32, 4)
                self.assertEqual([chunk.page_content for chunk in chunks], [text])
                self.assertEqual(chunks[0].metadata["start_index"], 0)
                self.assertEqual(chunks[0].metadata["end_index"], len(text))

    def test_empty_documents(self):
        """空列表、空文本和仅空白文档不生成无效块。"""
        for splitter in self.strategies:
            self.assertEqual(splitter([], 16, 0), [])
            self.assertEqual(splitter([Document(page_content=text) for text in ("", " \n\t")], 16, 0), [])

    def test_invalid_parameters(self):
        """拒绝不合法大小和重叠，避免零步长或负数切片。"""
        for splitter in self.strategies:
            for size, overlap in ((0, 0), (-1, 0), (4, -1), (4, 4), (4, 5),
                                  (4.5, 0), (4, 1.5), (True, 0), (4, False)):
                with self.subTest(strategy=splitter.__name__, size=size, overlap=overlap):
                    with self.assertRaises(ValueError):
                        splitter([], size, overlap)

    def test_fixed_exact_overlap(self):
        """固定窗口按字符数切分，相邻正文块精确重叠。"""
        text = "甲乙丙丁戊己庚辛壬癸"
        chunks = split_fixed([Document(page_content=text)], 4, 1)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲乙丙丁", "丁戊己庚", "庚辛壬癸"])
        self.assertEqual([chunk.metadata["start_index"] for chunk in chunks], [0, 3, 6])

    def test_fixed_no_redundant_tail(self):
        """末尾已被覆盖时不再输出完全包含在上一块中的尾块。"""
        for text, expected in (("12345678", ["12345678"]), ("1234567890", ["12345678", "67890"])):
            chunks = split_fixed([Document(page_content=text)], 8, 3)
            self.assertEqual([chunk.page_content for chunk in chunks], expected)

    def test_fixed_course_sizes(self):
        """课程规定的 256/512/1024 都可切分，正文长度不超过设置。"""
        text = "科研文本" * 600
        for size in (256, 512, 1024):
            chunks = split_fixed([Document(page_content=text)], size, 64)
            self.assertGreater(len(chunks), 1)
            self.assertTrue(all(len(chunk.page_content) <= size for chunk in chunks))
            self.assertEqual(chunks[-1].metadata["end_index"], len(text))

    def test_recursive_paragraph_boundaries(self):
        """优先按段落分隔符切分，并保留分隔符原文。"""
        text = "甲段内容。\n\n乙段内容。\n\n丙段内容。"
        chunks = split_recursive([Document(page_content=text)], 9, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["甲段内容。", "\n\n乙段内容。", "\n\n丙段内容。"])
        self.assertEqual("".join(chunk.page_content for chunk in chunks), text)

    def test_recursive_repeated_punctuation_advances_and_covers_source(self):
        """默认重叠下，短重复分隔符不能回指旧片段或漏掉句号。"""
        text = "\n\n。\n。。" + "xy" * 400
        chunks = split_recursive([Document(page_content=text)], 512, 64)
        covered = set()
        previous_end = 0
        for chunk in chunks:
            start, end = chunk.metadata["start_index"], chunk.metadata["end_index"]
            self.assertGreater(end, previous_end)
            self.assertEqual(chunk.page_content, text[start:end])
            covered.update(range(start, end))
            previous_end = end
        self.assertTrue(all(i in covered for i, char in enumerate(text) if not char.isspace()))

    def test_recursive_character_fallback(self):
        """没有分隔符的长文本仍能限长并保留重叠。"""
        chunks = split_recursive([Document(page_content="ABCDEFGHIJ")], 4, 1)
        self.assertEqual([chunk.page_content for chunk in chunks], ["ABCD", "DEFG", "GHIJ"])

    def test_semantic_keeps_paragraphs(self):
        """短段落整体保留，空行跟随原段落。"""
        text = "甲段内容。\r\n\r\n乙段内容。\r\n\r\n丙段内容。"
        chunks = split_semantic([Document(page_content=text)], 10, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["甲段内容。\r\n\r\n", "乙段内容。\r\n\r\n", "丙段内容。"])

    def test_semantic_chinese_sentences_and_quotes(self):
        """长段落拆句，结束引号与标点跟随原句。"""
        text = "第一句。”第二句！第三句？第四句。"
        chunks = split_semantic([Document(page_content=text)], 8, 0)
        self.assertEqual([chunk.page_content for chunk in chunks], ["第一句。”", "第二句！第三句？", "第四句。"])

    def test_semantic_english_decimal_and_doi(self):
        """英文句号可拆句，小数和 DOI 内部的点不误拆。"""
        text = "Value 3.14. DOI 10.1000/xyz. Result good."
        chunks = split_semantic([Document(page_content=text)], 27, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["Value 3.14.", " DOI 10.1000/xyz.", " Result good."])

    def test_semantic_complete_unit_overlap(self):
        """重叠复用完整句子；不足容纳整句时不截出半句。"""
        document = Document(page_content="甲甲。乙乙。丙丙。丁丁。")
        chunks = split_semantic([document], 6, 3)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲甲。乙乙。", "乙乙。丙丙。", "丙丙。丁丁。"])
        chunks = split_semantic([document], 6, 2)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲甲。乙乙。", "丙丙。丁丁。"])

    def test_semantic_long_sentence_fallback(self):
        """超长单句按字符兜底，最后短片段不丢失。"""
        text = "没有任何句号的超长文本用于测试"
        chunks = split_semantic([Document(page_content=text)], 5, 2)
        self.assertEqual([chunk.page_content for chunk in chunks], [text[i:i + 5] for i in range(0, len(text), 5)])

    def test_ranges_cover_original_content(self):
        """不同大小和重叠下，区间准确、向前推进，所有非空白字符均被覆盖。"""
        texts = ("重复段落。\n\n" * 30, "ABCD" * 50, "甲。乙！丙？\r\n\r\n" * 20)
        for splitter in self.strategies:
            for text in texts:
                for size, overlap in ((1, 0), (8, 0), (8, 2), (16, 15), (64, 10)):
                    with self.subTest(strategy=splitter.__name__, size=size, overlap=overlap, text=text[:12]):
                        chunks = splitter([Document(page_content=text)], size, overlap)
                        covered = set()
                        previous_start = -1
                        for chunk in chunks:
                            start, end = chunk.metadata["start_index"], chunk.metadata["end_index"]
                            self.assertGreater(start, previous_start)
                            self.assertLessEqual(len(chunk.page_content), size)
                            self.assertEqual(chunk.page_content, text[start:end])
                            covered.update(range(start, end))
                            previous_start = start
                        self.assertTrue(all(i in covered for i, char in enumerate(text) if not char.isspace()))

    def test_preserves_metadata_without_mutating_input(self):
        """保留 PDF 页码、版面和公式原位置，并复制元数据。"""
        metadata = {"doc_id": "论文ID", "source_file": "论文.pdf", "page": 2, "page_number": 3,
                    "layout": "two_column", "formula_layout": '[{"text":"x²"}]'}
        document = Document(page_content="公式 x^{2}。\n实验结果。" * 5, metadata=metadata)
        for splitter in self.strategies:
            chunks = splitter([document], 20, 2)
            for chunk in chunks:
                for key, value in metadata.items():
                    if key != "formula_layout":
                        self.assertEqual(chunk.metadata[key], value)
            # 此旧布局的x²与原文x^{2}不一致，无法定位时仍完整保留在首块。
            self.assertEqual(json.loads(chunks[0].metadata["formula_layout"]), json.loads(metadata["formula_layout"]))
            self.assertTrue(all(chunk.metadata["formula_layout_origin_chunk_id"] == chunks[0].metadata["chunk_id"] for chunk in chunks[1:]))
            chunks[0].metadata["source_file"] = "改名.pdf"
            self.assertEqual(document.metadata, metadata)
            self.assertNotIn("chunk_id", document.metadata)

    def test_image_layout_is_stored_once_per_page_with_stable_source(self):
        """三种分块不重复整页图片坐标；每页来源独立，重做ID稳定，原文不修改。"""
        images = json.dumps([{"page_number": 1, "bbox": [0, i, 10, i + 1]} for i in range(5)])
        documents = [Document(page_content="  \n" + "图像附近的说明。" * 20,
                              metadata={"doc_id": "paper", "page": page, "image_regions": images})
                     for page in (0, 1)]
        for splitter in self.strategies:
            chunks = splitter(documents, 30, 4)
            for page in (0, 1):
                group = [chunk for chunk in chunks if chunk.metadata["page"] == page]
                self.assertGreater(len(group), 1)
                self.assertEqual(group[0].metadata["image_regions"], images)
                self.assertTrue(all("image_regions" not in chunk.metadata for chunk in group[1:]))
                self.assertTrue(all(chunk.metadata["image_regions_origin_chunk_id"] == group[0].metadata["chunk_id"] for chunk in group[1:]))
                stored = sum(len(chunk.metadata.get("image_regions", "")) + len(chunk.metadata.get("image_regions_origin_chunk_id", "")) for chunk in group)
                self.assertLess(stored, len(images) * len(group) / 2)
            again = splitter(documents, 30, 4)
            self.assertEqual([chunk.metadata for chunk in chunks], [chunk.metadata for chunk in again])
            self.assertTrue(all(document.metadata["image_regions"] == images for document in documents))
            self.assertTrue(all("image_regions_origin_chunk_id" not in document.metadata for document in documents))

    def test_stable_distinct_chunk_ids(self):
        """重复执行 ID 不变，重复正文在不同页/段落及不同配置中不会混淆。"""
        documents = [Document(page_content="重复内容" * 8, metadata={"doc_id": "id", "page": page})
                     for page in (0, 1)]
        for splitter in self.strategies:
            chunks = splitter(documents, 12, 2)
            ids = [chunk.metadata["chunk_id"] for chunk in chunks]
            self.assertEqual(ids, [chunk.metadata["chunk_id"] for chunk in splitter(documents, 12, 2)])
            self.assertEqual(len(set(ids)), len(ids))
            changed = {chunk.metadata["chunk_id"] for chunk in splitter(documents, 16, 2)}
            self.assertTrue(set(ids).isdisjoint(changed))
        word = [Document(page_content="重复段落", metadata={"doc_id": "word", "block_index": index})
                for index in (1, 2)]
        self.assertEqual(len({chunk.metadata["chunk_id"] for chunk in split_fixed(word, 8)}), 2)

    def test_layout_is_localized_without_losing_cross_chunk_or_unlocated_lines(self):
        """布局按字符相交保留，跨块公式、重复文本和不在正文的表格坐标均不丢。"""
        lines = [{"text": f"公式{i}: x^{{2}} = {i}，附加说明。", "bbox": [1, i, 10, i + 1]} for i in range(20)]
        unlocated = {"text": "独立表格中的公式", "bbox": [1, 40, 10, 41]}
        layout = json.dumps(lines + [unlocated], ensure_ascii=False)
        text = "\n".join(line["text"] for line in lines)
        document = Document(page_content=text, metadata={"formula_layout": layout, "page_number": 1})
        for splitter in self.strategies:
            chunks = splitter([document], 40, 8)
            actual = [json.loads(chunk.metadata["formula_layout"]) for chunk in chunks]
            self.assertLess(sum(len(chunk.metadata["formula_layout"]) for chunk in chunks), len(layout) * len(chunks) / 2)
            self.assertIn(unlocated, actual[0])
            self.assertFalse(any(unlocated in entries for entries in actual[1:]))
            self.assertTrue(all(chunk.metadata["formula_layout_origin_chunk_id"] == chunks[0].metadata["chunk_id"] for chunk in chunks[1:]))
            for chunk, entries in zip(chunks, actual):
                start, end = chunk.metadata["start_index"], chunk.metadata["end_index"]
                for line in lines:
                    offset = text.index(line["text"])
                    self.assertEqual(line in entries, offset < end and offset + len(line["text"]) > start)
            self.assertEqual(document.metadata["formula_layout"], layout)

    def test_pdf_and_word_tables_remain_whole(self):
        """独立表格不拆，即使超长也保留跨页行信息和保护标记。"""
        rows = '[{"cells":["数据"],"page_number":4},{"cells":["结果"],"page_number":5}]'
        tables = [Document(page_content="| 表头 |\n| --- |\n| 数据 |\n| 结果 |", metadata={
            "doc_id": "id", "content_type": "table", "table_id": "id:table:1",
            "page": 3, "page_number": 4, "page_end": 5, "table_rows": rows,
        }), Document(page_content="表头\t值\n数据\t结果", metadata={"block_type": "table", "block_index": 2})]
        for splitter in self.strategies:
            chunks = splitter(tables, 4, 1)
            self.assertEqual([chunk.page_content for chunk in chunks], [table.page_content for table in tables])
            self.assertTrue(all(chunk.metadata["chunk_preserved"] == "table" for chunk in chunks))
            self.assertEqual(chunks[0].metadata["table_rows"], rows)
            self.assertEqual(chunks[0].metadata["page_end"], 5)

    def test_loaded_text_line_numbers(self):
        """真实 TXT 加载后细化 CRLF 行范围；块内原文和文件来源保持一致。"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "文本.txt"
            path.write_bytes("甲乙\r\n丙丁\r\n戊己".encode("utf-8"))
            documents = load_text(path)
            chunks = split_fixed(documents, 4, 0)
            self.assertEqual([(chunk.metadata["line_start"], chunk.metadata["line_end"]) for chunk in chunks],
                             [(1, 1), (2, 2), (3, 3)])
            self.assertTrue(all(chunk.metadata["source"] == str(path.resolve()) for chunk in chunks))
            self.assertEqual(documents[0].metadata["line_end"], 3)
            # 恰好从 CRLF 中间开始的块仍归原行，不提前增加行号。
            chunks = split_fixed(documents, 3, 0)
            self.assertEqual([(chunk.metadata["line_start"], chunk.metadata["line_end"]) for chunk in chunks],
                             [(1, 1), (1, 2), (2, 3), (3, 3)])

    def test_loaded_word_and_pdf(self):
        """真实 Word/PDF 加载结果可直接分块，来源位置与表格不丢失。"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "论文.docx"
            word = WordDocument()
            word.add_paragraph("研究背景。实验方法。研究结论。" * 3)
            word.add_table(rows=1, cols=2).cell(0, 0).text = "完整表格"
            word.save(path)
            for splitter in self.strategies:
                chunks = splitter(load_docx(path), 10, 2)
                self.assertTrue(all("block_index" in chunk.metadata and "page" not in chunk.metadata for chunk in chunks))
                self.assertEqual(chunks[-1].page_content, "完整表格\t")
                self.assertEqual(chunks[-1].metadata["chunk_preserved"], "table")
            path = Path(directory) / "论文.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "研究背景。实验方法。研究结论。", fontname="china-s")
                pdf.save(path)
            for splitter in self.strategies:
                chunks = splitter(load_pdf(path), 10, 2)
                self.assertTrue(all(chunk.metadata["page_number"] == 1 for chunk in chunks))

    def test_config_defaults_and_overrides(self):
        """统一入口读取唯一配置，显式参数可覆盖实验值。"""
        documents = [Document(page_content="1234567890")]
        with patch("src.chunking.load_config", return_value={
            "chunking": {"strategy": "fixed", "chunk_size": 4, "chunk_overlap": 1},
        }):
            self.assertEqual([chunk.page_content for chunk in split_documents(documents)], ["1234", "4567", "7890"])
            chunks = split_documents(documents, strategy="semantic", chunk_size=8, chunk_overlap=0)
            self.assertEqual([chunk.page_content for chunk in chunks], ["12345678", "90"])
            self.assertTrue(all(chunk.metadata["chunk_strategy"] == "semantic" for chunk in chunks))

    def test_unknown_strategy(self):
        """未知策略明确报错，不悄悄回退。"""
        with self.assertRaisesRegex(ValueError, "不支持的分块策略"):
            split_documents([], "unknown", 8, 0)


if __name__ == "__main__":
    unittest.main()
