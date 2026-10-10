"""5.1.1 文档加载与批量导入：TestDOCXLoader。"""

import sys
from pathlib import Path

# 从其他目录直接运行测试文件时，也能定位 src 和共用样例。
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import os
import tempfile
import unittest
from docx import Document as WordDocument
from langchain_core.documents import Document
from src.data_loader.docx_loader import load_docx


class TestDOCXLoader(unittest.TestCase):
    """通过实际 DOCX 验证正文、表格顺序和位置，不修改原始课程资料。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "论文.DOCX"

    def test_paragraph_table_order_and_metadata(self):
        """表格保留在正文原位置，中文段落与单元格内容均可读取。"""
        word = WordDocument()
        word.add_paragraph("摘要")
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "方法"
        table.cell(0, 1).text = "准确率"
        table.cell(1, 0).text = "Transformer"
        table.cell(1, 1).text = "90%"
        word.add_paragraph("结论")
        word.save(self.path)

        documents = load_docx(os.path.relpath(self.path))
        self.assertEqual([d.page_content for d in documents], [
            "摘要", "方法\t准确率\nTransformer\t90%", "结论",
        ])
        self.assertEqual([d.metadata["block_index"] for d in documents], [1, 2, 3])
        self.assertEqual([d.metadata["block_type"] for d in documents], [
            "paragraph", "table", "paragraph",
        ])
        self.assertEqual(documents[0].metadata["paragraph_index"], 1)
        self.assertEqual(documents[1].metadata["table_index"], 1)
        self.assertEqual(documents[2].metadata["paragraph_index"], 2)
        for document in documents:
            self.assertIsInstance(document, Document)
            self.assertEqual(document.metadata["source"], str(self.path.resolve()))
            self.assertEqual(document.metadata["source_file"], "论文.DOCX")
            self.assertEqual(document.metadata["file_type"], ".docx")
            self.assertNotIn("page", document.metadata)
            self.assertNotIn("page_number", document.metadata)


if __name__ == "__main__":
    unittest.main()
