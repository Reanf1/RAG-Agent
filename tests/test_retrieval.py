"""模块一测试：临时文档与真实临时 Chroma；大模型只在独立实验中运行。"""

import hashlib
from copy import deepcopy
import io
import json
import math
import os
import sys
import tempfile
import unittest
from importlib import import_module
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

# 直接运行测试文件时，按文件位置加入项目根目录，不依赖当前工作目录。
if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.data_loader.pdf_loader import load_pdf
from src.data_loader.docx_loader import load_docx
from src.data_loader.text_loader import load_text
from src.data_loader import batch_import, create_import_tasks, load_document
from src.chunking import split_documents
from src.chunking.fixed_chunk import split_fixed
from src.chunking.recursive_chunk import split_recursive
from src.chunking.semantic_chunk import split_semantic
from src.retrieval.vector_store import VectorStore, batch_build_index, get_embeddings
from src.retrieval.bm25_retriever import BM25Retriever, tokenize
from src.retrieval.hybrid_retriever import HybridRetriever, rrf_fusion
from src.retrieval.reranker import Reranker, get_reranker


class SmallEmbeddings(Embeddings):
    """测试使用明确的二维向量，不用于实验效果或速度结论。"""

    def __init__(self):
        self.document_calls = []
        self.query_calls = []

    def _vector(self, text):
        if "农业" in text:
            return [0.0, 1.0]
        if "反向" in text:
            return [-1.0, 0.0]
        return [1.0, 0.0]

    def embed_documents(self, texts):
        self.document_calls.append(list(texts))
        return [self._vector(text) for text in texts]

    def embed_query(self, text):
        self.query_calls.append(text)
        return self._vector(text)


class TestVectorStore(unittest.TestCase):
    """验证选定库的实际写入、查重、余弦查询、持久化和删除。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.embeddings = SmallEmbeddings()
        self.store = VectorStore(self.directory.name, self.embeddings)
        self.chunks = [
            Document(page_content="神经网络论文", metadata={
                "chunk_id": "a1", "doc_id": "a", "source_file": "论文A.pdf",
                "page_number": 2, "start_index": 0, "end_index": 7,
                "formula_layout": '[{"text":"x²","bbox":[1,2,3,4]}]',
            }),
            Document(page_content="农业数据论文", metadata={
                "chunk_id": "b1", "doc_id": "b", "source_file": "论文B.pdf", "page_number": 1,
            }),
            Document(page_content="反向向量论文", metadata={
                "chunk_id": "a2", "doc_id": "a", "source_file": "论文A.pdf", "page_number": 3,
            }),
        ]

    def test_empty_store_and_empty_query(self):
        self.assertEqual(self.store.count(), 0)
        self.assertEqual(self.store.search("神经网络"), [])
        self.store.add_chunks(self.chunks)
        self.assertEqual(self.store.search(" \n"), [])
        self.assertEqual(self.store.search("神经网络", doc_id="missing"), [])
        self.assertEqual(self.embeddings.query_calls, [])

    def test_text_and_source_roundtrip(self):
        self.assertEqual(self.store.add_chunks(self.chunks), 3)
        actual = {chunk.metadata["chunk_id"]: chunk for chunk in self.store.list_chunks()}
        self.assertEqual(actual, {chunk.metadata["chunk_id"]: chunk for chunk in self.chunks})
        self.assertEqual(len(self.store.list_chunks("a")), 2)
        self.assertEqual(self.store.list_chunks("missing"), [])

    def test_cosine_order_and_negative_score(self):
        self.store.add_chunks(self.chunks)
        found = self.store.search("神经网络", k=10)
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in found], ["a1", "b1", "a2"])
        for (_, score), expected in zip(found, (1, 0, -1)):
            self.assertAlmostEqual(score, expected)

    def test_document_filter_and_top_k(self):
        self.store.add_chunks(self.chunks)
        self.assertEqual(len(self.store.search("神经网络", k=1)), 1)
        found = self.store.search("神经网络", doc_id="b")
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in found], ["b1"])

    def test_default_top_k_and_query_change_without_reembedding_documents(self):
        """默认 K 来自配置；新查询只编码问题，沿用持久化文档向量。"""
        from src.utils.config import load_config
        config = load_config()
        config["retrieval"]["top_k"] = 2
        self.store.add_chunks(self.chunks)
        with patch("src.retrieval.vector_store.load_config", return_value=config):
            reopened = VectorStore(self.directory.name, self.embeddings)
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in reopened.search("神经网络")],
                         ["a1", "b1"])
        self.assertEqual(reopened.search("农业", k=1)[0][0], self.chunks[1])
        self.assertEqual(self.embeddings.query_calls, ["神经网络", "农业"])
        self.assertEqual(self.embeddings.document_calls, [[chunk.page_content for chunk in self.chunks]])

    def test_blank_query_does_not_read_database(self):
        """空问题直接结束，不依赖数据库连接或查询编码。"""
        with patch.object(self.store, "count", side_effect=RuntimeError("数据库不可用")):
            self.assertEqual(self.store.search(" \n\t"), [])
        self.assertEqual(self.embeddings.query_calls, [])

    def test_readonly_operations_do_not_load_model(self):
        """只读正文、重复块和无有效向量查询无需权重，真正查询才加载。"""
        self.store.add_chunks(self.chunks)
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=FileNotFoundError("模型未准备")) as model:
            readonly = VectorStore(self.directory.name)
            self.assertEqual(readonly.count(), 3)
            self.assertEqual(len(readonly.list_chunks()), 3)
            self.assertEqual(readonly.add_chunks(self.chunks), 0)
            self.assertEqual(readonly.delete_document("missing"), 0)
            self.assertEqual(readonly.search("  "), [])
            self.assertEqual(readonly.search("问题", doc_id="missing"), [])
            model.assert_not_called()
            with self.assertRaisesRegex(FileNotFoundError, "模型未准备"):
                readonly.search("神经网络")
            model.assert_called_once()

    def test_duplicate_import_does_not_encode_again(self):
        self.assertEqual(self.store.add_chunks(self.chunks + self.chunks), 3)
        self.assertEqual(self.store.add_chunks(self.chunks), 0)
        self.assertEqual(self.store.add_chunks([]), 0)
        self.assertEqual(self.embeddings.document_calls, [[chunk.page_content for chunk in self.chunks]])
        self.assertEqual(self.store.count(), 3)

    def test_incremental_import_only_encodes_new_chunks(self):
        self.store.add_chunks(self.chunks[:1])
        self.assertEqual(self.store.add_chunks(self.chunks), 2)
        self.assertEqual(self.embeddings.document_calls, [["神经网络论文"], ["农业数据论文", "反向向量论文"]])
        self.assertEqual(self.store.count(), 3)

    def test_delete_document_and_repeat_delete(self):
        self.store.add_chunks(self.chunks)
        self.assertEqual(self.store.delete_document("a"), 2)
        self.assertEqual(self.store.delete_document("a"), 0)
        self.assertEqual(self.store.count(), 1)
        self.assertEqual(self.store.list_chunks("a"), [])
        self.assertEqual([doc.metadata["doc_id"] for doc, _ in self.store.search("神经网络")], ["b"])

    def test_reopen_preserves_content_without_reembedding(self):
        self.store.add_chunks(self.chunks)
        other_embeddings = SmallEmbeddings()
        reopened = VectorStore(self.directory.name, other_embeddings)
        self.assertEqual(reopened.count(), 3)
        self.assertEqual(reopened.search("神经网络", k=1)[0][0], self.chunks[0])
        self.assertEqual(other_embeddings.document_calls, [])

    def test_changed_model_or_index_parameters_are_rejected(self):
        from src.utils.config import load_config

        for section, key, value in (("embedding", "revision", "different-version"),
                                    ("retrieval", "search_ef", 10)):
            config = load_config()
            config[section][key] = value
            with self.subTest(key=key), patch("src.retrieval.vector_store.load_config", return_value=config):
                with self.assertRaisesRegex(ValueError, "使用新索引目录重建"):
                    VectorStore(self.directory.name, SmallEmbeddings())

    def test_invalid_ids_are_rejected_before_any_write(self):
        for metadata in ({}, {"chunk_id": "x"}, {"chunk_id": "", "doc_id": "a"}):
            with self.subTest(metadata=metadata), self.assertRaises(ValueError):
                self.store.add_chunks([self.chunks[0], Document(page_content="坏块", metadata=metadata)])
        self.assertEqual(self.store.count(), 0)
        self.assertEqual(self.embeddings.document_calls, [])
        for value in (None, "", "  "):
            with self.assertRaises(ValueError):
                self.store.delete_document(value)
        for k in (0, -1, 1.5, True):
            with self.assertRaises(ValueError):
                self.store.search("问题", k=k)

    def test_same_text_in_different_documents_keeps_both_sources(self):
        second = Document(page_content=self.chunks[0].page_content,
                          metadata={"chunk_id": "c1", "doc_id": "c", "source_file": "论文C.pdf"})
        self.assertEqual(self.store.add_chunks([self.chunks[0], second]), 2)
        self.assertEqual({doc.metadata["doc_id"] for doc, _ in self.store.search("神经网络")}, {"a", "c"})

    def test_real_text_loading_and_batches_over_500(self):
        path = Path(self.directory.name) / "真实样本.md"
        path.write_text("# 神经网络\n\n论文使用农业数据。", encoding="utf-8")
        chunks = split_documents(load_document(path))
        self.assertEqual(self.store.add_chunks(chunks), len(chunks))
        self.assertEqual(self.store.search("农业")[0][0].metadata["source_file"], path.name)
        extra = [Document(page_content="神经网络", metadata={"chunk_id": f"extra-{i}", "doc_id": "extra"})
                 for i in range(501)]
        self.assertEqual(self.store.add_chunks(extra), 501)
        self.assertEqual(self.store.add_chunks(extra), 0)
        self.assertEqual(self.store.count(), 501 + len(chunks))


class TestBM25Retriever(unittest.TestCase):
    """真实 BM25 与临时 Chroma；数值按公式独立核对，不调用大模型。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.embeddings = SmallEmbeddings()
        self.store = VectorStore(self.directory.name, self.embeddings)
        self.chunks = [
            Document(page_content="BatchNormalization BatchNormalization", metadata={
                "chunk_id": "a1", "doc_id": "a", "source_file": "论文A.pdf", "page_number": 2}),
            Document(page_content="batchnormalization training training training training training", metadata={
                "chunk_id": "a2", "doc_id": "a", "source_file": "论文A.pdf", "page_number": 3}),
            Document(page_content="Adam optimizer", metadata={
                "chunk_id": "b1", "doc_id": "b", "source_file": "论文B.docx", "paragraph_index": 1}),
            Document(page_content="注意力机制", metadata={
                "chunk_id": "c1", "doc_id": "c", "source_file": "论文C.txt", "line_start": 1, "line_end": 1}),
            Document(page_content="卷积网络", metadata={
                "chunk_id": "d1", "doc_id": "d", "source_file": "论文D.md", "line_start": 1, "line_end": 1}),
        ]
        self.store.add_chunks(self.chunks)
        self.retriever = BM25Retriever(self.store)

    def test_tokenize_chinese_english_and_numbers(self):
        """沿用上游分词，保留完整英文术语，忽略标点，文档/问题统一小写。"""
        self.assertEqual(tokenize("BatchNormalization 注意力，BERT-base 2024!"),
                         ["batchnormalization", "bert", "base", "2024", "注", "意", "力"])

    def test_scores_follow_term_frequency_and_length_normalization(self):
        """用已知词频、文档频次和长度独立计算 BM25 分数，核验排序。"""
        found = self.retriever.search("BATCHNORMALIZATION", k=10)
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in found], ["a1", "a2"])
        average_length = (2 + 6 + 2 + 5 + 4) / 5
        idf = math.log((5 - 2 + 0.5) / (2 + 0.5))
        for (document, score), frequency, length in zip(found, (2, 1), (2, 6)):
            expected = idf * frequency * 2.5 / (frequency + 1.5 * (0.25 + 0.75 * length / average_length))
            self.assertAlmostEqual(score, expected)
            self.assertEqual(document, self.chunks[0 if frequency == 2 else 1])

    def test_chinese_and_source_preservation(self):
        """中文单字组合可检索，对应原文、位置和稳定 ID 全部保留。"""
        found = self.retriever.search("注意力")
        self.assertEqual([doc for doc, _ in found], [self.chunks[3]])
        self.assertGreater(found[0][1], 0)

    def test_fullwidth_pdf_terms_match_without_changing_original_text(self):
        """全角英文/数字常见于论文 PDF，词项归一化后可查且原文不变。"""
        fullwidth = Document(page_content="ＢＥＲＴ ２０２４", metadata={
            "chunk_id": "fullwidth", "doc_id": "fullwidth", "source_file": "全角论文.pdf", "page_number": 1})
        self.store.add_chunks([fullwidth])
        self.retriever.rebuild()
        for query in ("bert 2024", "ｂｅｒｔ ２０２４"):
            self.assertEqual(self.retriever.search(query)[0][0], fullwidth)

    def test_default_k_override_and_document_filter(self):
        """继承存储配置 K；文档过滤先于截断，分数仍基于完整语料。"""
        self.store.top_k = 1
        retriever = BM25Retriever(self.store)
        self.assertEqual(len(retriever.search("BatchNormalization")), 1)
        self.assertEqual(len(retriever.search("BatchNormalization", k=10)), 2)
        query = "Adam BatchNormalization"
        all_scores = {doc.metadata["chunk_id"]: score for doc, score in retriever.search(query, k=10)}
        found = retriever.search(query, doc_id="a")
        self.assertEqual(found[0][0].metadata["doc_id"], "a")
        self.assertEqual(found[0][1], all_scores[found[0][0].metadata["chunk_id"]])
        self.assertEqual(retriever.search(query, doc_id="missing"), [])

    def test_empty_and_unknown_queries_do_not_return_unrelated_documents(self):
        """无匹配时返回空列表，英文术语按词匹配而非子串匹配。"""
        for query in ("  \n", "！？…", "unknownkeyword", "Normalization"):
            with self.subTest(query=query):
                self.assertEqual(self.retriever.search(query), [])
        self.assertEqual(self.embeddings.query_calls, [])
        self.assertEqual(self.embeddings.document_calls, [[doc.page_content for doc in self.chunks]])

    def test_empty_and_punctuation_only_corpus(self):
        """空库和全无词语料不构造 BM25，混入标点块不影响有效块查询。"""
        store = VectorStore(Path(self.directory.name) / "empty", self.embeddings)
        retriever = BM25Retriever(store)
        self.assertEqual(retriever.search("Adam"), [])
        store.add_chunks([Document(page_content="！？…", metadata={"chunk_id": "p1", "doc_id": "p"})])
        self.assertEqual(retriever.rebuild(), 0)
        self.assertEqual(retriever.search("Adam"), [])
        store.add_chunks([self.chunks[2]])
        self.assertEqual(retriever.rebuild(), 1)
        self.assertEqual(retriever.search("Adam")[0][0], self.chunks[2])

    def test_zero_and_negative_scores_still_keep_keyword_matches(self):
        """复现上游 >0 过滤的漏检：单文档负分、两文档零分均应命中。"""
        for index, texts in enumerate((["BatchNormalization"], ["BatchNormalization", "Adam"])):
            with self.subTest(texts=texts):
                store = VectorStore(Path(self.directory.name) / f"small-{index}", self.embeddings)
                store.add_chunks([Document(page_content=text, metadata={"chunk_id": str(i), "doc_id": str(i)})
                                  for i, text in enumerate(texts)])
                found = BM25Retriever(store).search("BatchNormalization")
                self.assertEqual(len(found), 1)
                self.assertEqual(found[0][0].page_content, "BatchNormalization")
                if index == 0:
                    self.assertLess(found[0][1], 0)
                else:
                    self.assertEqual(found[0][1], 0)

    def test_rebuild_after_add_delete_and_reopen(self):
        """新增/删除后从当前正文重建，重开自动恢复，不重新编码文档。"""
        new = Document(page_content="Reranker", metadata={"chunk_id": "new", "doc_id": "new"})
        self.store.add_chunks([new])
        self.assertEqual(self.retriever.rebuild(), 6)
        self.assertEqual(self.retriever.search("Reranker")[0][0], new)
        self.store.delete_document("a")
        self.assertEqual(self.retriever.rebuild(), 4)
        self.assertEqual(self.retriever.search("BatchNormalization"), [])
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("BM25 不应加载模型")):
            reopened = BM25Retriever(VectorStore(self.directory.name))
            self.assertEqual(reopened.search("Reranker"), self.retriever.search("Reranker"))
        self.assertEqual(self.embeddings.query_calls, [])
        self.assertEqual(len(self.embeddings.document_calls), 2)

    def test_same_content_keeps_separate_document_sources(self):
        """不同来源的相同正文保留各自 ID，供后续按块融合使用。"""
        self.store.add_chunks([Document(page_content="Reranker", metadata={"chunk_id": f"r{i}", "doc_id": f"r{i}"})
                               for i in range(2)])
        self.retriever.rebuild()
        self.assertEqual({doc.metadata["doc_id"] for doc, _ in self.retriever.search("Reranker")}, {"r0", "r1"})

    def test_invalid_k_and_database_errors_are_not_hidden(self):
        """参数错误与正文读取失败直接上报，不静默回退向量查询。"""
        for k in (0, -1, 1.5, True):
            with self.subTest(k=k), self.assertRaises(ValueError):
                self.retriever.search("Adam", k=k)
        with patch.object(self.store, "list_chunks", side_effect=RuntimeError("数据库不可用")):
            with self.assertRaisesRegex(RuntimeError, "数据库不可用"):
                BM25Retriever(self.store)


class TestRRFFusion(unittest.TestCase):
    """用明确排名和手算分数核验融合，不用原始分数冒充排名贡献。"""

    def setUp(self):
        self.documents = {key: Document(page_content=f"论文 {key}", metadata={
            "chunk_id": key, "doc_id": key, "source_file": f"{key}.pdf", "page_number": 1})
            for key in ("a", "b", "c", "d")}

    def test_one_based_rank_formula_and_shared_chunk_bonus(self):
        """两路共同命中的块分数累加，第一名贡献为 1/61。"""
        d = self.documents
        found = rrf_fusion([(d["a"], 0.9), (d["b"], 0.8), (d["c"], -0.2)],
                           [(d["b"], 8.0), (d["c"], 2.0), (d["d"], 0.0)])
        expected = {"a": 1 / 61, "b": 1 / 62 + 1 / 61,
                    "c": 1 / 63 + 1 / 62, "d": 1 / 63}
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in found], ["b", "c", "a", "d"])
        for document, score in found:
            self.assertAlmostEqual(score, expected[document.metadata["chunk_id"]])
            self.assertIs(document, d[document.metadata["chunk_id"]])

    def test_raw_score_scale_does_not_affect_fusion(self):
        """负分、零分与大数的原始值都不参与 RRF，排名相同则结果相同。"""
        d = self.documents
        original = rrf_fusion([(d["a"], 0.9), (d["b"], 0.2)], [(d["b"], 5), (d["c"], 0)])
        changed = rrf_fusion([(d["a"], 999), (d["b"], -1)], [(d["b"], -200), (d["c"], -800)])
        self.assertEqual(original, changed)

    def test_duplicate_within_one_route_counts_only_once(self):
        """单路重复不多次加分，保留首次出现的位置，两路贡献仍可叠加。"""
        d = self.documents
        found = rrf_fusion([(d["a"], 1), (d["a"], 1), (d["b"], 0.5)],
                           [(d["a"], 2), (d["a"], 2)])
        self.assertEqual(len(found), 2)
        self.assertAlmostEqual(found[0][1], 2 / 61)
        self.assertAlmostEqual(found[1][1], 1 / 63)

    def test_same_content_in_different_sources_is_not_merged(self):
        """合并键是 chunk_id，正文相同也保留不同文件与页码。"""
        a = Document(page_content="相同正文", metadata={"chunk_id": "a", "doc_id": "a",
                     "source_file": "A.pdf", "page_number": 2, "page_end": 3})
        b = Document(page_content="相同正文", metadata={"chunk_id": "b", "doc_id": "b",
                     "source_file": "B.docx", "table_index": 1})
        self.assertEqual([doc for doc, _ in rrf_fusion([(a, 1)], [(b, 2)])], [a, b])

    def test_equal_scores_have_stable_chunk_id_order(self):
        """两路交换同分候选的输入顺序，仍按稳定块 ID 打破平分。"""
        d = self.documents
        first = rrf_fusion([(d["b"], 1)], [(d["a"], 2)])
        second = rrf_fusion([(d["a"], 2)], [(d["b"], 1)])
        self.assertEqual(first, second)
        self.assertEqual([doc.metadata["chunk_id"] for doc, _ in first], ["a", "b"])

    def test_empty_routes_and_zero_smoothing(self):
        """单路为空时只使用另一支，双空返回空；k=0 时分母从 1 开始。"""
        d = self.documents
        route = [(d["a"], -1), (d["b"], 0)]
        self.assertEqual(rrf_fusion([], []), [])
        self.assertEqual(rrf_fusion(route, []), rrf_fusion([], route))
        self.assertEqual([score for _, score in rrf_fusion(route, [], rrf_k=0)], [1, 0.5])

    def test_missing_ids_and_invalid_smoothing_are_rejected(self):
        """身份缺失不能退回按正文合并；平滑常数的错误输入明确报错。"""
        for value in (None, "", "  ", 1):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "chunk_id"):
                rrf_fusion([(Document(page_content="正文", metadata={"chunk_id": value}), 1)], [])
        for value in (-1, 1.5, True):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "rrf_k"):
                rrf_fusion([], [], rrf_k=value)


class TestHybridRetriever(unittest.TestCase):
    """真实 Chroma/BM25 联调，两维向量仅用于确定行为与模型调用次数。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.embeddings = SmallEmbeddings()
        self.store = VectorStore(self.directory.name, self.embeddings)
        self.chunks = [Document(page_content=text, metadata={
            "chunk_id": str(i), "doc_id": str(i), "source_file": f"论文{i}.pdf", "page_number": i + 1})
            for i, text in enumerate(("神经网络研究", "BatchNormalization BatchNormalization",
                                     "BatchNormalization training training training", "农业数据", "反向向量"))]
        self.store.add_chunks(self.chunks)
        self.retriever = HybridRetriever(self.store)

    def test_both_routes_take_candidates_before_final_top_k(self):
        """最终只取 1 个也先各召回 20 个候选，不提前把两路截为 Top-1。"""
        original_bm25_search = BM25Retriever.search
        with (patch.object(self.store, "search", wraps=self.store.search) as vector,
              patch.object(BM25Retriever, "search", autospec=True, side_effect=original_bm25_search) as bm25):
            found = self.retriever.search("BatchNormalization", k=1)
        vector.assert_called_once_with("BatchNormalization", k=20, doc_id=None)
        self.assertEqual(bm25.call_args.kwargs, {"k": 20, "doc_id": None})
        self.assertIs(bm25.call_args.args[0].vector_store, self.store)
        self.assertEqual(len(found), 1)
        self.assertIn(found[0][0], self.chunks[1:3])
        self.assertGreater(found[0][1], 1 / 61)
        self.assertEqual(self.embeddings.query_calls, ["BatchNormalization"])
        self.assertEqual(self.embeddings.document_calls, [[doc.page_content for doc in self.chunks]])

    def test_document_filter_and_empty_bm25_route(self):
        """两路共同过滤文档，BM25 无匹配时保留向量路的排名贡献。"""
        found = self.retriever.search("BatchNormalization", doc_id="3")
        self.assertEqual([doc for doc, _ in found], [self.chunks[3]])
        self.assertAlmostEqual(found[0][1], 1 / 61)
        before = len(self.embeddings.query_calls)
        self.assertEqual(self.retriever.search("BatchNormalization", doc_id="missing"), [])
        self.assertEqual(len(self.embeddings.query_calls), before)

    def test_blank_query_and_empty_store_do_not_load_model(self):
        """空问题不调用两路，空库不会初始化 Embedding。"""
        with patch.object(self.store, "search", side_effect=AssertionError("不应查询")):
            self.assertEqual(self.retriever.search(" \n"), [])
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("不应加载模型")):
            empty = HybridRetriever(VectorStore(Path(self.directory.name) / "empty"))
            self.assertEqual(empty.search("BERT"), [])

    def test_reused_retriever_reads_current_corpus_after_add_and_delete(self):
        """同一个混合检索实例也读取最新正文，新增和删除后无陈旧 BM25 块。"""
        new = Document(page_content="Reranker", metadata={"chunk_id": "new", "doc_id": "new"})
        self.store.add_chunks([new])
        self.assertEqual(self.retriever.search("Reranker", k=1)[0][0], new)
        self.store.delete_document("new")
        found = self.retriever.search("Reranker", k=20)
        self.assertEqual({doc.metadata["chunk_id"] for doc, _ in found}, {str(i) for i in range(5)})
        self.assertEqual(len(self.embeddings.document_calls), 2)

    def test_default_k_large_k_and_yaml_smoothing(self):
        """默认 K 与存储一致；更大的 K 扩大召回，平滑常数从 YAML 读取。"""
        from src.utils.config import load_config
        config = load_config()
        config["retrieval"]["rrf_k"] = 0
        self.store.top_k = 2
        with patch("src.retrieval.hybrid_retriever.load_config", return_value=config):
            retriever = HybridRetriever(self.store)
        self.assertEqual(len(retriever.search("unknownkeyword")), 2)
        self.assertEqual([score for _, score in retriever.search("unknownkeyword")], [1, 0.5])
        extra = [Document(page_content="BatchNormalization", metadata={"chunk_id": f"extra-{i}", "doc_id": "extra"})
                 for i in range(41)]
        self.store.add_chunks(extra)
        found = retriever.search("BatchNormalization", k=100)
        self.assertEqual(len(found), 46)
        self.assertEqual(len({doc.metadata["chunk_id"] for doc, _ in found}), 46)

    def test_invalid_parameters_and_route_errors_do_not_silently_degrade(self):
        """参数错误直接拒绝，两路异常明确上报，不能以半成品冒充混合检索。"""
        from src.utils.config import load_config
        for value in (0, -1, True, 1.5):
            with self.subTest(k=value), self.assertRaises(ValueError):
                self.retriever.search("BERT", k=value)
        for key, value in (("candidate_k", 0), ("candidate_k", True), ("rrf_k", -1)):
            config = load_config()
            config["retrieval"][key] = value
            with self.subTest(key=key), patch("src.retrieval.hybrid_retriever.load_config", return_value=config):
                with self.assertRaisesRegex(ValueError, key):
                    HybridRetriever(self.store)
        with patch.object(self.store, "search", side_effect=RuntimeError("向量查询失败")):
            with self.assertRaisesRegex(RuntimeError, "向量查询失败"):
                self.retriever.search("BERT")
        with patch.object(BM25Retriever, "search", side_effect=RuntimeError("BM25 查询失败")):
            with self.assertRaisesRegex(RuntimeError, "BM25 查询失败"):
                self.retriever.search("BERT")

    def test_model_reranks_fused_candidates_before_final_top_k(self):
        """模型先看到融合 Top-20；候选末项也能经模型升到最终 Top-1。"""
        self.store.add_chunks([Document(page_content=f"论文候选 {i}", metadata={
            "chunk_id": f"extra-{i}", "doc_id": "extra"}) for i in range(25)])
        candidates = self.retriever.search("unknownkeyword", k=20)
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.01] * 19 + [0.99]
            found = self.retriever.search("unknownkeyword", k=1, rerank=True)
            self.assertEqual(found, [(candidates[-1][0], 0.99)])
            model.return_value.predict.assert_called_once_with(
                [["unknownkeyword", doc.page_content] for doc, _ in candidates],
                batch_size=8, show_progress_bar=False)
            # K 大于默认候选数时相应扩大，而不是返回不足 K 的人为截断结果。
            model.return_value.predict.return_value = [0.5] * 25
            found = self.retriever.search("unknownkeyword", k=25, rerank=True)
            self.assertEqual(len(found), 25)
        self.assertEqual(len(self.embeddings.document_calls), 2)

    def test_model_reranking_respects_filter_and_empty_results(self):
        """过滤先应用于两路，单条候选也精排；空结果和空问题不加载模型。"""
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.7]
            found = self.retriever.search("BatchNormalization", doc_id="3", rerank=True)
            self.assertEqual(found, [(self.chunks[3], 0.7)])
            model.return_value.predict.assert_called_once_with(
                [["BatchNormalization", self.chunks[3].page_content]], batch_size=8, show_progress_bar=False)
            model.reset_mock()
            self.assertEqual(self.retriever.search("BatchNormalization", doc_id="missing", rerank=True), [])
            self.assertEqual(self.retriever.search("  ", rerank=True), [])
            model.assert_not_called()

    def test_plain_rrf_does_not_load_model_and_model_failure_is_reported(self):
        """不启用精排时不依赖重排权重；启用后失败不得返回未经重排的候选。"""
        with patch("src.retrieval.reranker.get_reranker", side_effect=RuntimeError("重排推理失败")) as model:
            self.assertTrue(self.retriever.search("BERT"))
            model.assert_not_called()
            with self.assertRaisesRegex(RuntimeError, "重排推理失败"):
                self.retriever.search("BERT", rerank=True)


class TestReranker(unittest.TestCase):
    """大模型隔离，核验成对输入、排序与来源，不将模拟分数当实验效果。"""

    def setUp(self):
        from src.utils.config import load_config
        self.config = load_config()
        self.config["retrieval"]["top_k"] = 2
        self.config["retrieval"]["reranker_batch_size"] = 2
        patcher = patch("src.retrieval.reranker.load_config", return_value=self.config)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.reranker = Reranker()
        self.candidates = [(Document(page_content=text, metadata={
            "chunk_id": str(i), "source_file": f"论文{i}.pdf", "page_number": i + 1}), 10 - i)
            for i, text in enumerate(("无关材料", "注意力机制", "Attention mechanism"))]

    def test_pairs_model_scores_order_and_sources(self):
        """只用模型分数精排，同分保持原候选顺序，正文和来源不修改。"""
        original = [(doc.page_content, dict(doc.metadata)) for doc, _ in self.candidates]
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.1, 0.9, 0.9]
            found = self.reranker.rerank("注意力是什么？", self.candidates)
            model.return_value.predict.assert_called_once_with(
                [["注意力是什么？", doc.page_content] for doc, _ in self.candidates],
                batch_size=2, show_progress_bar=False)
        self.assertEqual(found, [(self.candidates[1][0], 0.9), (self.candidates[2][0], 0.9)])
        self.assertIs(found[0][0], self.candidates[1][0])
        self.assertEqual(original, [(doc.page_content, doc.metadata) for doc, _ in self.candidates])

    def test_single_candidate_zero_score_and_large_k(self):
        """单候选结果仍为列表，零分保留，K 超出候选数不补齐或重复。"""
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.0]
            self.assertEqual(self.reranker.rerank("query", self.candidates[:1], k=100),
                             [(self.candidates[0][0], 0.0)])

    def test_empty_query_or_candidates_do_not_load_model(self):
        with patch("src.retrieval.reranker.get_reranker") as model:
            self.assertEqual(self.reranker.rerank("  \n", self.candidates), [])
            self.assertEqual(self.reranker.rerank("query", []), [])
            model.assert_not_called()

    def test_invalid_k_is_rejected_before_loading_model(self):
        for value in (0, -1, True, 1.5):
            with self.subTest(k=value), patch("src.retrieval.reranker.get_reranker") as model:
                with self.assertRaisesRegex(ValueError, "k"):
                    self.reranker.rerank("query", self.candidates, k=value)
                model.assert_not_called()

    def test_invalid_batch_and_token_length_are_rejected(self):
        for key in ("reranker_batch_size", "reranker_max_length"):
            old = self.config["retrieval"][key]
            for value in (0, -1, True, 1.5):
                with self.subTest(key=key, value=value):
                    self.config["retrieval"][key] = value
                    with self.assertRaisesRegex(ValueError, key):
                        Reranker()
            self.config["retrieval"][key] = old

    def test_score_count_and_nonfinite_values_are_rejected(self):
        for scores in ([0.1], [float("nan"), 0.1, 0.2], [float("inf"), 0.1, 0.2]):
            with self.subTest(scores=scores), patch("src.retrieval.reranker.get_reranker") as model:
                model.return_value.predict.return_value = scores
                with self.assertRaisesRegex(ValueError, "分数"):
                    self.reranker.rerank("query", self.candidates)

    def test_inference_error_is_not_silently_replaced_by_rrf(self):
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.side_effect = RuntimeError("推理失败")
            with self.assertRaisesRegex(RuntimeError, "推理失败"):
                self.reranker.rerank("query", self.candidates)


class TestLocalReranker(unittest.TestCase):
    """核验本地权重加载、缓存与失败恢复，测试不下载真实模型。"""

    def setUp(self):
        from src.utils.config import load_config
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.config = load_config()
        self.config["retrieval"]["reranker_local_path"] = self.directory.name
        patcher = patch.dict(os.environ, {"HF_HOME": str(Path(self.directory.name) / "cache")})
        patcher.start()
        self.addCleanup(patcher.stop)
        get_reranker.cache_clear()
        self.addCleanup(get_reranker.cache_clear)
        patcher = patch("src.retrieval.reranker.load_config", return_value=self.config)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_local_loading_and_reuse(self):
        """配置/模型/Tokenizer 均只读本地文件，不执行远程代码，重复调用只加载一次。"""
        from torch.nn import Sigmoid
        with patch("sentence_transformers.CrossEncoder") as model:
            self.assertIs(get_reranker(), model.return_value)
            self.assertIs(get_reranker(), model.return_value)
            model.assert_called_once()
            self.assertEqual(model.call_args.args, (self.directory.name,))
            kwargs = dict(model.call_args.kwargs)
            self.assertIsInstance(kwargs.pop("default_activation_function"), Sigmoid)
            self.assertEqual(kwargs, {"device": "cpu", "max_length": 512,
                                     "local_files_only": True, "trust_remote_code": False})

    def test_missing_model_does_not_initialize_remote_client(self):
        self.config["retrieval"]["reranker_local_path"] = str(Path(self.directory.name) / "missing")
        with patch("sentence_transformers.CrossEncoder") as model:
            with self.assertRaisesRegex(FileNotFoundError, "请先按用户手册下载权重"):
                get_reranker()
            model.assert_not_called()

    def test_failed_loading_is_not_cached_and_can_retry(self):
        """损坏权重错误保留，修复后下次调用重新加载而不是缓存失败。"""
        with patch("sentence_transformers.CrossEncoder") as model:
            model.side_effect = OSError("损坏的权重")
            with self.assertRaisesRegex(OSError, "损坏的权重"):
                get_reranker()
            model.side_effect = None
            self.assertIs(get_reranker(), model.return_value)
            self.assertEqual(model.call_count, 2)


class TestLocalEmbeddings(unittest.TestCase):
    """隔离大模型，验证配置、本地约束和延迟加载缓存。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        get_embeddings.cache_clear()
        self.addCleanup(get_embeddings.cache_clear)
        self.config = {"embedding": {
            "local_path": self.directory.name, "device": "cpu", "batch_size": 32,
        }}
        patcher = patch("src.retrieval.vector_store.load_config", return_value=self.config)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_local_loading_and_normalization_arguments(self):
        """只从本地目录加载，关闭远程代码，文档编码开启归一化。"""
        with patch("src.retrieval.vector_store.HuggingFaceEmbeddings") as model:
            self.assertIs(get_embeddings(), model.return_value)
            model.assert_called_once_with(
                model_name=self.directory.name,
                model_kwargs={"device": "cpu", "local_files_only": True, "trust_remote_code": False},
                encode_kwargs={"normalize_embeddings": True, "batch_size": 32},
            )

    def test_repeated_calls_reuse_model(self):
        """多次调用复用已加载模型，不在每次提问时加载权重。"""
        with patch("src.retrieval.vector_store.HuggingFaceEmbeddings") as model:
            first = get_embeddings()
            self.assertIs(get_embeddings(), first)
            model.assert_called_once()

    def test_missing_model_does_not_initialize_remote_client(self):
        """缺少权重明确报错，不构造模型客户端或回退在线模式。"""
        self.config["embedding"]["local_path"] = str(Path(self.directory.name) / "missing")
        with patch("src.retrieval.vector_store.HuggingFaceEmbeddings") as model:
            with self.assertRaisesRegex(FileNotFoundError, "请先按用户手册下载权重"):
                get_embeddings()
            model.assert_not_called()


class TestEmbeddingEvaluation(unittest.TestCase):
    """用已知排名验证指标公式，模型效果来自独立的真实实验。"""

    def test_hit_recall_and_mrr_cutoffs(self):
        """排名六的相关项不算 Hit@5，排名十一的项不算 MRR@10。"""
        import numpy as np
        evaluate_rankings = import_module("reports.5_1_3 向量化与存储.compare_embeddings").evaluate_rankings

        sample = {"corpus": [{"id": str(index)} for index in range(11)], "queries": [
            {"id": "q1", "relevant_ids": ["0"]},
            {"id": "q2", "relevant_ids": ["5", "10"]},
        ]}
        scores = np.asarray([list(range(11, 0, -1))] * 2)
        result = evaluate_rankings(scores, sample)
        self.assertEqual(result["hit_at_5"], 0.5)
        self.assertEqual(result["recall_at_5"], 0.5)
        self.assertAlmostEqual(result["mrr_at_10"], (1 + 1 / 6) / 2)
        self.assertEqual(result["queries"][1]["top10_ids"], [str(index) for index in range(10)])

    def test_language_groups_use_equal_weight_macro_average(self):
        """问题数不等时宏平均仍等权，并保留表现较差的语言组。"""
        import numpy as np
        evaluate_rankings = import_module("reports.5_1_3 向量化与存储.compare_embeddings").evaluate_rankings

        sample = {"corpus": [{"id": str(index)} for index in range(6)], "queries": [
            {"id": "zh1", "group": "zh->zh", "relevant_ids": ["0"]},
            {"id": "zh2", "group": "zh->zh", "relevant_ids": ["5"]},
            {"id": "en1", "group": "en->en", "relevant_ids": ["5"]},
        ]}
        result = evaluate_rankings(np.asarray([list(range(6, 0, -1))] * 3), sample)
        self.assertAlmostEqual(result["hit_at_5"], 1 / 3)
        self.assertEqual(result["groups"]["zh->zh"]["query_count"], 2)
        self.assertEqual(result["groups"]["zh->zh"]["hit_at_5"], 0.5)
        self.assertEqual(result["macro"]["hit_at_5"], 0.25)
        self.assertEqual(result["worst_group_hit_at_5"], 0)


class TestRetrievalEvaluation(unittest.TestCase):
    """手算 Top-5 边界、相关块集合与宏平均，实验质量来自真实模型。"""

    def test_rank_five_counts_and_rank_six_does_not(self):
        evaluate_ranking = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").evaluate_ranking
        ranking = [str(i) for i in range(8)]
        found = evaluate_ranking(ranking, ["4", "7"])
        self.assertTrue(found["hit_at_5"])
        self.assertEqual(found["recall_at_5"], 0.5)
        self.assertEqual(found["mrr_at_5"], 1 / 5)
        self.assertEqual(found["first_relevant_rank"], 5)
        self.assertFalse(evaluate_ranking(ranking, ["5"])["hit_at_5"])
        self.assertEqual(evaluate_ranking(ranking, ["5"])["mrr_at_5"], 0)

    def test_multiple_relevant_short_and_empty_results(self):
        evaluate_ranking = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").evaluate_ranking
        found = evaluate_ranking(["a", "b", "c"], ["b", "c", "z"])
        self.assertEqual(found["relevant_returned"], 2)
        self.assertEqual(found["recall_at_5"], 2 / 3)
        self.assertEqual(found["mrr_at_5"], 1 / 2)
        self.assertEqual(evaluate_ranking([], ["a"])["recall_at_5"], 0)

    def test_empty_annotations_and_duplicate_results_are_errors(self):
        evaluate_ranking = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").evaluate_ranking
        for ranking, labels in ((["a"], []), (["a", "a"], ["a"])):
            with self.subTest(ranking=ranking), self.assertRaises(ValueError):
                evaluate_ranking(ranking, labels)

    def test_unequal_groups_and_repeated_timing_are_not_duplicate_queries(self):
        evaluate_ranking = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").evaluate_ranking
        summarize = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").summarize
        rows = [{**evaluate_ranking(ids, ["a"]), "group": group, "latency_ms_runs": times}
                for ids, group, times in ((["a"], "zh->zh", [10, 20]),
                                         (["b"], "zh->zh", [30, 40]),
                                         (["a"], "en->en", [50, 60]))]
        result = summarize(rows)
        self.assertEqual(result["query_count"], 3)
        self.assertEqual(result["hit_count"], 2)
        self.assertEqual(result["hit_at_5"], 2 / 3)
        self.assertEqual(result["macro"]["hit_at_5"], 0.75)
        self.assertEqual(result["worst_group_hit_at_5"], 0.5)
        self.assertEqual(result["latency_mean_ms"], 35)
        self.assertEqual(result["latency_p95_ms"], 57.5)

    def test_sample_ids_and_annotation_membership_are_validated(self):
        validate_sample = import_module("reports.5_1_4 混合检索与重排序.compare_retrieval").validate_sample
        sample = {"corpus": [{"id": "a"}], "queries": [{"id": "q", "relevant_ids": ["a"]}]}
        validate_sample(sample)
        sample["queries"][0]["relevant_ids"] = ["missing"]
        with self.assertRaisesRegex(ValueError, "相关标注"):
            validate_sample(sample)
        sample["queries"][0]["relevant_ids"] = ["a"]
        sample["corpus"].append({"id": "a"})
        with self.assertRaisesRegex(ValueError, "重复 ID"):
            validate_sample(sample)


class TestChunking(unittest.TestCase):
    """用真实切分器检查内容、边界与溯源，不依赖模型和检索库。"""

    strategies = (split_fixed, split_recursive, split_semantic)

    def test_short_document(self):
        """短文档保持原文，包括缩进与尾部换行。"""
        text = "  中文论文摘要。\r\n"
        for splitter in self.strategies:
            with self.subTest(strategy=splitter.__name__):
                chunks = splitter([Document(page_content=text)], 32, 4)
                self.assertEqual([chunk.page_content for chunk in chunks], [text])
                self.assertEqual(chunks[0].metadata["start_index"], 0)
                self.assertEqual(chunks[0].metadata["end_index"], len(text))

    def test_empty_documents(self):
        """空列表、空文本和仅空白文档不生成无效块。"""
        for splitter in self.strategies:
            self.assertEqual(splitter([], 16, 0), [])
            self.assertEqual(splitter([Document(page_content=text) for text in ("", " \n\t")], 16, 0), [])

    def test_invalid_parameters(self):
        """拒绝不合法大小和重叠，避免零步长或负数切片。"""
        for splitter in self.strategies:
            for size, overlap in ((0, 0), (-1, 0), (4, -1), (4, 4), (4, 5),
                                  (4.5, 0), (4, 1.5), (True, 0), (4, False)):
                with self.subTest(strategy=splitter.__name__, size=size, overlap=overlap):
                    with self.assertRaises(ValueError):
                        splitter([], size, overlap)

    def test_fixed_exact_overlap(self):
        """固定窗口按字符数切分，相邻正文块精确重叠。"""
        text = "甲乙丙丁戊己庚辛壬癸"
        chunks = split_fixed([Document(page_content=text)], 4, 1)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲乙丙丁", "丁戊己庚", "庚辛壬癸"])
        self.assertEqual([chunk.metadata["start_index"] for chunk in chunks], [0, 3, 6])

    def test_fixed_no_redundant_tail(self):
        """末尾已被覆盖时不再输出完全包含在上一块中的尾块。"""
        for text, expected in (("12345678", ["12345678"]), ("1234567890", ["12345678", "67890"])):
            chunks = split_fixed([Document(page_content=text)], 8, 3)
            self.assertEqual([chunk.page_content for chunk in chunks], expected)

    def test_fixed_course_sizes(self):
        """课程规定的 256/512/1024 都可切分，正文长度不超过设置。"""
        text = "科研文本" * 600
        for size in (256, 512, 1024):
            chunks = split_fixed([Document(page_content=text)], size, 64)
            self.assertGreater(len(chunks), 1)
            self.assertTrue(all(len(chunk.page_content) <= size for chunk in chunks))
            self.assertEqual(chunks[-1].metadata["end_index"], len(text))

    def test_recursive_paragraph_boundaries(self):
        """优先按段落分隔符切分，并保留分隔符原文。"""
        text = "甲段内容。\n\n乙段内容。\n\n丙段内容。"
        chunks = split_recursive([Document(page_content=text)], 9, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["甲段内容。", "\n\n乙段内容。", "\n\n丙段内容。"])
        self.assertEqual("".join(chunk.page_content for chunk in chunks), text)

    def test_recursive_character_fallback(self):
        """没有分隔符的长文本仍能限长并保留重叠。"""
        chunks = split_recursive([Document(page_content="ABCDEFGHIJ")], 4, 1)
        self.assertEqual([chunk.page_content for chunk in chunks], ["ABCD", "DEFG", "GHIJ"])

    def test_semantic_keeps_paragraphs(self):
        """短段落整体保留，空行跟随原段落。"""
        text = "甲段内容。\r\n\r\n乙段内容。\r\n\r\n丙段内容。"
        chunks = split_semantic([Document(page_content=text)], 10, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["甲段内容。\r\n\r\n", "乙段内容。\r\n\r\n", "丙段内容。"])

    def test_semantic_chinese_sentences_and_quotes(self):
        """长段落拆句，结束引号与标点跟随原句。"""
        text = "第一句。”第二句！第三句？第四句。"
        chunks = split_semantic([Document(page_content=text)], 8, 0)
        self.assertEqual([chunk.page_content for chunk in chunks], ["第一句。”", "第二句！第三句？", "第四句。"])

    def test_semantic_english_decimal_and_doi(self):
        """英文句号可拆句，小数和 DOI 内部的点不误拆。"""
        text = "Value 3.14. DOI 10.1000/xyz. Result good."
        chunks = split_semantic([Document(page_content=text)], 27, 0)
        self.assertEqual([chunk.page_content for chunk in chunks],
                         ["Value 3.14.", " DOI 10.1000/xyz.", " Result good."])

    def test_semantic_complete_unit_overlap(self):
        """重叠复用完整句子；不足容纳整句时不截出半句。"""
        document = Document(page_content="甲甲。乙乙。丙丙。丁丁。")
        chunks = split_semantic([document], 6, 3)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲甲。乙乙。", "乙乙。丙丙。", "丙丙。丁丁。"])
        chunks = split_semantic([document], 6, 2)
        self.assertEqual([chunk.page_content for chunk in chunks], ["甲甲。乙乙。", "丙丙。丁丁。"])

    def test_semantic_long_sentence_fallback(self):
        """超长单句按字符兜底，最后短片段不丢失。"""
        text = "没有任何句号的超长文本用于测试"
        chunks = split_semantic([Document(page_content=text)], 5, 2)
        self.assertEqual([chunk.page_content for chunk in chunks], [text[i:i + 5] for i in range(0, len(text), 5)])

    def test_ranges_cover_original_content(self):
        """不同大小和重叠下，区间准确、向前推进，所有非空白字符均被覆盖。"""
        texts = ("重复段落。\n\n" * 30, "ABCD" * 50, "甲。乙！丙？\r\n\r\n" * 20)
        for splitter in self.strategies:
            for text in texts:
                for size, overlap in ((1, 0), (8, 0), (8, 2), (16, 15), (64, 10)):
                    with self.subTest(strategy=splitter.__name__, size=size, overlap=overlap, text=text[:12]):
                        chunks = splitter([Document(page_content=text)], size, overlap)
                        covered = set()
                        previous_start = -1
                        for chunk in chunks:
                            start, end = chunk.metadata["start_index"], chunk.metadata["end_index"]
                            self.assertGreater(start, previous_start)
                            self.assertLessEqual(len(chunk.page_content), size)
                            self.assertEqual(chunk.page_content, text[start:end])
                            covered.update(range(start, end))
                            previous_start = start
                        self.assertTrue(all(i in covered for i, char in enumerate(text) if not char.isspace()))

    def test_preserves_metadata_without_mutating_input(self):
        """保留 PDF 页码、版面和公式原位置，并复制元数据。"""
        metadata = {"doc_id": "论文ID", "source_file": "论文.pdf", "page": 2, "page_number": 3,
                    "layout": "two_column", "formula_layout": '[{"text":"x²"}]'}
        document = Document(page_content="公式 x^{2}。\n实验结果。" * 5, metadata=metadata)
        for splitter in self.strategies:
            chunks = splitter([document], 20, 2)
            for chunk in chunks:
                for key, value in metadata.items():
                    self.assertEqual(chunk.metadata[key], value)
            chunks[0].metadata["source_file"] = "改名.pdf"
            self.assertEqual(document.metadata, metadata)
            self.assertNotIn("chunk_id", document.metadata)

    def test_stable_distinct_chunk_ids(self):
        """重复执行 ID 不变，重复正文在不同页/段落及不同配置中不会混淆。"""
        documents = [Document(page_content="重复内容" * 8, metadata={"doc_id": "id", "page": page})
                     for page in (0, 1)]
        for splitter in self.strategies:
            chunks = splitter(documents, 12, 2)
            ids = [chunk.metadata["chunk_id"] for chunk in chunks]
            self.assertEqual(ids, [chunk.metadata["chunk_id"] for chunk in splitter(documents, 12, 2)])
            self.assertEqual(len(set(ids)), len(ids))
            changed = {chunk.metadata["chunk_id"] for chunk in splitter(documents, 16, 2)}
            self.assertTrue(set(ids).isdisjoint(changed))
        word = [Document(page_content="重复段落", metadata={"doc_id": "word", "block_index": index})
                for index in (1, 2)]
        self.assertEqual(len({chunk.metadata["chunk_id"] for chunk in split_fixed(word, 8)}), 2)

    def test_pdf_and_word_tables_remain_whole(self):
        """独立表格不拆，即使超长也保留跨页行信息和保护标记。"""
        rows = '[{"cells":["数据"],"page_number":4},{"cells":["结果"],"page_number":5}]'
        tables = [Document(page_content="| 表头 |\n| --- |\n| 数据 |\n| 结果 |", metadata={
            "doc_id": "id", "content_type": "table", "table_id": "id:table:1",
            "page": 3, "page_number": 4, "page_end": 5, "table_rows": rows,
        }), Document(page_content="表头\t值\n数据\t结果", metadata={"block_type": "table", "block_index": 2})]
        for splitter in self.strategies:
            chunks = splitter(tables, 4, 1)
            self.assertEqual([chunk.page_content for chunk in chunks], [table.page_content for table in tables])
            self.assertTrue(all(chunk.metadata["chunk_preserved"] == "table" for chunk in chunks))
            self.assertEqual(chunks[0].metadata["table_rows"], rows)
            self.assertEqual(chunks[0].metadata["page_end"], 5)

    def test_loaded_text_line_numbers(self):
        """真实 TXT 加载后细化 CRLF 行范围；块内原文和文件来源保持一致。"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "文本.txt"
            path.write_bytes("甲乙\r\n丙丁\r\n戊己".encode("utf-8"))
            documents = load_text(path)
            chunks = split_fixed(documents, 4, 0)
            self.assertEqual([(chunk.metadata["line_start"], chunk.metadata["line_end"]) for chunk in chunks],
                             [(1, 1), (2, 2), (3, 3)])
            self.assertTrue(all(chunk.metadata["source"] == str(path.resolve()) for chunk in chunks))
            self.assertEqual(documents[0].metadata["line_end"], 3)
            # 恰好从 CRLF 中间开始的块仍归原行，不提前增加行号。
            chunks = split_fixed(documents, 3, 0)
            self.assertEqual([(chunk.metadata["line_start"], chunk.metadata["line_end"]) for chunk in chunks],
                             [(1, 1), (1, 2), (2, 3), (3, 3)])

    def test_loaded_word_and_pdf(self):
        """真实 Word/PDF 加载结果可直接分块，来源位置与表格不丢失。"""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "论文.docx"
            word = WordDocument()
            word.add_paragraph("研究背景。实验方法。研究结论。" * 3)
            word.add_table(rows=1, cols=2).cell(0, 0).text = "完整表格"
            word.save(path)
            for splitter in self.strategies:
                chunks = splitter(load_docx(path), 10, 2)
                self.assertTrue(all("block_index" in chunk.metadata and "page" not in chunk.metadata for chunk in chunks))
                self.assertEqual(chunks[-1].page_content, "完整表格\t")
                self.assertEqual(chunks[-1].metadata["chunk_preserved"], "table")
            path = Path(directory) / "论文.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "研究背景。实验方法。研究结论。", fontname="china-s")
                pdf.save(path)
            for splitter in self.strategies:
                chunks = splitter(load_pdf(path), 10, 2)
                self.assertTrue(all(chunk.metadata["page_number"] == 1 for chunk in chunks))

    def test_config_defaults_and_overrides(self):
        """统一入口读取唯一配置，显式参数可覆盖实验值。"""
        documents = [Document(page_content="1234567890")]
        with patch("src.chunking.load_config", return_value={
            "chunking": {"strategy": "fixed", "chunk_size": 4, "chunk_overlap": 1},
        }):
            self.assertEqual([chunk.page_content for chunk in split_documents(documents)], ["1234", "4567", "7890"])
            chunks = split_documents(documents, strategy="semantic", chunk_size=8, chunk_overlap=0)
            self.assertEqual([chunk.page_content for chunk in chunks], ["12345678", "90"])
            self.assertTrue(all(chunk.metadata["chunk_strategy"] == "semantic" for chunk in chunks))

    def test_unknown_strategy(self):
        """未知策略明确报错，不悄悄回退。"""
        with self.assertRaisesRegex(ValueError, "不支持的分块策略"):
            split_documents([], "unknown", 8, 0)


class TestPDFLoader(unittest.TestCase):
    """验证实际 PDF 解析、页码和常见导入失败，不用 mock 替代解析器。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "论文.PDF"

    def write_pdf(self, texts):
        """生成可重复的中文/英文测试页面，空字符串对应空白页。"""
        with pymupdf.open() as pdf:
            for text in texts:
                page = pdf.new_page()
                if text:
                    page.insert_text((72, 72), text, fontname="china-s")
            pdf.save(self.path)

    def test_chinese_text_and_page_metadata(self):
        """每页正文与来源完整，页码同时提供索引和展示值。"""
        self.write_pdf(["中文论文摘要", "实验结果与结论"])
        documents = load_pdf(self.path)
        self.assertEqual(len(documents), 2)
        self.assertEqual([d.page_content for d in documents], ["中文论文摘要", "实验结果与结论"])
        for index, document in enumerate(documents):
            self.assertIsInstance(document, Document)
            self.assertEqual(document.metadata, {
                "source": str(self.path.resolve()),
                "source_file": "论文.PDF",
                "file_type": ".pdf",
                "doc_id": hashlib.sha256(self.path.read_bytes()).hexdigest(),
                "page": index,
                "page_number": index + 1,
                "total_pages": 2,
            })

    def test_relative_string_path(self):
        """字符串相对路径也能加载，来源统一为绝对路径。"""
        self.write_pdf(["摘要"])
        documents = load_pdf(os.path.relpath(self.path))
        self.assertEqual(documents[0].metadata["source"], str(self.path.resolve()))

    def test_document_id_is_stable_and_shared(self):
        """重复导入与文件副本共用内容标识，各页不会随机生成不同 ID。"""
        self.write_pdf(["摘要", "结论"])
        copy_path = self.path.with_name("副本.pdf")
        copy_path.write_bytes(self.path.read_bytes())
        documents = load_pdf(self.path) + load_pdf(self.path) + load_pdf(copy_path)
        self.assertEqual(len({d.metadata["doc_id"] for d in documents}), 1)

    def test_document_id_changes_with_file_content(self):
        """文件内容改变后标识也改变，避免误用旧索引。"""
        self.write_pdf(["初始内容"])
        original_id = load_pdf(self.path)[0].metadata["doc_id"]
        self.path = self.path.with_name("更新.pdf")
        self.write_pdf(["更新内容"])
        self.assertNotEqual(load_pdf(self.path)[0].metadata["doc_id"], original_id)

    def test_position_sorting(self):
        """内容流先写下方再写上方时，正文仍按页面位置排序。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 144), "BOTTOM")
            page.insert_text((72, 72), "TOP")
            pdf.save(self.path)
        text = load_pdf(self.path)[0].page_content
        self.assertLess(text.index("TOP"), text.index("BOTTOM"))

    def test_blank_page_does_not_shift_page_numbers(self):
        """跳过中间空白页仍保留原始物理页码，避免引用错页。"""
        self.write_pdf(["第一页", "", "第三页"])
        with self.assertLogs("src.data_loader.pdf_loader", level="WARNING") as logs:
            documents = load_pdf(self.path)
        self.assertEqual([d.metadata["page_number"] for d in documents], [1, 3])
        self.assertEqual([d.metadata["total_pages"] for d in documents], [3, 3])
        self.assertIn("第 2 页", logs.output[0])

    def test_image_only_pdf_is_not_silently_accepted(self):
        """将文字渲染为图片，验证没有文本层的扫描页不会伪装为导入成功。"""
        with pymupdf.open() as original:
            page = original.new_page()
            page.insert_text((72, 72), "Image-only paper")
            image = page.get_pixmap().tobytes("png")
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_image(page.rect, stream=image)
            pdf.save(self.path)
        with self.assertLogs("src.data_loader.pdf_loader", level="WARNING"):
            with self.assertRaisesRegex(ValueError, "未提取到文本"):
                load_pdf(self.path)

    def test_missing_file(self):
        """不存在的文件明确报错。"""
        with self.assertRaises(FileNotFoundError):
            load_pdf(self.path)

    def test_unsupported_extension(self):
        """其他加载器的格式不能误入 PDF 解析流程。"""
        path = self.path.with_suffix(".txt")
        path.write_text("摘要", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "不支持此文件格式"):
            load_pdf(path)

    def test_damaged_pdf(self):
        """损坏文件的解析错误向上传递，不返回空列表掩盖失败。"""
        self.path.write_bytes(b"not a valid PDF")
        with self.assertRaises(pymupdf.FileDataError):
            load_pdf(self.path)

    def test_password_protected_pdf(self):
        """需要打开密码的文件明确提示先解密。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 72), "Protected paper")
            pdf.save(self.path, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                     owner_pw="owner", user_pw="reader")
        with self.assertRaisesRegex(ValueError, "需要密码"):
            load_pdf(self.path)


class TestAcademicPDF(unittest.TestCase):
    """用实际绘制的论文版面验证阅读顺序、续表与公式，不 mock 解析器。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "academic.pdf"

    def draw_table(self, page, top, rows, caption="Table 1", three_lines=False):
        """绘制可重复的网格表或三线表，保留真正的文字和线段。"""
        if caption:
            page.insert_text((50, top - 15), caption)
        bottom = top + 30 * len(rows)
        columns = len(rows[0])
        if not three_lines:
            for index in range(columns + 1):
                x = 50 + 450 * index / columns
                page.draw_line((x, top), (x, bottom))
        ys = [top, top + 30, bottom] if three_lines else range(top, bottom + 1, 30)
        for y in ys:
            page.draw_line((50, y), (500, y))
        for row_index, row in enumerate(rows):
            for column, value in enumerate(row):
                page.insert_text((60 + 450 * column / columns, top + 20 + row_index * 30), value)

    def write_columns(self, page, top):
        """同一高度左右栏同时存在，朴素位置排序会交错读取。"""
        for index in range(3):
            for x, label in [(50, "LEFT"), (330, "RIGHT")]:
                page.insert_text((x, top + index * 25),
                                 f"{label}-{top}-{index} paragraph with many words", fontsize=9)

    def test_two_columns_with_full_width_heading(self):
        """先通栏标题，再完整左栏、完整右栏，最后页脚。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            page.insert_text((190, 60), "ACADEMIC PAPER TITLE", fontsize=16)
            self.write_columns(page, 100)
            page.insert_text((270, 760), "FOOTER TEXT")
            pdf.save(self.path)
        text = load_pdf(self.path)[0].page_content
        self.assertLess(text.index("TITLE"), text.index("LEFT-100-0"))
        self.assertLess(text.index("LEFT-100-2"), text.index("RIGHT-100-0"))
        self.assertLess(text.index("RIGHT-100-2"), text.index("FOOTER"))

    def test_full_width_separator_between_column_sections(self):
        """页中通栏内容将双栏分为上下区域，避免跨区域串读。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            self.write_columns(page, 100)
            page.insert_text((180, 220), "FULL WIDTH SECTION HEADING", fontsize=14)
            self.write_columns(page, 270)
            pdf.save(self.path)
        text = load_pdf(self.path)[0].page_content
        expected = ["LEFT-100-2", "RIGHT-100-0", "HEADING", "LEFT-270-0", "RIGHT-270-0"]
        self.assertEqual([text.index(token) for token in expected], sorted(text.index(token) for token in expected))

    def test_indented_single_column_is_not_reordered(self):
        """上下分离的缩进段落不能被误当作同时并排的两栏。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            for y, x, label in [(100,330,"FIRST"),(125,330,"SECOND"),(300,50,"THIRD"),(325,50,"FOURTH")]:
                page.insert_text((x,y), label + " paragraph with enough words", fontsize=9)
            pdf.save(self.path)
        text = load_pdf(self.path)[0].page_content
        self.assertLess(text.index("FIRST"), text.index("THIRD"))

    def test_cross_page_table_merges_rows_and_keeps_row_sources(self):
        """相邻页同一续表合并，重复表头去重，每行记录真实来源页。"""
        with pymupdf.open() as pdf:
            self.draw_table(pdf.new_page(width=600,height=800), 650, [["Method","Score"],["A","90"],["B","95"]])
            self.draw_table(pdf.new_page(width=600,height=800), 70, [["Method","Score"],["C","96"],["D","97"]], "Table 1 (continued)")
            pdf.save(self.path)
        tables = [d for d in load_pdf(self.path) if d.metadata.get("content_type") == "table"]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual(table.page_content.count("Method"), 1)
        self.assertIn("| C | 96 |", table.page_content)
        self.assertEqual((table.metadata["page_number"], table.metadata["page_end"]), (1,2))
        rows = json.loads(table.metadata["table_rows"])
        self.assertEqual([row["page_number"] for row in rows], [1,1,1,2,2])
        self.assertEqual(len(json.loads(table.metadata["table_regions"])), 2)

    def test_continuation_without_repeated_header(self):
        """明确同编号续表可无重复表头，下一页首行数据不可丢失。"""
        with pymupdf.open() as pdf:
            self.draw_table(pdf.new_page(width=600,height=800),650,[["Method","Score"],["A","90"],["B","91"]])
            self.draw_table(pdf.new_page(width=600,height=800),70,[["C","92"],["D","93"]],"Table 1 (continued)")
            pdf.save(self.path)
        tables = [d for d in load_pdf(self.path) if d.metadata.get("content_type") == "table"]
        self.assertEqual(len(tables), 1)
        self.assertIn("| C | 92 |", tables[0].page_content)
        self.assertIn("| D | 93 |", tables[0].page_content)

    def test_different_numbered_tables_do_not_merge(self):
        """几何位置和表头相同也不能合并不同编号的独立表。"""
        with pymupdf.open() as pdf:
            self.draw_table(pdf.new_page(width=600,height=800),650,[["Method","Score"],["A","90"]],"Table 1")
            self.draw_table(pdf.new_page(width=600,height=800),70,[["Method","Score"],["B","91"]],"Table 2")
            pdf.save(self.path)
        self.assertEqual(sum(d.metadata.get("content_type") == "table" for d in load_pdf(self.path)), 2)

    def test_nonadjacent_or_incompatible_tables_do_not_merge(self):
        """页码断开、表头或列结构不同、非页边续接均分别保留。"""
        for middle_page, top, rows in [(True,70,[["Method","Score"],["B","91"]]),
                                      (False,70,[["Data","Count"],["B","91"]]),
                                      (False,70,[["Method","Score","Time"],["B","91","2"]]),
                                      (False,350,[["Method","Score"],["B","91"]])]:
            with self.subTest(middle_page=middle_page,top=top,rows=rows):
                with pymupdf.open() as pdf:
                    self.draw_table(pdf.new_page(width=600,height=800),650,[["Method","Score"],["A","90"]])
                    if middle_page:
                        pdf.new_page(width=600,height=800).insert_text((50,100),"Intermediate page")
                    self.draw_table(pdf.new_page(width=600,height=800),top,rows,"Table 1")
                    pdf.save(self.path)
                self.assertEqual(sum(d.metadata.get("content_type") == "table" for d in load_pdf(self.path)), 2)

    def test_three_line_table_and_empty_cells(self):
        """三线表在规则围出的区域内识别，空单元格不从相邻行填值。"""
        with pymupdf.open() as pdf:
            self.draw_table(pdf.new_page(width=600,height=800),150,[["Method","Score"],["A","90"],["B",""],["C","95"]],three_lines=True)
            pdf.save(self.path)
        tables = [d for d in load_pdf(self.path) if d.metadata.get("content_type") == "table"]
        self.assertEqual(len(tables), 1)
        self.assertIn("| B |  |", tables[0].page_content)
        self.assertEqual(json.loads(tables[0].metadata["table_rows"])[2]["cells"], ["B", ""])

    def test_formula_symbols_and_scripts_are_preserved(self):
        """公式符号与上下标保持，不能把 x 的平方压成普通 x2。"""
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_text((60,80),"y = x",fontsize=14)
            page.insert_text((90,74),"2",fontsize=8)
            page.insert_text((100,80)," + a",fontsize=14)
            page.insert_text((124,85),"i",fontsize=8)
            page.insert_font(fontname="math",fontbuffer=pymupdf.Font("cjk").buffer)
            page.insert_text((60,120),"α + β = ∑ x / √n",fontname="math")
            pdf.save(self.path)
        document = load_pdf(self.path)[0]
        self.assertIn("x^{2}", document.page_content)
        self.assertIn("a_{i}", document.page_content)
        for symbol in ["α","β","∑","√"]:
            self.assertIn(symbol, document.page_content)
        layout = json.loads(document.metadata["formula_layout"])
        spans = [span for line in layout for span in line["spans"]]
        self.assertTrue(any(span["text"] == "2" and span["origin"][1] == 74 for span in spans))

    def test_multiple_table_styles_on_same_page(self):
        """同页网格表与三线表分别保留，不把横线范围或单元格混在一起。"""
        with pymupdf.open() as pdf:
            page=pdf.new_page(width=600,height=800)
            self.draw_table(page,150,[["Method","Score"],["GRID","90"]],"Table 1")
            self.draw_table(page,400,[["Method","Score"],["PLAIN","95"],["OTHER","96"]],"Table 2",three_lines=True)
            pdf.save(self.path)
        tables=[d for d in load_pdf(self.path) if d.metadata.get("content_type")=="table"]
        self.assertEqual(len(tables),2)
        self.assertIn("GRID",tables[0].page_content)
        self.assertNotIn("PLAIN",tables[0].page_content)
        self.assertIn("PLAIN",tables[1].page_content)

    def test_three_line_formula_cells_and_header_words(self):
        """多词表头和公式中的间隔不拆成虚假的列。"""
        with pymupdf.open() as pdf:
            page=pdf.new_page(width=600,height=800)
            self.draw_table(page,150,[["Layer Type","Complexity per Layer"],["Attention","O(n * d)"],["Recurrent","O(n * d * d)"]],three_lines=True)
            pdf.save(self.path)
        table=next(d for d in load_pdf(self.path) if d.metadata.get("content_type")=="table")
        self.assertEqual(len(json.loads(table.metadata["table_rows"])[0]["cells"]),2)
        self.assertIn("Complexity per Layer",table.page_content)
        self.assertIn("O(n * d * d)",table.page_content)

    def test_chinese_continuation_and_unnumbered_tables(self):
        """中文续表无重复表头时可合并；无编号的独立表不凭同表头合并。"""
        for numbered in [True,False]:
            with self.subTest(numbered=numbered):
                with pymupdf.open() as pdf:
                    first=pdf.new_page(width=600,height=800)
                    self.draw_table(first,650,[["Method","Score"],["A","90"]],"")
                    if numbered:
                        first.insert_text((50,635),"表 1",fontname="china-s")
                    second=pdf.new_page(width=600,height=800)
                    self.draw_table(second,70,[["B","91"],["C","92"]] if numbered else [["Method","Score"],["B","91"]],"")
                    if numbered:
                        second.insert_text((50,55),"续表 1",fontname="china-s")
                    pdf.save(self.path)
                tables=[d for d in load_pdf(self.path) if d.metadata.get("content_type")=="table"]
                self.assertEqual(len(tables),1 if numbered else 2)
                self.assertIn("B",tables[-1].page_content)

    def test_dense_formula_rows_do_not_leak_into_neighbors(self):
        """上标边缘与邻行相交时，不能向表头或邻行复制残留字符。"""
        with pymupdf.open() as pdf:
            page=pdf.new_page(width=600,height=800)
            page.insert_text((50,85),"Table 1")
            for y in [100,120,158]:page.draw_line((50,y),(500,y))
            page.insert_text((60,115),"Method",fontsize=11)
            page.insert_text((300,115),"Complexity",fontsize=11)
            for y,label in [(130,"Attention"),(142,"Recurrent")]:
                page.insert_text((60,y),label,fontsize=11)
                page.insert_text((300,y),"x",fontsize=11)
                page.insert_text((306,y-4),"2",fontsize=7)
            pdf.save(self.path)
        table=next(d for d in load_pdf(self.path) if d.metadata.get("content_type")=="table")
        rows=json.loads(table.metadata["table_rows"])
        self.assertEqual([row["cells"] for row in rows],[["Method","Complexity"],["Attention","x^{2}"],["Recurrent","x^{2}"]])

    def test_ambiguous_three_line_header_keeps_original_text(self):
        """父表头横跨两个子列时不输出错误结构，正文和原文位置仍保留。"""
        with pymupdf.open() as pdf:
            page=pdf.new_page(width=600,height=800)
            page.insert_text((50,85),"Table 1")
            for y in [100,145,220]:page.draw_line((50,y),(500,y))
            page.insert_text((60,120),"Method")
            page.insert_text((285,115),"Combined long parent heading")
            page.insert_text((285,138),"LEFT")
            page.insert_text((400,138),"RIGHT")
            for y,label in [(165,"A"),(190,"B")]:
                page.insert_text((60,y),label)
                page.insert_text((285,y),"90")
                page.insert_text((400,y),"95")
            pdf.save(self.path)
        with self.assertLogs("src.data_loader.pdf_loader",level="WARNING"):
            documents=load_pdf(self.path)
        self.assertFalse(any(d.metadata.get("content_type")=="table" for d in documents))
        self.assertIn("LEFT",documents[0].page_content)
        self.assertIn("RIGHT",documents[0].page_content)

    def test_three_page_table_keeps_first_caption_identity(self):
        """中间续页不重复标题时，仍保留整张表的编号和第三页数据。"""
        with pymupdf.open() as pdf:
            self.draw_table(pdf.new_page(width=600,height=800),650,[["Method","Score"],["A","90"],["B","91"]])
            self.draw_table(pdf.new_page(width=600,height=800),70,[["Method","Score"]]+[[f"M{i}",str(i)] for i in range(22)],"")
            self.draw_table(pdf.new_page(width=600,height=800),70,[["Method","Score"],["FINAL","99"]],"")
            pdf.save(self.path)
        tables=[d for d in load_pdf(self.path) if d.metadata.get("content_type")=="table"]
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0].metadata["page_end"],3)
        self.assertIn("FINAL",tables[0].page_content)

    def test_fraction_layout_retains_original_positions(self):
        """不猜测二维分式的 LaTeX，保留分子/分母位置以回看原文。"""
        with pymupdf.open() as pdf:
            page=pdf.new_page()
            page.insert_text((60,130),"f =",fontsize=12)
            page.insert_text((95,119),"1",fontsize=10)
            page.draw_line((92,125),(119,125))
            page.insert_text((95,140),"1 + x",fontsize=10)
            pdf.save(self.path)
        document = load_pdf(self.path)[0]
        spans = [s for line in json.loads(document.metadata["formula_layout"]) for s in line["spans"]]
        self.assertTrue(any(s["text"] == "1" and s["origin"][1] == 119 for s in spans))
        self.assertTrue(any("1 + x" in s["text"] and s["origin"][1] == 140 for s in spans))

    def test_image_formula_is_locatable_after_batch_save(self):
        """混合文本页的公式图片保留原文定位，批量保存后仍可裁剪查看。"""
        with pymupdf.open() as picture:
            image_page=picture.new_page(width=180,height=50)
            image_page.insert_text((10,30),"y = (a+b)/c",fontsize=16)
            image=image_page.get_pixmap().tobytes("png")
        with pymupdf.open() as pdf:
            page=pdf.new_page()
            page.insert_text((60,80),"Image equation follows:")
            page.insert_image(pymupdf.Rect(60,100,240,150),stream=image)
            pdf.save(self.path)
        tasks=create_import_tasks([("formula.pdf",self.path.read_bytes())])
        list(batch_import(tasks, Path(self.directory.name)/"raw"))
        self.assertEqual(tasks[0]["status"],"success")
        document=tasks[0]["documents"][0]
        region=json.loads(document.metadata["image_regions"])[0]
        self.assertIn("图像区域",document.page_content)
        with pymupdf.open(document.metadata["source"]) as saved:
            crop=saved[region["page_number"]-1].get_pixmap(clip=region["bbox"])
            self.assertEqual((crop.width,crop.height),(180,50))
            self.assertGreater(len(crop.tobytes("png")),100)


class TestDOCXLoader(unittest.TestCase):
    """通过实际 DOCX 验证正文、表格顺序和位置，不修改原始课程资料。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "论文.DOCX"

    def test_paragraph_table_order_and_metadata(self):
        """表格保留在正文原位置，中文段落与单元格内容均可读取。"""
        word = WordDocument()
        word.add_paragraph("摘要")
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "方法"
        table.cell(0, 1).text = "准确率"
        table.cell(1, 0).text = "Transformer"
        table.cell(1, 1).text = "90%"
        word.add_paragraph("结论")
        word.save(self.path)

        documents = load_docx(os.path.relpath(self.path))
        self.assertEqual([d.page_content for d in documents], [
            "摘要", "方法\t准确率\nTransformer\t90%", "结论",
        ])
        self.assertEqual([d.metadata["block_index"] for d in documents], [1, 2, 3])
        self.assertEqual([d.metadata["block_type"] for d in documents], [
            "paragraph", "table", "paragraph",
        ])
        self.assertEqual(documents[0].metadata["paragraph_index"], 1)
        self.assertEqual(documents[1].metadata["table_index"], 1)
        self.assertEqual(documents[2].metadata["paragraph_index"], 2)
        for document in documents:
            self.assertIsInstance(document, Document)
            self.assertEqual(document.metadata["source"], str(self.path.resolve()))
            self.assertEqual(document.metadata["source_file"], "论文.DOCX")
            self.assertEqual(document.metadata["file_type"], ".docx")
            self.assertNotIn("page", document.metadata)
            self.assertNotIn("page_number", document.metadata)

    def test_blank_blocks_do_not_shift_positions(self):
        """空段落和空表格跳过，但后续段落/表格编号不偏移。"""
        word = WordDocument()
        word.add_paragraph("")
        word.add_table(rows=1, cols=1)
        word.add_paragraph("正文")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "实验数据"
        word.save(self.path)
        documents = load_docx(self.path)
        self.assertEqual([d.metadata["block_index"] for d in documents], [3, 4])
        self.assertEqual(documents[0].metadata["paragraph_index"], 2)
        self.assertEqual(documents[1].metadata["table_index"], 2)

    def test_table_empty_cells_and_multiple_paragraphs(self):
        """空单元格保留分隔符，单元格内的多个段落不丢失。"""
        word = WordDocument()
        table = word.add_table(rows=1, cols=3)
        table.cell(0, 1).text = "第一段"
        table.cell(0, 1).add_paragraph("第二段")
        word.save(self.path)
        self.assertEqual(load_docx(self.path)[0].page_content, "\t第一段\n第二段\t")

    def test_document_id_is_stable_shared_and_changes(self):
        """同文件各块共用内容 ID，重复读取稳定，内容变化后更新。"""
        word = WordDocument()
        word.add_paragraph("摘要")
        word.add_paragraph("结论")
        word.save(self.path)
        documents = load_docx(self.path) + load_docx(self.path)
        ids = {d.metadata["doc_id"] for d in documents}
        self.assertEqual(ids, {hashlib.sha256(self.path.read_bytes()).hexdigest()})
        word.add_paragraph("新增实验")
        word.save(self.path)
        self.assertNotIn(load_docx(self.path)[0].metadata["doc_id"], ids)

    def test_empty_document(self):
        """仅有空段落和空表格时明确报错。"""
        word = WordDocument()
        word.add_paragraph(" \t ")
        word.add_table(rows=1, cols=2)
        word.save(self.path)
        with self.assertRaisesRegex(ValueError, "未提取到文本"):
            load_docx(self.path)

    def test_missing_file(self):
        """缺失文件不会被当成空文档。"""
        with self.assertRaises(FileNotFoundError):
            load_docx(self.path)

    def test_legacy_doc_extension(self):
        """旧版 .doc 格式需转换后导入。"""
        path = self.path.with_suffix(".doc")
        path.write_bytes(b"legacy document")
        with self.assertRaisesRegex(ValueError, "仅支持 .docx"):
            load_docx(path)

    def test_damaged_document(self):
        """损坏文件的解析错误向上传递。"""
        self.path.write_bytes(b"not a DOCX archive")
        with self.assertRaises(PackageNotFoundError):
            load_docx(self.path)


class TestTextLoader(unittest.TestCase):
    """验证纯文本保真、中文编码和行范围。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "文献.txt"

    def test_utf8_text_and_line_metadata(self):
        """中文正文保留首尾空白与空行，来源和行号完整。"""
        text = "\n  中文正文\n\n第二段\n"
        self.path.write_bytes(text.encode("utf-8"))
        documents = load_text(os.path.relpath(self.path))
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].page_content, text)
        self.assertEqual(documents[0].metadata, {
            "source": str(self.path.resolve()),
            "source_file": "文献.txt",
            "file_type": ".txt",
            "doc_id": hashlib.sha256(self.path.read_bytes()).hexdigest(),
            "line_start": 1,
            "line_end": 4,
        })

    def test_utf8_bom_and_crlf(self):
        """移除 BOM，但保留 Windows 的 CRLF 换行。"""
        text = "第一行\r\n第二行\r\n"
        self.path.write_bytes(text.encode("utf-8-sig"))
        document = load_text(self.path)[0]
        self.assertEqual(document.page_content, text)
        self.assertEqual(document.metadata["line_end"], 2)

    def test_markdown_preserves_syntax_and_indentation(self):
        """Markdown 标题、链接、表格和代码块不被转换或清理。"""
        text = "# 标题\n\n[论文](https://example.com)\n|方法|结果|\n|---|---|\n```python\n    print('中文')\n```\n"
        for suffix in [".MD", ".markdown"]:
            with self.subTest(suffix=suffix):
                path = self.path.with_suffix(suffix)
                path.write_bytes(text.encode("utf-8"))
                document = load_text(path)[0]
                self.assertEqual(document.page_content, text)
                self.assertEqual(document.metadata["file_type"], suffix.lower())
                self.assertNotIn("page_number", document.metadata)

    def test_document_id_is_stable_and_changes(self):
        """内容指纹由原始字节生成，重复读取稳定，文本变化后更新。"""
        self.path.write_text("原始正文", encoding="utf-8")
        original_id = load_text(self.path)[0].metadata["doc_id"]
        self.assertEqual(load_text(self.path)[0].metadata["doc_id"], original_id)
        self.path.write_text("更新正文", encoding="utf-8")
        self.assertNotEqual(load_text(self.path)[0].metadata["doc_id"], original_id)

    def test_empty_or_whitespace_only_file(self):
        """空字节、仅 BOM 和纯空白都不视为有效内容。"""
        for data in [b"", b"\xef\xbb\xbf", b" \t\r\n"]:
            with self.subTest(data=data):
                self.path.write_bytes(data)
                with self.assertRaisesRegex(ValueError, "没有有效内容"):
                    load_text(self.path)

    def test_invalid_utf8(self):
        """不静默替换错误字符，也不猜测 GBK 等其他编码。"""
        self.path.write_bytes("中文".encode("gbk"))
        with self.assertRaises(UnicodeDecodeError):
            load_text(self.path)

    def test_missing_file(self):
        """缺失文件明确报错。"""
        with self.assertRaises(FileNotFoundError):
            load_text(self.path)

    def test_unsupported_extension(self):
        """二进制格式不能被当作纯文本导入。"""
        path = self.path.with_suffix(".pdf")
        path.write_bytes(b"PDF")
        with self.assertRaisesRegex(ValueError, "不支持此文件格式"):
            load_text(path)


class TestBatchImport(unittest.TestCase):
    """验证批量任务的独立状态、真实保存和有界重试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.raw_dir = Path(self.directory.name) / "raw"

    def test_all_formats_saved_with_correct_sources(self):
        """四类文档均实际解析和保存，来源指向持久文件而非临时文件。"""
        word = WordDocument()
        word.add_paragraph("Word 摘要")
        buffer = io.BytesIO()
        word.save(buffer)
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "PDF abstract")
            pdf_data = pdf.tobytes()
        files = [("paper.pdf", pdf_data), ("paper.docx", buffer.getvalue()),
                 ("说明.txt", "文本摘要".encode("utf-8")), ("说明.md", b"# Abstract")]
        tasks = create_import_tasks(files)
        progress = list(batch_import(tasks, self.raw_dir))
        self.assertEqual(progress[-1], {"completed": 4, "total": 4})
        self.assertEqual([t["status"] for t in tasks], ["success"] * 4)
        for task in tasks:
            path = Path(task["path"])
            self.assertEqual(path.read_bytes(), task["data"])
            self.assertEqual(task["attempts"], 1)
            for document in task["documents"]:
                self.assertEqual(document.metadata["source"], str(path))
                self.assertEqual(document.metadata["source_file"], task["name"])

    def test_failure_does_not_stop_later_files_or_leave_failed_file(self):
        """中间文件损坏也继续导入后续文件，完成进度包含失败项。"""
        tasks = create_import_tasks([("first.txt", b"First"), ("bad.pdf", b"broken"),
                                     ("last.md", b"# Last")])
        events = batch_import(tasks, self.raw_dir)
        next(events)
        self.assertEqual(tasks[0]["status"], "pending")
        next(events)
        self.assertEqual(tasks[0]["status"], "loading")
        progress = list(events)
        self.assertEqual([t["status"] for t in tasks], ["success", "failed", "success"])
        self.assertIn("FileDataError", tasks[1]["error"])
        self.assertEqual(tasks[1]["documents"], [])
        self.assertEqual(tasks[1]["path"], "")
        self.assertEqual(progress[-1], {"completed": 3, "total": 3})
        self.assertEqual(len(list(self.raw_dir.rglob("*.*"))), 2)

    def test_retry_failed_recovers_without_reloading_success(self):
        """模拟一次临时错误，手动重试仅失败项，成功项尝试次数不增加。"""
        tasks = create_import_tasks([("good.md", b"Good"), ("retry.txt", b"Retry")])

        def fail_once(path):
            if path.name == "retry.txt":
                raise OSError("临时读取失败")
            return load_document(path)

        with patch("src.data_loader.load_document", side_effect=fail_once):
            list(batch_import(tasks, self.raw_dir))
        self.assertEqual([t["status"] for t in tasks], ["success", "failed"])
        progress = list(batch_import(tasks, self.raw_dir, retry_failed=True))
        self.assertEqual(progress[-1], {"completed": 1, "total": 1})
        self.assertEqual([t["status"] for t in tasks], ["success", "success"])
        self.assertEqual([t["attempts"] for t in tasks], [1, 2])
        self.assertEqual(tasks[1]["error"], "")

    def test_persistent_failure_is_bounded(self):
        """每次重试只处理一轮，永久错误不会产生无限循环。"""
        tasks = create_import_tasks([("bad.txt", b"\xff")])
        list(batch_import(tasks, self.raw_dir))
        list(batch_import(tasks, self.raw_dir, retry_failed=True))
        self.assertEqual(tasks[0]["attempts"], 2)
        self.assertEqual(tasks[0]["status"], "failed")
        self.assertIn("UnicodeDecodeError", tasks[0]["error"])

    def test_repeated_run_skips_completed_tasks(self):
        """成功后再次调用不会重复加载或重复保存。"""
        tasks = create_import_tasks([("ok.txt", b"Ready")])
        list(batch_import(tasks, self.raw_dir))
        with patch("src.data_loader.load_document") as loader:
            self.assertEqual(list(batch_import(tasks, self.raw_dir)), [{"completed": 0, "total": 0}])
            list(batch_import(tasks, self.raw_dir, retry_failed=True))
            loader.assert_not_called()
        self.assertEqual(tasks[0]["attempts"], 1)

    def test_interrupted_loading_task_can_resume(self):
        """页面中断后遗留的处理中任务可继续处理。"""
        tasks = create_import_tasks([("ok.txt", b"Ready")])
        events = batch_import(tasks, self.raw_dir)
        next(events)
        next(events)
        events.close()
        self.assertEqual(tasks[0]["status"], "loading")
        list(batch_import(tasks, self.raw_dir))
        self.assertEqual(tasks[0]["status"], "success")

    def test_empty_batch_and_duplicate_inputs(self):
        """空批次不创建目录，相同文件名和内容的重复项合并。"""
        self.assertEqual(list(batch_import([], self.raw_dir)), [{"completed": 0, "total": 0}])
        self.assertFalse(self.raw_dir.exists())
        tasks = create_import_tasks([("same.txt", b"Same"), ("same.txt", b"Same")])
        self.assertEqual(len(tasks), 1)

    def test_same_name_different_content_is_preserved(self):
        """不同内容的同名文献独立保存，不覆盖旧内容。"""
        tasks = create_import_tasks([("same.txt", b"First"), ("same.txt", b"Second")])
        list(batch_import(tasks, self.raw_dir))
        self.assertNotEqual(tasks[0]["path"], tasks[1]["path"])
        self.assertEqual([Path(t["path"]).read_bytes() for t in tasks], [b"First", b"Second"])

    def test_unsupported_format_and_unsafe_names(self):
        """路径名和不支持格式记录失败，不写到导入目录之外。"""
        names = ["../escape.txt", "/tmp/escape.txt", "folder\\escape.txt", "paper.exe"]
        tasks = create_import_tasks([(name, b"Text") for name in names])
        list(batch_import(tasks, self.raw_dir))
        self.assertTrue(all(t["status"] == "failed" for t in tasks))
        self.assertFalse(self.raw_dir.exists())

    def test_size_limit(self):
        """前端之外的调用也核验单份文档大小。"""
        tasks = create_import_tasks([("large.txt", b"A" * (1024 * 1024 + 1))])
        list(batch_import(tasks, self.raw_dir, max_file_size_mb=1))
        self.assertEqual(tasks[0]["status"], "failed")
        self.assertIn("文件过大", tasks[0]["error"])

    def test_existing_file_with_different_bytes_is_not_overwritten(self):
        """已有保存位置被外部修改时明确失败，保留其内容。"""
        tasks = create_import_tasks([("paper.txt", b"Original")])
        list(batch_import(tasks, self.raw_dir))
        path = Path(tasks[0]["path"])
        path.write_bytes(b"Changed externally")
        tasks = create_import_tasks([("paper.txt", b"Original")])
        list(batch_import(tasks, self.raw_dir))
        self.assertEqual(tasks[0]["status"], "failed")
        self.assertEqual(path.read_bytes(), b"Changed externally")

    def test_storage_failure_is_retryable(self):
        """保存位置暂时不可用时记录错误，恢复后可重试成功。"""
        self.raw_dir.write_bytes(b"not a directory")
        tasks = create_import_tasks([("ok.txt", b"Text")])
        list(batch_import(tasks, self.raw_dir))
        self.assertEqual(tasks[0]["status"], "failed")
        self.raw_dir.unlink()
        list(batch_import(tasks, self.raw_dir, retry_failed=True))
        self.assertEqual(tasks[0]["status"], "success")


class TestBatchIndex(unittest.TestCase):
    """真实加载器/分块/Chroma 联调；二维向量用于核对是否重复编码。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.raw_dir = Path(self.directory.name) / "raw"
        self.embeddings = SmallEmbeddings()
        self.store = VectorStore(Path(self.directory.name) / "index", self.embeddings)

    def run_batch(self, tasks, **kwargs):
        """复用实际入库流程与真实临时数据库。"""
        return list(batch_build_index(tasks, self.raw_dir, vector_store=self.store, **kwargs))

    def test_size_boundary_rejects_before_parse_and_keeps_next_valid_document(self):
        """真实PDF刚好1MiB可入库；超1字节先拒绝，重试不污染已有向量。"""
        limit = 1024 * 1024
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Size boundary fixture")
            raw = pdf.tobytes()
        exact = raw + b" " * (limit - len(raw))
        tasks = create_import_tasks([("too_large.pdf", exact + b" "), ("accepted.pdf", exact)])
        with patch("src.data_loader.load_document", wraps=load_document) as loader:
            progress = self.run_batch(tasks, max_file_size_mb=1)
        self.assertEqual([call.args[0].name for call in loader.call_args_list], ["accepted.pdf"])
        self.assertEqual(progress[-1], {"completed": 2, "total": 2})
        self.assertEqual([task["status"] for task in tasks], ["failed", "success"])
        self.assertEqual(tasks[0]["path"], "")
        self.assertEqual(tasks[0]["documents"], [])
        self.assertFalse(tasks[0]["indexed"])
        self.assertEqual(Path(tasks[1]["path"]).stat().st_size, limit)
        before = self.store.list_chunks()
        calls = deepcopy(self.embeddings.document_calls)
        with patch("src.data_loader.load_document", side_effect=AssertionError("超限文件不应进入解析")):
            self.run_batch(tasks, max_file_size_mb=1, retry_failed=True)
        self.assertEqual(tasks[0]["attempts"], 2)
        self.assertIn("文件过大", tasks[0]["error"])
        self.assertEqual(self.store.list_chunks(), before)
        self.assertEqual(self.embeddings.document_calls, calls)

    def test_all_formats_complete_index_and_keep_sources(self):
        """四类真实文件完成整个流程，只有写入索引后才标记成功。"""
        word = WordDocument()
        word.add_paragraph("Word 神经网络摘要")
        buffer = io.BytesIO()
        word.save(buffer)
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((72, 72), "PDF neural network")
            pdf_data = pdf.tobytes()
        tasks = create_import_tasks([("论文.pdf", pdf_data), ("论文.docx", buffer.getvalue()),
                                     ("论文.txt", "农业论文正文".encode()), ("论文.md", b"# AI paper")])
        progress = self.run_batch(tasks)
        self.assertEqual(progress[-1], {"completed": 4, "total": 4})
        self.assertTrue(all(task["status"] == "success" and task["indexed"] for task in tasks))
        self.assertEqual(self.store.count(), sum(task["chunk_count"] for task in tasks))
        for task in tasks:
            self.assertEqual(task["attempts"], 1)
            self.assertEqual(task["processed_chunks"], task["chunk_count"])
            self.assertEqual(task["added_chunks"], task["chunk_count"])
            for document in self.store.list_chunks(task["documents"][0].metadata["doc_id"]):
                self.assertEqual(document.metadata["source"], task["path"])
                self.assertEqual(document.metadata["source_file"], task["name"])

    def test_new_document_keeps_old_index_and_only_encodes_new(self):
        """新增一篇文献保留旧块，重新构造客户端不重算旧向量。"""
        first = create_import_tasks([("first.md", b"# Neural network")])
        self.run_batch(first)
        old_chunks = self.store.list_chunks()
        other_embeddings = SmallEmbeddings()
        reopened = VectorStore(Path(self.directory.name) / "index", other_embeddings)
        second = create_import_tasks([("second.txt", "农业论文".encode())])
        list(batch_build_index(second, self.raw_dir, vector_store=reopened))
        self.assertEqual(other_embeddings.document_calls, [["农业论文"]])
        self.assertEqual(reopened.count(), 2)
        self.assertEqual(reopened.list_chunks(old_chunks[0].metadata["doc_id"]), old_chunks)

    def test_repeated_and_recreated_tasks_do_not_encode_old_chunks(self):
        """页面重跑跳过完成任务，重新上传相同原文也不重新编码。"""
        files = [("paper.txt", b"Neural network")]
        tasks = create_import_tasks(files)
        self.run_batch(tasks)
        with patch("src.data_loader.load_document") as loader:
            self.assertEqual(self.run_batch(tasks), [{"completed": 0, "total": 0}])
            loader.assert_not_called()
        again = create_import_tasks(files)
        self.run_batch(again)
        self.assertTrue(again[0]["indexed"])
        self.assertEqual(again[0]["added_chunks"], 0)
        self.assertEqual(self.store.count(), 1)
        self.assertEqual(self.embeddings.document_calls, [["Neural network"]])

    def test_index_error_keeps_loaded_file_and_retry_skips_loading(self):
        """一个索引错误不阻断下一文档，重试保留加载结果并跳过成功任务。"""
        tasks = create_import_tasks([("retry.txt", b"First"), ("good.txt", b"Second")])
        real_add = self.store.add_chunks

        def fail_first(chunks):
            if chunks[0].metadata["source_file"] == "retry.txt":
                raise OSError("索引暂时不可写")
            return real_add(chunks)

        with patch.object(self.store, "add_chunks", side_effect=fail_first):
            self.run_batch(tasks)
        self.assertEqual([task["status"] for task in tasks], ["failed", "success"])
        self.assertTrue(Path(tasks[0]["path"]).is_file())
        self.assertTrue(tasks[0]["documents"])
        with patch("src.data_loader.load_document", side_effect=AssertionError("不应重载")):
            progress = self.run_batch(tasks, retry_failed=True)
        self.assertEqual(progress[-1], {"completed": 1, "total": 1})
        self.assertEqual([task["attempts"] for task in tasks], [2, 1])
        self.assertEqual(tasks[0]["error"], "")
        self.assertTrue(all(task["indexed"] for task in tasks))
        self.assertEqual(self.store.count(), 2)

    def test_partial_batch_failure_only_retries_missing_chunks(self):
        """第 501 块故障后保留前 500 块，下一次只编码剩余块。"""
        tasks = create_import_tasks([("large.txt", ("神经网络。\n\n" * 200).encode())])
        from src.chunking.fixed_chunk import split_fixed

        # 小窗口获得 501+ 个真实原文块，便于验证批次边界而不加载真实大模型。
        with patch("src.chunking.split_documents", side_effect=lambda docs: split_fixed(docs, 2, 0)):
            real_add = self.store.add_chunks
            calls = 0

            def fail_second(chunks):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise OSError("第二批失败")
                return real_add(chunks)

            with patch.object(self.store, "add_chunks", side_effect=fail_second):
                self.run_batch(tasks)
            self.assertEqual(tasks[0]["status"], "failed")
            self.assertFalse(tasks[0]["indexed"])
            self.assertEqual(tasks[0]["processed_chunks"], 500)
            self.assertEqual(self.store.count(), 500)
            self.run_batch(tasks, retry_failed=True)
        count = tasks[0]["chunk_count"]
        self.assertGreater(count, 500)
        self.assertEqual(tasks[0]["status"], "success")
        self.assertEqual(self.store.count(), count)
        self.assertEqual(tasks[0]["added_chunks"], count - 500)
        self.assertEqual(sum(len(call) for call in self.embeddings.document_calls), count)

    def test_stage_progress_and_interrupted_index_resume(self):
        """观察分块/索引阶段；进度中断后可恢复，不误标为成功。"""
        tasks = create_import_tasks([("paper.txt", b"Neural network")])
        events = batch_build_index(tasks, self.raw_dir, vector_store=self.store)
        statuses = []
        for _ in events:
            statuses.append(tasks[0]["status"])
            if tasks[0]["status"] == "indexing":
                break
        events.close()
        self.assertIn("loading", statuses)
        self.assertIn("chunking", statuses)
        self.assertFalse(tasks[0]["indexed"])
        self.run_batch(tasks)
        self.assertTrue(tasks[0]["indexed"])
        self.assertEqual(self.store.count(), 1)

    def test_loading_failures_and_empty_batch_do_not_initialize_model(self):
        """空批次/解码失败无需模型，也不创建索引；加载错误仍可有界重试。"""
        with patch("src.retrieval.vector_store.VectorStore") as factory:
            self.assertEqual(list(batch_build_index([], self.raw_dir)), [{"completed": 0, "total": 0}])
            tasks = create_import_tasks([("bad.txt", b"\xff")])
            list(batch_build_index(tasks, self.raw_dir))
            list(batch_build_index(tasks, self.raw_dir, retry_failed=True))
            factory.assert_not_called()
        self.assertEqual(tasks[0]["attempts"], 2)
        self.assertIn("UnicodeDecodeError", tasks[0]["error"])

    def test_loaded_documents_can_be_indexed_without_reload(self):
        """原始加载接口的成功任务也可继续进入索引阶段。"""
        tasks = create_import_tasks([("paper.txt", b"Neural network")])
        list(batch_import(tasks, self.raw_dir))
        with patch("src.data_loader.load_document", side_effect=AssertionError("不应重载")):
            self.run_batch(tasks)
        self.assertTrue(tasks[0]["indexed"])
        self.assertEqual(tasks[0]["attempts"], 2)

    def test_empty_content_after_loading_is_not_index_success(self):
        """加载后没有有效分块不能算索引成功；保留原文供用户核对。"""
        tasks = create_import_tasks([("paper.txt", b"Neural network")])
        with patch("src.chunking.split_documents", return_value=[]):
            self.run_batch(tasks)
        self.assertEqual(tasks[0]["status"], "failed")
        self.assertFalse(tasks[0]["indexed"])
        self.assertTrue(Path(tasks[0]["path"]).is_file())
        self.assertIn("没有可索引", tasks[0]["error"])
        self.assertEqual(self.embeddings.document_calls, [])


class TestImportFrontend(unittest.TestCase):
    """真实上传组件与 Chroma，隔离原文/索引；小型向量隔离大模型。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        from src.utils.config import load_config
        config = load_config()
        config["paths"]["raw_documents"] = str(Path(self.directory.name) / "raw")
        config["paths"]["session_db"] = str(Path(self.directory.name) / "memory.sqlite3")
        config["paths"]["vector_index"] = str(Path(self.directory.name) / "index")
        config["paths"]["logs"] = str(Path(self.directory.name) / "logs")
        self.embeddings = SmallEmbeddings()
        for target, value in (("src.utils.config.load_config", config),
                              ("src.utils.logger.load_config", config),
                              ("src.retrieval.vector_store.load_config", config),
                              ("src.retrieval.hybrid_retriever.load_config", config),
                              ("src.retrieval.reranker.load_config", config),
                              ("src.retrieval.vector_store.get_embeddings", self.embeddings)):
            patcher = patch(target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)
        from streamlit.testing.v1 import AppTest
        app_path = Path(__file__).resolve().parents[1] / "src/frontend/app.py"
        # 新环境首次加载界面依赖较慢，避免默认 3 秒等待导致误报。
        self.app = AppTest.from_file(str(app_path), default_timeout=10).run()

    def test_batch_upload_progress_state_and_rerun(self):
        """实际操作上传/按钮，核验混合结果、完整进度和重跑不重复执行。"""
        app = self.app
        self.assertTrue(app.button(key="start_import").disabled)
        # 开发过程中页面可能保留上一阶段“只加载成功”的任务，不得显示索引成功。
        loaded = create_import_tasks([("good.txt", "中文正文".encode("utf-8"))])
        list(batch_import(loaded, Path(self.directory.name) / "raw"))
        app.session_state["import_tasks"] = loaded
        app.file_uploader[0].set_value([("good.txt", "中文正文".encode("utf-8"), "text/plain")]).run()
        self.assertEqual(list(app.sidebar.dataframe[0].value["状态"]), ["待索引"])
        self.assertFalse(app.button(key="start_import").disabled)
        app.file_uploader[0].set_value([
            ("good.txt", "中文正文".encode("utf-8"), "text/plain"),
            ("bad.txt", b"\xff", "text/plain"),
        ]).run()
        app.button(key="start_import").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(list(app.sidebar.dataframe[0].value["状态"]), ["成功", "失败"])
        self.assertEqual(list(app.sidebar.dataframe[0].value["本次新增块"]), [1, 0])
        self.assertTrue(app.session_state["import_tasks"][0]["indexed"])
        self.assertEqual(app.session_state["import_progress"], {"completed": 2, "total": 2})
        self.assertEqual(app.get("progress")[0].proto.value, 100)
        self.assertTrue(app.button(key="start_import").disabled)
        self.assertFalse(app.button(key="retry_import").disabled)
        app.button(key="retry_import").click().run()
        self.assertEqual([t["attempts"] for t in app.session_state["import_tasks"]], [1, 2])
        app.run()
        self.assertEqual([t["attempts"] for t in app.session_state["import_tasks"]], [1, 2])

    def test_oversized_upload_shows_failure_and_valid_file_still_indexes(self):
        """AppTest绕过浏览器大小限制，验证20MiB后端保护与页面失败/重试状态。"""
        app = self.app
        app.file_uploader[0].set_value([("oversized.pdf", b"x" * (20 * 1024 * 1024 + 1), "application/pdf"),
                                       ("valid.txt", b"Neural network", "text/plain")]).run()
        with patch("src.data_loader.load_document", wraps=load_document) as loader:
            app.button(key="start_import").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([call.args[0].name for call in loader.call_args_list], ["valid.txt"])
        tasks = app.session_state["import_tasks"]
        self.assertEqual([task["status"] for task in tasks], ["failed", "success"])
        self.assertIn("文件过大", tasks[0]["error"])
        self.assertEqual(tasks[0]["path"], "")
        self.assertFalse(tasks[0]["indexed"])
        self.assertTrue(tasks[1]["indexed"])
        self.assertEqual(list(app.sidebar.dataframe[0].value["状态"]), ["失败", "成功"])
        self.assertFalse(app.button(key="retry_import").disabled)
        app.button(key="retry_import").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([task["attempts"] for task in app.session_state["import_tasks"]], [2, 1])
        self.assertEqual(self.embeddings.document_calls, [["Neural network"]])

    def test_failed_upload_recovers_and_disables_retry(self):
        """页面通过失败按钮恢复成功，之后重试按钮禁用且文献存在。"""
        app = self.app
        app.file_uploader[0].set_value([("retry.txt", b"Retry", "text/plain")]).run()
        with patch("src.data_loader.load_document", side_effect=OSError("暂时失败")):
            app.button(key="start_import").click().run()
        self.assertEqual(app.session_state["import_tasks"][0]["status"], "failed")
        app.button(key="retry_import").click().run()
        self.assertFalse(app.exception)
        task = app.session_state["import_tasks"][0]
        self.assertEqual(task["status"], "success")
        self.assertEqual(task["attempts"], 2)
        self.assertEqual(task["error"], "")
        self.assertTrue(Path(task["path"]).is_file())
        self.assertTrue(app.button(key="retry_import").disabled)

    def test_new_upload_is_incremental_and_repeat_upload_skips_encoding(self):
        """改变文件选择新增文献，重新选原文也不重算已有向量。"""
        app = self.app
        for filename, content in (("first.txt", b"First"), ("second.txt", b"Second"), ("first.txt", b"First")):
            app.file_uploader[0].set_value([(filename, content, "text/plain")]).run()
            app.button(key="start_import").click().run()
            self.assertFalse(app.exception)
        task = app.session_state["import_tasks"][0]
        self.assertEqual(task["added_chunks"], 0)
        self.assertEqual(task["index_total"], 2)
        self.assertEqual(self.embeddings.document_calls, [["First"], ["Second"]])

    def test_model_error_keeps_file_and_retry_completes_index(self):
        """缺失本地模型时不误报成功，恢复模型后仅重试索引。"""
        app = self.app
        app.file_uploader[0].set_value([("retry.txt", b"Retry", "text/plain")]).run()
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=FileNotFoundError("本地模型不存在")):
            app.button(key="start_import").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["import_tasks"][0]["status"], "failed")
        self.assertTrue(Path(app.session_state["import_tasks"][0]["path"]).is_file())
        with patch("src.data_loader.load_document", side_effect=AssertionError("不应重新加载")):
            app.button(key="retry_import").click().run()
        task = app.session_state["import_tasks"][0]
        self.assertEqual(task["status"], "success")
        self.assertTrue(task["indexed"])
        self.assertEqual(task["added_chunks"], 1)
        self.assertEqual(task["attempts"], 2)

    def test_vector_search_uses_persisted_index_and_top_k(self):
        """页面没有上传任务也能查旧库，展示排序、跨页来源与负相似度。"""
        chunks = [
            Document(page_content="神经网络实验", metadata={"chunk_id": "a1", "doc_id": "a",
                     "source_file": "论文A.pdf", "page_number": 2, "page_end": 3}),
            Document(page_content="农业实验", metadata={"chunk_id": "b1", "doc_id": "b",
                     "source_file": "论文B.pdf", "page_number": 1}),
            Document(page_content="反向向量实验", metadata={"chunk_id": "c1", "doc_id": "c",
                     "source_file": "论文C.pdf", "page_number": 4}),
        ]
        VectorStore().add_chunks(chunks)
        app = self.app
        self.assertEqual(app.number_input(key="vector_top_k").value, 5)
        self.assertEqual(app.session_state["import_tasks"], [])
        app.text_input(key="vector_query").set_value("神经网络")
        app.number_input(key="vector_top_k").set_value(2)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["神经网络实验", "农业实验"])
        self.assertTrue(any("1. 论文A.pdf · 余弦相似度 1.0000" in panel.label for panel in app.expander))
        self.assertIn("来源：论文A.pdf；物理页码：2–3", [element.value for element in app.caption])
        app.number_input(key="vector_top_k").set_value(10)
        app.button(key="vector_search").click().run()
        self.assertEqual(len(app.text), 3)
        self.assertTrue(any(panel.label.startswith("3. ") and "余弦相似度 -1.0000" in panel.label
                            for panel in app.expander))
        self.assertEqual(self.embeddings.query_calls, ["神经网络", "神经网络"])
        self.assertEqual(self.embeddings.document_calls, [[chunk.page_content for chunk in chunks]])

    def test_vector_search_document_filter_and_non_pdf_locations(self):
        """文档过滤生效，Word 显示段落/表格、TXT 显示行号，不伪造页码。"""
        VectorStore().add_chunks([
            Document(page_content="Word 神经网络正文", metadata={"chunk_id": "w1", "doc_id": "word",
                     "source_file": "论文.docx", "paragraph_index": 3}),
            Document(page_content="反向表格", metadata={"chunk_id": "w2", "doc_id": "word",
                     "source_file": "论文.docx", "table_index": 2}),
            Document(page_content="农业文本", metadata={"chunk_id": "t1", "doc_id": "text",
                     "source_file": "论文.txt", "line_start": 4, "line_end": 8}),
        ])
        app = self.app
        app.text_input(key="vector_query").set_value("神经网络")
        app.text_input(key="vector_doc_id").set_value(" word ")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["Word 神经网络正文", "反向表格"])
        captions = [element.value for element in app.caption]
        self.assertIn("来源：论文.docx；段落：3", captions)
        self.assertIn("来源：论文.docx；表格：2", captions)
        self.assertFalse(any("物理页码" in value for value in captions))
        app.text_input(key="vector_query").set_value("农业")
        app.text_input(key="vector_doc_id").set_value("text")
        app.button(key="vector_search").click().run()
        self.assertEqual([element.value for element in app.text], ["农业文本"])
        self.assertIn("来源：论文.txt；行范围：4–8", [element.value for element in app.caption])
        app.text_input(key="vector_doc_id").set_value("missing")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.text)
        self.assertTrue(any("没有可检索的文档块" in element.value for element in app.info))
        self.assertEqual(self.embeddings.query_calls, ["神经网络", "农业"])

    def test_vector_search_blank_input_and_empty_index(self):
        """启动/空问题不初始化模型，空库明确提示且不计算查询向量。"""
        app = self.app
        with patch("src.retrieval.vector_store.get_embeddings") as model:
            app.run()
            app.text_input(key="vector_query").set_value("  ")
            app.button(key="vector_search").click().run()
            model.assert_not_called()
        self.assertFalse(app.exception)
        self.assertTrue(any("请输入查询内容" in element.value for element in app.warning))
        app.text_input(key="vector_query").set_value("神经网络")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("没有可检索的文档块" in element.value for element in app.info))
        self.assertEqual(self.embeddings.query_calls, [])

    def test_vector_search_errors_are_visible_and_retryable(self):
        """模型/数据库失败明确报错，修复后重新提交可检索，不当作空库。"""
        VectorStore().add_chunks([Document(page_content="神经网络", metadata={
            "chunk_id": "a1", "doc_id": "a", "source_file": "论文.txt", "line_start": 1, "line_end": 1})])
        app = self.app
        app.text_input(key="vector_query").set_value("神经网络")
        for target, error in (("src.retrieval.vector_store.get_embeddings", FileNotFoundError("本地模型不存在")),
                              ("src.retrieval.vector_store.VectorStore.search", RuntimeError("数据库不可用"))):
            with self.subTest(target=target), patch(target, side_effect=error):
                app.button(key="vector_search").click().run()
            self.assertFalse(app.exception)
            self.assertTrue(any(str(error) in element.value for element in app.error))
            self.assertFalse(any("没有可检索的文档块" in element.value for element in app.info))
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual([element.value for element in app.text], ["神经网络"])

    def test_bm25_search_without_model_and_with_document_filter(self):
        """页面 BM25 不加载权重，单文献负分仍展示正文、位置与 ID。"""
        VectorStore().add_chunks([Document(page_content="BatchNormalization", metadata={
            "chunk_id": "bn1", "doc_id": "bn", "source_file": "论文.pdf", "page_number": 2})])
        app = self.app
        app.selectbox(key="retrieval_method").set_value("BM25 关键词")
        app.text_input(key="vector_query").set_value("BATCHNORMALIZATION")
        app.text_input(key="vector_doc_id").set_value(" bn ")
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("不能加载模型")) as model:
            app.button(key="vector_search").click().run()
            model.assert_not_called()
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["BatchNormalization"])
        self.assertTrue(any("BM25 分数 -" in panel.label for panel in app.expander))
        self.assertIn("来源：论文.pdf；物理页码：2", [element.value for element in app.caption])
        self.assertEqual(self.embeddings.query_calls, [])
        app.text_input(key="vector_query").set_value("Normalization")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.text)
        self.assertTrue(any("关键词无匹配" in element.value for element in app.info))

    def test_bm25_search_reflects_upload_and_delete(self):
        """上传后每次查询读取当前语料，新文档立即可查，已删除文档不再出现。"""
        app = self.app
        app.file_uploader[0].set_value([("first.txt", b"BatchNormalization", "text/plain")]).run()
        app.button(key="start_import").click().run()
        first_doc_id = app.session_state["import_tasks"][0]["documents"][0].metadata["doc_id"]
        app.selectbox(key="retrieval_method").set_value("BM25 关键词")
        app.text_input(key="vector_query").set_value("BatchNormalization")
        app.button(key="vector_search").click().run()
        self.assertEqual([element.value for element in app.text], ["BatchNormalization"])
        app.file_uploader[0].set_value([("second.txt", b"Adam", "text/plain")]).run()
        app.button(key="start_import").click().run()
        app.text_input(key="vector_query").set_value("Adam")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["Adam"])
        VectorStore().delete_document(first_doc_id)
        app.text_input(key="vector_query").set_value("BatchNormalization")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.text)
        self.assertEqual(self.embeddings.document_calls, [["BatchNormalization"], ["Adam"]])
        self.assertEqual(self.embeddings.query_calls, [])

    def test_bm25_database_error_is_visible_and_retryable(self):
        """语料读取异常明确提示，恢复后重新提交可查，不静默切换检索方式。"""
        VectorStore().add_chunks([Document(page_content="Adam", metadata={
            "chunk_id": "a1", "doc_id": "a", "source_file": "论文.txt", "line_start": 1, "line_end": 1})])
        app = self.app
        app.selectbox(key="retrieval_method").set_value("BM25 关键词")
        app.text_input(key="vector_query").set_value("Adam")
        with patch("src.retrieval.vector_store.VectorStore.list_chunks", side_effect=RuntimeError("语料读取失败")):
            app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("语料读取失败" in element.value for element in app.error))
        self.assertFalse(app.text)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual([element.value for element in app.text], ["Adam"])
        self.assertEqual(self.embeddings.query_calls, [])

    def test_rrf_search_displays_fused_score_and_original_source(self):
        """页面切换混合检索，共同命中的块排序提高，并保留 Word 位置。"""
        VectorStore().add_chunks([
            Document(page_content="神经网络", metadata={"chunk_id": "a", "doc_id": "a",
                     "source_file": "论文A.pdf", "page_number": 1}),
            Document(page_content="BatchNormalization", metadata={"chunk_id": "b", "doc_id": "b",
                     "source_file": "论文B.docx", "paragraph_index": 3}),
            Document(page_content="农业 BatchNormalization", metadata={"chunk_id": "c", "doc_id": "c",
                     "source_file": "论文C.txt", "line_start": 1, "line_end": 1}),
        ])
        app = self.app
        app.selectbox(key="retrieval_method").set_value("RRF 混合检索")
        app.text_input(key="vector_query").set_value("BatchNormalization")
        app.number_input(key="vector_top_k").set_value(1)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["BatchNormalization"])
        self.assertTrue(any("1. 论文B.docx · RRF 分数" in panel.label for panel in app.expander))
        self.assertIn("来源：论文B.docx；段落：3", [element.value for element in app.caption])
        self.assertEqual(self.embeddings.query_calls, ["BatchNormalization"])

    def test_rrf_empty_input_empty_store_and_model_error(self):
        """空问题/空库无需权重；有数据但模型失败不静默改为 BM25。"""
        app = self.app
        app.selectbox(key="retrieval_method").set_value("RRF 混合检索")
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("不应加载模型")) as model:
            app.button(key="vector_search").click().run()
            app.text_input(key="vector_query").set_value("BERT")
            app.button(key="vector_search").click().run()
            model.assert_not_called()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertFalse(app.text)
        VectorStore().add_chunks([Document(page_content="BERT", metadata={
            "chunk_id": "b", "doc_id": "b", "source_file": "论文.txt", "line_start": 1, "line_end": 1})])
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=FileNotFoundError("本地模型不存在")):
            app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("本地模型不存在" in element.value for element in app.error))
        self.assertFalse(app.text)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.error)
        self.assertEqual([element.value for element in app.text], ["BERT"])

    def test_rrf_search_after_upload_and_database_recovery(self):
        """新上传文档参加两路召回，数据库异常恢复后可查，删除过滤无陈旧块。"""
        app = self.app
        for filename, text in (("first.txt", b"BERT"), ("second.txt", b"Reranker")):
            app.file_uploader[0].set_value([(filename, text, "text/plain")]).run()
            app.button(key="start_import").click().run()
        doc_id = app.session_state["import_tasks"][0]["documents"][0].metadata["doc_id"]
        app.selectbox(key="retrieval_method").set_value("RRF 混合检索")
        app.text_input(key="vector_query").set_value("Reranker")
        app.number_input(key="vector_top_k").set_value(1)
        with patch("src.retrieval.vector_store.VectorStore.list_chunks", side_effect=RuntimeError("语料读取失败")):
            app.button(key="vector_search").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("语料读取失败" in element.value for element in app.error))
        self.assertFalse(app.text)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.error)
        self.assertEqual([element.value for element in app.text], ["Reranker"])
        VectorStore().delete_document(doc_id)
        app.text_input(key="vector_doc_id").set_value(doc_id)
        app.button(key="vector_search").click().run()
        self.assertFalse(app.text)
        self.assertEqual(self.embeddings.document_calls, [["BERT"], ["Reranker"]])

    def test_model_reranking_displays_new_order_score_and_sources(self):
        """页面使用模型分数重新排列，保留 Word 段落与 TXT 行范围。"""
        VectorStore().add_chunks([
            Document(page_content="神经网络", metadata={"chunk_id": "a", "doc_id": "a",
                     "source_file": "论文A.pdf", "page_number": 1}),
            Document(page_content="BatchNormalization", metadata={"chunk_id": "b", "doc_id": "b",
                     "source_file": "论文B.docx", "paragraph_index": 3}),
            Document(page_content="农业 BatchNormalization", metadata={"chunk_id": "c", "doc_id": "c",
                     "source_file": "论文C.txt", "line_start": 2, "line_end": 3}),
        ])
        app = self.app
        app.selectbox(key="retrieval_method").set_value("RRF + 模型重排")
        app.text_input(key="vector_query").set_value("BatchNormalization")
        app.number_input(key="vector_top_k").set_value(2)
        with patch("src.retrieval.reranker.get_reranker") as model:
            # 按正文给出明确测试分数，模型无需与实际论文质量等同。
            values = {"神经网络": 0.1, "BatchNormalization": 0.8, "农业 BatchNormalization": 0.9}
            model.return_value.predict.side_effect = lambda pairs, **kwargs: [values[text] for _, text in pairs]
            app.button(key="vector_search").click().run()
            self.assertEqual(len(model.return_value.predict.call_args.args[0]), 3)
        self.assertFalse(app.exception)
        self.assertEqual([element.value for element in app.text], ["农业 BatchNormalization", "BatchNormalization"])
        self.assertTrue(any("论文C.txt · 模型相关性分数 0.9000" in panel.label for panel in app.expander))
        self.assertIn("来源：论文C.txt；行范围：2–3", [element.value for element in app.caption])
        self.assertIn("来源：论文B.docx；段落：3", [element.value for element in app.caption])

    def test_model_reranking_empty_cases_and_error_recovery(self):
        """空问题/空库无需模型，模型缺失或推理失败有错误提示，修复后可重新提交。"""
        app = self.app
        app.selectbox(key="retrieval_method").set_value("RRF + 模型重排")
        with patch("src.retrieval.reranker.get_reranker", side_effect=AssertionError("不应加载模型")) as model:
            app.button(key="vector_search").click().run()
            app.text_input(key="vector_query").set_value("BERT")
            app.button(key="vector_search").click().run()
            model.assert_not_called()
        self.assertFalse(app.error)
        VectorStore().add_chunks([Document(page_content="BERT", metadata={
            "chunk_id": "b", "doc_id": "b", "source_file": "论文.txt", "line_start": 1, "line_end": 1})])
        for message in ("重排模型不存在", "重排推理失败"):
            with patch("src.retrieval.reranker.get_reranker", side_effect=RuntimeError(message)):
                app.button(key="vector_search").click().run()
            self.assertFalse(app.exception)
            self.assertTrue(any(message in element.value for element in app.error))
            self.assertFalse(app.text)
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.8]
            app.button(key="vector_search").click().run()
        self.assertFalse(app.error)
        self.assertEqual([element.value for element in app.text], ["BERT"])

    def test_model_reranking_uploaded_document_filter_and_delete(self):
        """新文献可精排，文档 ID 限定在模型评分前生效，删除后不加载重排模型。"""
        app = self.app
        for filename, text in (("first.txt", b"BERT"), ("second.txt", b"Reranker")):
            app.file_uploader[0].set_value([(filename, text, "text/plain")]).run()
            app.button(key="start_import").click().run()
        doc_id = app.session_state["import_tasks"][0]["documents"][0].metadata["doc_id"]
        app.selectbox(key="retrieval_method").set_value("RRF + 模型重排")
        app.text_input(key="vector_query").set_value("Reranker")
        app.text_input(key="vector_doc_id").set_value(doc_id)
        with patch("src.retrieval.reranker.get_reranker") as model:
            model.return_value.predict.return_value = [0.7]
            app.button(key="vector_search").click().run()
            self.assertEqual([element.value for element in app.text], ["Reranker"])
            model.return_value.predict.assert_called_once_with(
                [["Reranker", "Reranker"]], batch_size=8, show_progress_bar=False)
            VectorStore().delete_document(doc_id)
            model.reset_mock()
            app.button(key="vector_search").click().run()
            self.assertFalse(app.text)
            model.assert_not_called()
        self.assertFalse(app.exception)
        self.assertEqual(self.embeddings.document_calls, [["BERT"], ["Reranker"]])


    def test_layout_library_delete_cancel_restore_and_cache(self):
        """真实页面上传、取消删除、删除和恢复均核验原文及Chroma。"""
        from src.frontend.components.documents import list_documents
        app = self.app
        raw, index = Path(self.directory.name) / "raw", Path(self.directory.name) / "index"
        app.file_uploader[0].set_value([("paper.md", b"Transformer uses six layers.", "text/markdown")]).run()
        app.button(key="start_import").click().run()
        doc_id = list_documents(raw, index)[0]["doc_id"]
        self.assertEqual([tab.label for tab in app.tabs], ["Agent 科研助理", "RAG 流式问答"])
        self.assertEqual(app.sidebar.get("progress")[0].proto.value, 100)
        self.assertEqual(app.sidebar.selectbox(key="manage_doc_id").value, doc_id)
        app.button(key="delete_document").click().run()
        self.assertEqual(VectorStore().count(), 1)
        app.button(key="cancel_delete_document").click().run()
        self.assertTrue((raw / doc_id / "paper.md").is_file())
        app.session_state["rag_cache"].clear = unittest.mock.Mock()
        app.button(key="delete_document").click().run()
        app.button(key="confirm_delete_document").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(VectorStore().count(), 0)
        self.assertFalse((raw / doc_id).exists())
        self.assertEqual((raw / ".trash" / doc_id / "paper.md").read_bytes(), b"Transformer uses six layers.")
        self.assertEqual(app.session_state["import_tasks"], [])
        self.assertFalse(app.button(key="start_import").disabled)
        app.session_state["rag_cache"].clear.assert_called_once()
        app.button(key="restore_document").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(VectorStore().count(), 1)
        self.assertTrue(app.session_state["import_tasks"][0]["indexed"])
        self.assertTrue((raw / doc_id / "paper.md").exists())
        self.assertFalse((raw / ".trash" / doc_id).exists())
        app.run()
        self.assertEqual(VectorStore().count(), 1)

    def test_delete_only_selected_document_keeps_other_index_and_progress(self):
        """删除一份文献不会清空全库，保留任务的进度与真实列表一致。"""
        app = self.app
        app.file_uploader[0].set_value([("a.txt", b"Adam", "text/plain"),
                                       ("b.txt", b"Transformer", "text/plain")]).run()
        app.button(key="start_import").click().run()
        app.button(key="delete_document").click().run()
        app.button(key="confirm_delete_document").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(VectorStore().count(), 1)
        self.assertEqual(VectorStore().list_chunks()[0].page_content, "Transformer")
        self.assertEqual(app.session_state["import_progress"], {"completed": 1, "total": 1})
        self.assertEqual(len(app.selectbox(key="manage_doc_id").options), 1)
        app.selectbox(key="retrieval_method").set_value("BM25 关键词")
        app.text_input(key="vector_query").set_value("Adam")
        app.button(key="vector_search").click().run()
        self.assertFalse(app.text)

    def test_library_persists_without_upload_and_failure_keeps_original(self):
        """页面重载读全库；索引删除失败会回滚原文且允许重试。"""
        app = self.app
        app.file_uploader[0].set_value([("keep.txt", b"Adam", "text/plain")]).run()
        app.button(key="start_import").click().run()
        app.session_state["import_tasks"] = []
        app.file_uploader[0].set_value([]).run()
        doc_id = app.selectbox(key="manage_doc_id").value
        app.button(key="delete_document").click().run()
        with patch("src.retrieval.vector_store.VectorStore.delete_document", side_effect=RuntimeError("索引删除失败")):
            app.button(key="confirm_delete_document").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("索引删除失败" in e.value for e in app.error))
        self.assertEqual(VectorStore().count(), 1)
        self.assertTrue((Path(self.directory.name) / "raw" / doc_id / "keep.txt").exists())
        app.button(key="confirm_delete_document").click().run()
        self.assertEqual(VectorStore().count(), 0)

    def test_restore_index_failure_can_retry(self):
        """回收原文恢复后模型失败，不冒充入库成功；现有失败重试补全索引。"""
        app = self.app
        app.file_uploader[0].set_value([("restore.txt", b"Neural network", "text/plain")]).run()
        app.button(key="start_import").click().run()
        app.button(key="delete_document").click().run()
        app.button(key="confirm_delete_document").click().run()
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=FileNotFoundError("权重缺失")):
            app.button(key="restore_document").click().run()
        task = app.session_state["import_tasks"][0]
        self.assertEqual(task["status"], "failed")
        self.assertTrue(Path(task["path"]).exists())
        self.assertEqual(VectorStore().count(), 0)
        app.button(key="retry_import").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(VectorStore().count(), 1)


    def knowledge_rows(self, app=None):
        """按字段找到只读表格，避免依赖新增面板后的全局元素顺序。"""
        return next(table.value for table in (app or self.app).dataframe if "向量化状态" in table.value.columns)

    def test_knowledge_panel_empty_is_read_only(self):
        """空库显示真实零值，刷新不会建向量库、编码或改变当前会话。"""
        app = self.app
        self.assertEqual({m.label: m.value for m in app.metric if m.label.startswith("知识库")
                          or m.label == "已向量化文档数"},
                         {"知识库文档数": "0", "已向量化文档数": "0", "知识库索引块数": "0"})
        session = app.session_state["agent_session_id"]
        app.button(key="refresh_knowledge").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["agent_session_id"], session)
        self.assertFalse((Path(self.directory.name) / "index").exists())
        self.assertEqual(self.embeddings.document_calls, [])

    def test_knowledge_panel_tracks_failed_unindexed_and_refresh(self):
        """成功与写入失败并存，刷新后磁盘状态仍准确但不补造本页批次结果。"""
        from streamlit.testing.v1 import AppTest
        app = self.app
        app.file_uploader[0].set_value([("ok.txt", b"Adam", "text/plain"),
                                       ("failed.txt", b"Failure", "text/plain")]).run()
        original = VectorStore.add_chunks
        def write_or_fail(store, chunks):
            if chunks[0].metadata["source_file"] == "failed.txt":
                raise OSError("测试索引写入失败")
            return original(store, chunks)
        with patch.object(VectorStore, "add_chunks", autospec=True, side_effect=write_or_fail):
            app.button(key="start_import").click().run()
        rows = self.knowledge_rows(app).set_index("文件名")
        self.assertEqual(rows.loc["ok.txt", "向量化状态"], "已向量化")
        self.assertEqual(rows.loc["failed.txt", "向量化状态"], "未向量化")
        self.assertEqual(rows.loc["failed.txt", "导入结果（本页）"], "失败")
        self.assertIn("索引写入失败", rows.loc["failed.txt", "错误（本页）"])
        self.assertEqual({m.label: m.value for m in app.metric if m.label.startswith("知识库")
                          or m.label == "已向量化文档数"},
                         {"知识库文档数": "2", "已向量化文档数": "1", "知识库索引块数": "1"})
        calls = len(self.embeddings.document_calls)
        restored = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "src/frontend/app.py"),
                                     default_timeout=10).run()
        rows = self.knowledge_rows(restored)
        self.assertEqual(set(rows["导入结果（本页）"]), {"—"})
        self.assertEqual(set(rows["向量化状态"]), {"已向量化", "未向量化"})
        self.assertEqual(len(self.embeddings.document_calls), calls)
        app.button(key="retry_import").click().run()
        rows = self.knowledge_rows(app)
        self.assertEqual(set(rows["向量化状态"]), {"已向量化"})
        self.assertEqual(set(rows["导入结果（本页）"]), {"成功"})

    def test_knowledge_panel_partial_index_does_not_claim_import_success(self):
        """第二批失败时500块仍存在，面板同时保留失败/预期501块，重试补全。"""
        import hashlib
        app = self.app
        data = b"Paper blocks"
        doc_id = hashlib.sha256(data).hexdigest()
        chunks = [Document(page_content=f"block {i}", metadata={"doc_id": doc_id,
                  "chunk_id": f"block-{i}", "source_file": "partial.txt"}) for i in range(501)]
        app.file_uploader[0].set_value([("partial.txt", data, "text/plain")]).run()
        original = VectorStore.add_chunks
        def fail_last_batch(store, batch):
            if len(batch) == 1:
                raise TimeoutError("第二批失败")
            return original(store, batch)
        with patch("src.chunking.split_documents", return_value=chunks), \
                patch.object(VectorStore, "add_chunks", autospec=True, side_effect=fail_last_batch):
            app.button(key="start_import").click().run()
        row = self.knowledge_rows(app).iloc[0]
        self.assertEqual(row["索引块数"], 500)
        self.assertEqual(row["导入结果（本页）"], "失败")
        self.assertEqual(row["本次预期块数"], "501")
        self.assertTrue(any("不保证完整入库" in c.value for c in app.caption))
        with patch("src.chunking.split_documents", return_value=chunks):
            app.button(key="retry_import").click().run()
        self.assertEqual(self.knowledge_rows(app).iloc[0]["索引块数"], 501)
        self.assertEqual(len(self.embeddings.document_calls[-1]), 1)

    def test_knowledge_panel_missing_source_and_external_index_change(self):
        """只剩索引时明确原文缺失；外部移除块后刷新显示未向量化，不沿用旧批次成功。"""
        app = self.app
        app.file_uploader[0].set_value([("paper.txt", b"BERT", "text/plain")]).run()
        app.button(key="start_import").click().run()
        task = app.session_state["import_tasks"][0]
        source = Path(task["path"])
        source.unlink()
        app.button(key="refresh_knowledge").click().run()
        row = self.knowledge_rows(app).iloc[0]
        self.assertEqual(row["原文状态"], "缺失")
        self.assertEqual(row["向量化状态"], "已向量化")
        self.assertTrue(app.button(key="delete_document").disabled)
        source.write_bytes(b"BERT")
        VectorStore().delete_document(row["文档 ID"])
        app.button(key="refresh_knowledge").click().run()
        row = self.knowledge_rows(app).iloc[0]
        self.assertEqual(row["向量化状态"], "未向量化")
        self.assertEqual(row["索引块数"], 0)
        self.assertEqual(row["导入结果（本页）"], "成功")  # 原批次事实不能冒充当前索引状态。

    def test_knowledge_panel_content_alias_keeps_both_batch_results(self):
        """同内容异名按指纹合并，仍保留一个成功和另一个失败的实际结果。"""
        app = self.app
        app.file_uploader[0].set_value([("a.txt", b"BERT", "text/plain"),
                                       ("b.txt", b"BERT", "text/plain")]).run()
        original = VectorStore.add_chunks
        def fail_alias(store, chunks):
            if chunks[0].metadata["source_file"] == "b.txt":
                raise OSError("第二个别名失败")
            return original(store, chunks)
        with patch.object(VectorStore, "add_chunks", autospec=True, side_effect=fail_alias):
            app.button(key="start_import").click().run()
        rows = self.knowledge_rows(app)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows.iloc[0]["文件名"], "a.txt / b.txt")
        self.assertEqual(rows.iloc[0]["索引块数"], 1)
        self.assertEqual(rows.iloc[0]["导入结果（本页）"], "成功 / 失败")
        self.assertIn("第二个别名失败", rows.iloc[0]["错误（本页）"])

    def test_knowledge_panel_index_error_is_unknown_not_zero(self):
        """索引不可读不显示正常空库或捏造零统计，修复后刷新重新读取。"""
        app = self.app
        app.file_uploader[0].set_value([("paper.txt", b"BERT", "text/plain")]).run()
        app.button(key="start_import").click().run()
        with patch("src.frontend.components.documents.VectorStore.list_chunks", side_effect=RuntimeError("索引不可读")):
            app.button(key="refresh_knowledge").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("索引不可读" in e.value for e in app.error))
        self.assertFalse([m for m in app.metric if m.label.startswith("知识库")])
        self.assertFalse([d for d in app.dataframe if "向量化状态" in d.value.columns])
        app.button(key="refresh_knowledge").click().run()
        self.assertEqual(self.knowledge_rows(app).iloc[0]["索引块数"], 1)


class TestDocumentManagement(unittest.TestCase):
    """回收边界与实际加载任务验证；无需生成模型。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.raw = Path(self.directory.name) / "raw"
        self.index = Path(self.directory.name) / "index"
        self.data = b"Six encoder layers"
        import hashlib
        self.doc_id = hashlib.sha256(self.data).hexdigest()
        self.folder = self.raw / self.doc_id
        self.folder.mkdir(parents=True)
        (self.folder / "paper.txt").write_bytes(self.data)

    def test_unindexed_list_delete_restore_without_creating_database(self):
        from src.frontend.components.documents import list_documents, delete_document, restore_document
        with patch("src.retrieval.vector_store.get_embeddings", side_effect=AssertionError("不可加载模型")):
            self.assertEqual(list_documents(self.raw, self.index)[0]["chunks"], 0)
            self.assertEqual(delete_document(self.raw, self.index, self.doc_id), 0)
            self.assertEqual(list_documents(self.raw, self.index), [])
            tasks = restore_document(self.raw, self.doc_id)
        self.assertEqual(tasks[0]["data"], self.data)
        self.assertEqual(tasks[0]["status"], "pending")
        self.assertFalse(self.index.exists())

    def test_reject_path_and_symlink(self):
        from src.frontend.components.documents import delete_document, list_documents
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, "../outside")
        link = self.raw / ("a" * 64)
        link.symlink_to(self.folder, target_is_directory=True)
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, link.name)
        self.assertEqual(len(list_documents(self.raw, self.index)), 1)
        self.assertTrue((self.folder / "paper.txt").exists())

    def test_archive_conflict_preserves_source(self):
        from src.frontend.components.documents import delete_document
        (self.raw / ".trash" / self.doc_id).mkdir(parents=True)
        with self.assertRaises(ValueError):
            delete_document(self.raw, self.index, self.doc_id)
        self.assertEqual((self.folder / "paper.txt").read_bytes(), self.data)

    def test_restore_rejects_changed_content_and_no_overwrite(self):
        from src.frontend.components.documents import delete_document, restore_document
        delete_document(self.raw, self.index, self.doc_id)
        archived = self.raw / ".trash" / self.doc_id / "paper.txt"
        archived.write_bytes(b"Changed")
        with self.assertRaises(ValueError):
            restore_document(self.raw, self.doc_id)
        self.assertTrue(archived.exists())
        archived.write_bytes(self.data)
        self.folder.mkdir()
        with self.assertRaises(FileExistsError):
            restore_document(self.raw, self.doc_id)
        self.assertTrue(archived.exists())

    def test_missing_source_never_deletes_index(self):
        from src.frontend.components.documents import delete_document
        with patch("src.frontend.components.documents.VectorStore") as store:
            with self.assertRaises(FileNotFoundError):
                delete_document(self.raw, self.index, "b" * 64)
            store.assert_not_called()

    def test_graph_parallel_ids_and_join(self):
        from src.frontend.components.trace import trace_graph
        graph = trace_graph([
            {"type": "thought"},
            {"type": "tool_call", "name": 'a"tool', "call_id": "a"},
            {"type": "tool_call", "name": "b", "call_id": "b"},
            {"type": "tool_result", "name": "b", "call_id": "b"},
            {"type": "tool_result", "name": "a", "call_id": "a"},
            {"type": "observation", "decision": "continue"},
            {"type": "thought"}, {"type": "done", "stop_reason": "max_iterations"}])
        for edge in ("n0 -> n1", "n0 -> n2", "n2 -> n3", "n1 -> n4", "n3 -> n5", "n4 -> n5", "n5 -> n6"):
            self.assertIn(edge, graph)
        self.assertIn('a\\"tool', graph)
        self.assertNotIn("n1 -> n2", graph)

class TestHealthCheck(unittest.TestCase):
    """隔离Ollama HTTP，Chroma使用真实临时数据库；不生成回答或编码向量。"""

    def setUp(self):
        import httpx
        from src.utils.config import load_config
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.config = load_config()
        self.config["paths"]["vector_index"] = self.directory.name
        self.config["paths"]["logs"] = str(Path(self.directory.name) / "logs")
        self.embeddings = SmallEmbeddings()
        self.store = VectorStore(self.directory.name, self.embeddings)
        for name in ("src.utils.config.load_config", "src.retrieval.vector_store.load_config"):
            patcher = patch(name, return_value=self.config)
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = patch("src.utils.config.httpx.Client")
        self.http = patcher.start()
        self.addCleanup(patcher.stop)
        self.get = self.http.return_value.__enter__.return_value.get
        self.response = httpx.Response(200, json={"models": [{"name": self.config["llm"]["model"]}]},
                                       request=httpx.Request("GET", "http://localhost:11434/api/tags"))
        self.get.return_value = self.response

    def check(self):
        from src.utils.config import check_health
        return check_health()

    def test_existing_empty_database_is_available_not_missing(self):
        result = self.check()
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["vector_database"]["chunks"], 0)
        self.assertTrue(result["llm"]["service_reachable"])
        self.assertTrue(result["llm"]["model_available"])
        json.dumps(result, allow_nan=False)

    def test_nonempty_database_keeps_chunks_metadata_and_no_embedding_calls(self):
        chunks = [Document(page_content="实际临时论文块", metadata={"doc_id": "paper", "chunk_id": "chunk"})]
        self.store.add_chunks(chunks)
        calls = list(self.embeddings.document_calls)
        metadata = deepcopy(self.store._store._collection.metadata)
        with patch("src.retrieval.vector_store.get_embeddings") as model:
            result = self.check()
        model.assert_not_called()
        self.assertEqual(result["vector_database"]["chunks"], 1)
        self.assertEqual(self.store.list_chunks(), chunks)
        self.assertEqual(self.store._store._collection.metadata, metadata)
        self.assertEqual(self.embeddings.document_calls, calls)
        self.assertEqual(self.embeddings.query_calls, [])

    def test_connection_failure_does_not_skip_database_check(self):
        import httpx
        self.get.side_effect = httpx.ConnectError("服务未启动")
        result = self.check()
        self.assertEqual(result["status"], "degraded")
        self.assertEqual(result["llm"]["status"], "error")
        self.assertFalse(result["llm"]["service_reachable"])
        self.assertIsNone(result["llm"]["model_available"])
        self.assertEqual(result["vector_database"]["status"], "ok")
        self.assertIn("服务未启动", result["llm"]["detail"])

    def test_network_timeout_has_friendly_error_and_three_second_limit(self):
        import httpx
        self.get.side_effect = httpx.ReadTimeout("响应超时")
        result = self.check()
        self.assertIn("响应超时", result["llm"]["detail"])
        self.http.assert_called_once_with(timeout=3, trust_env=False, follow_redirects=False)
        self.get.assert_called_once_with("http://localhost:11434/api/tags")
        self.assertGreaterEqual(result["llm"]["seconds"], 0)

    def test_reachable_service_missing_configured_model_is_not_ready(self):
        import httpx
        self.get.return_value = httpx.Response(200, json={"models": []}, request=self.response.request)
        result = self.check()
        self.assertEqual(result["status"], "degraded")
        self.assertEqual(result["llm"]["status"], "model_missing")
        self.assertTrue(result["llm"]["service_reachable"])
        self.assertFalse(result["llm"]["model_available"])

    def test_model_without_tag_matches_latest_only(self):
        import httpx
        self.config["llm"]["model"] = "qwen2.5"
        for installed, expected in (("qwen2.5:latest", True), ("qwen2.5:7b", False)):
            with self.subTest(installed=installed):
                self.get.return_value = httpx.Response(200, json={"models": [{"name": installed}]}, request=self.response.request)
                self.assertEqual(self.check()["llm"]["model_available"], expected)

    def test_cloud_or_wrong_provider_is_rejected_without_network(self):
        for provider, address in (("ollama", "https://example.com"), ("openai", "http://localhost:11434")):
            with self.subTest(provider=provider):
                self.config["llm"].update(provider=provider, base_url=address)
                self.assertEqual(self.check()["llm"]["status"], "error")
        self.http.assert_not_called()

    def test_http_error_and_redirect_do_not_report_reachable_service(self):
        import httpx
        for status in (302, 500):
            with self.subTest(status=status):
                self.get.return_value = httpx.Response(status, request=self.response.request)
                result = self.check()
                self.assertEqual(result["llm"]["status"], "error")
                self.assertFalse(result["llm"]["service_reachable"])

    def test_invalid_json_or_model_list_is_error_not_model_missing(self):
        import httpx
        for body in (b"bad json", b'{}', b'{"models":null}', b'{"models":[{}]}', b'{"models":[42]}'):
            with self.subTest(body=body):
                self.get.return_value = httpx.Response(200, content=body, request=self.response.request)
                result = self.check()
                self.assertEqual(result["llm"]["status"], "error")
                self.assertTrue(result["llm"]["service_reachable"])
                self.assertIsNone(result["llm"]["model_available"])

    def test_missing_index_never_creates_directory_or_database(self):
        path = Path(self.directory.name) / "not-created"
        self.config["paths"]["vector_index"] = str(path)
        result = self.check()
        self.assertEqual(result["vector_database"]["status"], "not_initialized")
        self.assertIsNone(result["vector_database"]["chunks"])
        self.assertFalse(path.exists())

    def test_missing_collection_never_creates_a_new_collection(self):
        self.config["retrieval"]["collection_name"] = "missing_collection"
        result = self.check()
        self.assertEqual(result["vector_database"]["status"], "error")
        self.assertEqual([c.name for c in self.store._store._client.list_collections()], ["paper_chunks"])

    def test_index_version_or_parameters_mismatch_is_not_available(self):
        for section, key, value in (("embedding", "revision", "different"), ("retrieval", "search_ef", 101)):
            old = self.config[section][key]
            with self.subTest(key=key):
                self.config[section][key] = value
                result = self.check()
                self.assertEqual(result["vector_database"]["status"], "error")
                self.assertIn("不一致", result["vector_database"]["detail"])
            self.config[section][key] = old

    def test_database_heartbeat_failure_keeps_independent_llm_status(self):
        with patch("chromadb.api.client.Client.heartbeat", side_effect=RuntimeError("数据库连接失败")):
            result = self.check()
        self.assertEqual(result["llm"]["status"], "ok")
        self.assertEqual(result["vector_database"]["status"], "error")
        self.assertIn("数据库连接失败", result["vector_database"]["detail"])

    def test_corrupt_sqlite_file_is_reported_as_database_error(self):
        path = Path(self.directory.name) / "corrupt"
        path.mkdir()
        (path / "chroma.sqlite3").write_bytes(b"corrupt sqlite")
        self.config["paths"]["vector_index"] = str(path)
        result = self.check()
        self.assertEqual(result["vector_database"]["status"], "error")
        self.assertIsNone(result["vector_database"]["chunks"])

    def test_invalid_directory_or_backend_is_reported_as_database_error(self):
        path = Path(self.directory.name) / "file"
        path.write_text("这不是索引目录")
        self.config["paths"]["vector_index"] = str(path)
        self.assertIn("不是目录", self.check()["vector_database"]["detail"])
        self.config["retrieval"]["vector_store"] = "faiss"
        self.assertIn("Chroma", self.check()["vector_database"]["detail"])


class TestHealthCheckPage(unittest.TestCase):
    """页面按需调用，实际检查函数另有真实数据库测试。"""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        from src.utils.config import load_config
        config = load_config()
        config["paths"]["session_db"] = str(Path(self.directory.name) / "memory.sqlite3")
        config_patcher = patch("src.utils.config.load_config", return_value=config)
        config_patcher.start()
        self.addCleanup(config_patcher.stop)
        self.result = {"status": "ok", "checked_at": "2026-10-03T12:00:00+08:00",
                       "llm": {"status": "ok", "detail": "模型服务正常，未执行推理。", "seconds": 0.01},
                       "vector_database": {"status": "ok", "detail": "集合可读取。", "seconds": 0.02,
                                           "chunks": 0, "collection": "paper_chunks"}}
        patcher = patch("src.utils.config.check_health", side_effect=lambda: deepcopy(self.result))
        self.checker = patcher.start()
        self.addCleanup(patcher.stop)

    def page(self):
        from streamlit.testing.v1 import AppTest
        return AppTest.from_file(str(Path(__file__).resolve().parents[1] / "src/frontend/app.py"), default_timeout=10).run()

    def test_page_startup_does_not_probe_service_or_create_health_snapshot(self):
        app = self.page()
        self.assertFalse(app.exception)
        self.checker.assert_not_called()
        self.assertTrue(any("尚未检查" in row.value for row in app.info))

    def test_manual_check_shows_both_results_and_rerun_does_not_probe_again(self):
        app = self.page()
        app.button(key="check_health").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.success), 2)
        self.assertTrue(any("文档块数：0" in c.value for c in app.caption))
        app.run()
        self.checker.assert_called_once()
        app.button(key="check_health").click().run()
        self.assertEqual(self.checker.call_count, 2)

    def test_failed_llm_and_missing_index_are_visible_independently(self):
        self.result["status"] = "degraded"
        self.result["llm"].update(status="error", detail="连接失败，请启动Ollama。")
        self.result["vector_database"].update(status="not_initialized", detail="请先导入文档。", chunks=None)
        app = self.page()
        app.button(key="check_health").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("启动Ollama" in c.value for c in app.error))
        self.assertTrue(any("导入文档" in c.value for c in app.warning))

    def test_missing_model_and_database_failure_are_not_green(self):
        self.result["status"] = "degraded"
        self.result["llm"].update(status="model_missing", detail="配置模型未安装。")
        self.result["vector_database"].update(status="error", detail="索引损坏。", chunks=None)
        app = self.page()
        app.button(key="check_health").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.success)
        self.assertTrue(any("未安装" in c.value for c in app.warning))
        self.assertTrue(any("索引损坏" in c.value for c in app.error))


if __name__ == "__main__":
    unittest.main()
