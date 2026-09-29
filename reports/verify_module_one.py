"""模块一真实链路核验：多格式导入、五组分块、增量索引与四种检索。

只使用本地权重与临时索引；生成的 Word/TXT/Markdown 用于功能验证，
公开论文用于真实 PDF 解析，不把这些样例当作质量评测数据。
"""

import hashlib
import io
import json
import math
import os
import platform
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))


def verify_ingestion(directory):
    """每个检查都有断言；任一步失败就停止，不生成“通过”的报告。"""
    import numpy as np
    import torch
    from docx import Document as WordDocument
    from langchain_core.embeddings import Embeddings
    from src.chunking import split_documents
    from src.data_loader import create_import_tasks
    from src.retrieval.bm25_retriever import BM25Retriever
    from src.retrieval.hybrid_retriever import HybridRetriever
    from src.retrieval.vector_store import VectorStore, batch_build_index, get_embeddings
    from src.utils.config import load_config

    torch.set_num_threads(4)
    encoded_batches = []
    actual_embeddings = get_embeddings()

    class RecordedEmbeddings(Embeddings):
        """仅记录实际文档编码批次；所有向量仍由真实 M3E 产生。"""
        def embed_documents(self, texts):
            encoded_batches.append(len(texts))
            return actual_embeddings.embed_documents(texts)

        def embed_query(self, text):
            return actual_embeddings.embed_query(text)

    manifest = json.loads((ROOT / "reports/Embedding论文双语评测集.json").read_text())
    paper = next(row for row in manifest["papers"] if row["id"] == "attention")
    pdf = (ROOT / "data/raw/embedding_papers/attention.pdf").read_bytes()
    assert hashlib.sha256(pdf).hexdigest() == paper["pdf_sha256"]
    word = WordDocument()
    word.add_paragraph("论文方法：Transformer 使用自注意力处理序列，能够并行训练。")
    table = word.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("模型", "主要机制"), ("Transformer", "Self Attention"))):
        for cell, text in zip(row.cells, values):
            cell.text = text
    word.add_paragraph("实验指标：BLEU 用于机器翻译，准确率用于分类评估。")
    buffer = io.BytesIO()
    word.save(buffer)
    files = [("attention.pdf", pdf), ("方法说明.docx", buffer.getvalue()),
             ("训练记录.txt", b"\xef\xbb\xbf" + ("训练设置：Dropout 用于缓解过拟合。\r\n" * 40).encode()),
             ("索引笔记.md", ("# 检索笔记\n\nRRF 融合向量与 BM25 排名。\n" * 30).encode()),
             ("实验补充.markdown", "# 引用\n\n科研文献需要保留文档来源。\n".encode()),
             ("损坏文件.pdf", b"not a pdf")]
    checks = {}
    directory = Path(directory)
    raw, index = Path(directory) / "raw", Path(directory) / "index"
    store = VectorStore(index, embeddings=RecordedEmbeddings())
    assert store.search("空库问题") == []
    assert HybridRetriever(store).search("空库问题", rerank=True) == []
    tasks = create_import_tasks(files + [files[0]])
    assert len(tasks) == 6
    progress = []
    phases = set()
    for event in batch_build_index(tasks, raw, vector_store=store):
        progress.append(event)
        phases.update(task["status"] for task in tasks)
    good = tasks[:-1]
    assert all(task["status"] == "success" and task["indexed"] for task in good)
    assert tasks[-1]["status"] == "failed" and not tasks[-1]["path"]
    assert progress[-1] == {"completed": 6, "total": 6}
    assert {"loading", "chunking", "indexing", "success", "failed"} <= phases
    expected_count = sum(task["chunk_count"] for task in good)
    assert store.count() == expected_count == sum(encoded_batches)
    stored = {doc.metadata["chunk_id"]: doc for doc in store.list_chunks()}
    for task in good:
        assert Path(task["path"]).read_bytes() == task["data"]
        for chunk in split_documents(task["documents"]):
            found = stored[chunk.metadata["chunk_id"]]
            assert found.page_content == chunk.page_content and found.metadata == chunk.metadata
    checks["initial_import"] = {
        "successes": len(good), "failures": 1, "stored_chunks": expected_count,
        "real_encoding_batches": list(encoded_batches), "progress": progress[-1], "phases": sorted(phases)}
    print(f"多格式真实入库通过：5 份成功、1 份损坏 PDF 失败，{expected_count} 块", flush=True)

    # 同一原文比较固定三档，另外两种只用 YAML 默认参数；不计算质量指标。
    settings = []
    for strategy in ("fixed", "recursive", "semantic"):
        for size in ((256, 512, 1024) if strategy == "fixed" else (None,)):
            chunks = split_documents(tasks[0]["documents"], strategy=strategy, chunk_size=size)
            repeated = split_documents(tasks[0]["documents"], strategy=strategy, chunk_size=size)
            assert [doc.metadata["chunk_id"] for doc in chunks] == [doc.metadata["chunk_id"] for doc in repeated]
            body = [doc for doc in chunks if doc.metadata.get("chunk_preserved") != "table"]
            tables = [doc for doc in chunks if doc.metadata.get("chunk_preserved") == "table"]
            for chunk in body:
                meta = chunk.metadata
                parent = next(doc for doc in tasks[0]["documents"] if doc.metadata.get("page") == meta["page"]
                              and doc.metadata.get("content_type") != "table")
                assert parent.page_content[meta["start_index"]:meta["end_index"]] == chunk.page_content
                assert len(chunk.page_content) <= meta["chunk_size"]
            assert len(tables) == 3
            settings.append({"strategy": strategy, "chunk_size": chunks[0].metadata["chunk_size"],
                             "body_chunks": len(body), "table_chunks": len(tables)})
    checks["five_chunk_settings"] = settings
    attempts = [task["attempts"] for task in tasks]
    list(batch_build_index(tasks, raw, retry_failed=True, vector_store=store))
    assert [task["attempts"] for task in good] == attempts[:-1]
    assert tasks[-1]["attempts"] == 2 and tasks[-1]["status"] == "failed"
    # 损坏字节未修复，重试应继续失败，不能伪装成恢复成功。
    checks["failed_only_retry"] = {"successful_attempts_unchanged": True, "corrupt_pdf_attempts": 2}
    batches_before = list(encoded_batches)
    duplicates = create_import_tasks(files[:-1])
    list(batch_build_index(duplicates, raw, vector_store=store))
    assert all(task["added_chunks"] == 0 and task["indexed"] for task in duplicates)
    assert encoded_batches == batches_before and store.count() == expected_count
    checks["duplicate_import"] = {"added_chunks": 0, "additional_document_encoding": 0}

    old_vectors = store._store.get(include=["embeddings"])
    old_vectors = dict(zip(old_vectors["ids"], old_vectors["embeddings"]))
    added = create_import_tasks([("追加资料.txt", "新文档：科研系统需要增量索引而非重建全量向量。".encode())])
    list(batch_build_index(added, raw, vector_store=store))
    assert added[0]["indexed"] and added[0]["added_chunks"] == added[0]["chunk_count"]
    assert sum(encoded_batches[len(batches_before):]) == added[0]["chunk_count"]
    assert store.count() == expected_count + added[0]["chunk_count"]
    current = store._store.get(ids=list(old_vectors), include=["embeddings"])
    assert all(np.array_equal(old_vectors[key], vector) for key, vector in zip(current["ids"], current["embeddings"]))
    checks["incremental_import"] = {"added_chunks": added[0]["chunk_count"], "old_vectors_unchanged": True,
                                     "only_new_chunks_encoded": True}

    search = {}
    bm25, hybrid = BM25Retriever(store), HybridRetriever(store)
    query = "Transformer Self Attention 论文方法"
    calls = {"vector": lambda **kw: store.search(query, **kw),
             "bm25": lambda **kw: bm25.search(query, **kw),
             "hybrid": lambda **kw: hybrid.search(query, **kw),
             "hybrid_rerank": lambda **kw: hybrid.search(query, rerank=True, **kw)}
    for name, call in calls.items():
        found = call(k=5)
        assert 0 < len(found) <= 5
        assert all(math.isfinite(score) for _, score in found)
        assert all(found[i][1] >= found[i+1][1] for i in range(len(found)-1))
        for task in good:
            doc_id = task["documents"][0].metadata["doc_id"]
            filtered = call(k=5, doc_id=doc_id)
            assert all(doc.metadata["doc_id"] == doc_id for doc, _ in filtered)
        search[name] = [doc.metadata["chunk_id"] for doc, _ in found]
    for task in good:
        doc_id = task["documents"][0].metadata["doc_id"]
        assert store.search("文档来源", doc_id=doc_id)
    pool = hybrid.search(query, k=20)
    assert len(pool) == 20
    assert set(search["hybrid_rerank"]) <= {doc.metadata["chunk_id"] for doc, _ in pool}
    assert hybrid.search(query, doc_id="不存在的文档", rerank=True) == []
    checks["retrieval"] = {"top5_ids": search, "rrf_candidate_count": 20,
                           "filtered_sources_preserved": True, "no_matching_document_returns_empty": True}

    doc_id = added[0]["documents"][0].metadata["doc_id"]
    assert store.delete_document(doc_id) == added[0]["chunk_count"]
    assert not store.list_chunks(doc_id) and not hybrid.search("增量索引", doc_id=doc_id)
    assert Path(added[0]["path"]).exists()
    checks["delete_from_index"] = {"source_file_retained": True, "deleted_document_unretrievable": True}

    config = load_config()
    report = {"verified_at": datetime.now().astimezone().isoformat(), "python": platform.python_version(),
              "platform": platform.platform(), "offline_models": True, "temporary_index": True,
              "embedding": config["embedding"], "retrieval": config["retrieval"],
              "paper_sha256": paper["pdf_sha256"], "checks": checks,
              "scope": "功能链路验证，不替代分块策略召回实验、独立人工标注或完整系统评测"}
    # 临时快照交给下一进程；只有重开核验也通过才写正式报告。
    state = {"report": report, "query": query,
             "vector_ids": [doc.metadata["chunk_id"] for doc, _ in store.search(query, k=5)],
             "chunks": {doc.metadata["chunk_id"]: {"text": doc.page_content, "metadata": doc.metadata}
                        for doc in store.list_chunks()}}
    (directory / "state.json").write_text(json.dumps(state, ensure_ascii=False))
    print("入库、增量、分块与检索通过；结束进程后再核验重开。", flush=True)


def verify_reopen(directory):
    """前一模型进程已退出，新进程真实加载 M3E 并查询持久化索引。"""
    import torch
    from src.retrieval.vector_store import VectorStore

    torch.set_num_threads(4)
    directory = Path(directory)
    state = json.loads((directory / "state.json").read_text())
    store = VectorStore(directory / "index")
    actual = {doc.metadata["chunk_id"]: {"text": doc.page_content, "metadata": doc.metadata}
              for doc in store.list_chunks()}
    assert actual == state["chunks"]
    assert [doc.metadata["chunk_id"] for doc, _ in store.search(state["query"], k=5)] == state["vector_ids"]
    report = state["report"]
    report["checks"]["new_process_reopen"] = {
        "stored_chunks": len(actual), "content_and_metadata_equal": True, "vector_top5_equal": True}
    (ROOT / "reports/模块一完整性验证结果.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print("独立进程重开与真实模型检索通过，完整核验结果已保存。", flush=True)


if __name__ == "__main__":
    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"
    os.environ["TOKENIZERS_PARALLELISM"] = "false"
    os.environ.setdefault("HF_HOME", str(ROOT / "data/models/.hf-runtime"))
    if len(sys.argv) == 3 and sys.argv[1] == "--ingestion":
        verify_ingestion(sys.argv[2])
    elif len(sys.argv) == 3 and sys.argv[1] == "--reopen":
        verify_reopen(sys.argv[2])
    else:
        # 调度进程只导入标准库。先退出入库进程，再启动重开进程，模拟应用重启。
        with tempfile.TemporaryDirectory(prefix="rag-module-one-") as directory:
            for stage in ("--ingestion", "--reopen"):
                subprocess.run([sys.executable, str(Path(__file__).resolve()), stage, directory],
                               cwd=ROOT, env=os.environ.copy(), check=True, timeout=180)
